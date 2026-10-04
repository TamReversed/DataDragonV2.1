"""Tools on the scaffold (datadragon_tools.py): each one against a plain row-by-row reference, and the routes."""
import io
import json
import re
import unicodedata

import openpyxl
import pandas as pd
import pytest

import datadragon_tools as tools
from datadragon_tools import TOOLS, ToolError, describe_change, parse_options


def run(slug, df, **raw):
    tool = TOOLS[slug]
    return tool.run(df, parse_options(tool, raw, df))


MESSY = pd.DataFrame({
    "name": ["  Acme   Corp ", "acme corp", "Émile​  Zola", None, "DON'T, stop!", 12],
    "city": ["new york", "NEW YORK", " Paris", "paris ", None, "Rome"],
    "amount": [1.5, 2.0, None, 4.0, 5.0, 6.0],
})


# ------------------------------------------------------------------------------------------------ Text Cleaner
def reference_clean(value, o):
    if not isinstance(value, str):
        return value
    if o.get("remove_nonprinting", True):
        value = "".join(ch for ch in value if ch in "\t\n\r" or unicodedata.category(ch) not in ("Cc", "Cf"))
    if o.get("trim", True):
        value = value.strip()
    if o.get("collapse_spaces", True):
        value = re.sub(r"\s+", " ", value)
    if o.get("remove_accents"):
        value = "".join(ch for ch in unicodedata.normalize("NFKD", value) if not unicodedata.combining(ch))
    if o.get("remove_punctuation"):
        value = re.sub(r"[^\w\s]", "", value)
    case = o.get("case", "keep")
    if case == "lower":
        value = value.lower()
    elif case == "upper":
        value = value.upper()
    elif case == "title":
        value = value.title()
    return value


@pytest.mark.parametrize("options", [
    {}, {"case": "lower"}, {"case": "upper", "remove_accents": True}, {"case": "title", "remove_punctuation": True},
    {"trim": False, "collapse_spaces": False, "remove_nonprinting": False, "case": "lower"},
])
def test_text_cleaner_matches_a_row_by_row_reference(options):
    result = run("text-cleaner", MESSY, columns=["name", "city"], **options)
    for column in ("name", "city"):
        expected = [reference_clean(v, options) for v in MESSY[column]]
        got = list(result.df[column])
        assert all((a == b) or (pd.isna(a) and pd.isna(b)) for a, b in zip(got, expected)), (column, got, expected)
    assert result.df["amount"].equals(MESSY["amount"]) and len(result.df) == len(MESSY)
    changed = sum(1 for c in ("name", "city") for a, b in zip(MESSY[c], result.df[c]) if isinstance(a, str) and a != b)
    assert dict(result.summary)["Cells changed"] == changed


def test_text_cleaner_does_not_touch_text_that_is_a_number():
    csv_like = pd.DataFrame({"v": ["1,234.50", " 12 ", "-0.5", "45%", "1e3", "a.b", "12 apples"]})    # CSV cells are all text
    result = run("text-cleaner", csv_like, columns=["v"], remove_punctuation=True, case="upper")
    assert list(result.df["v"]) == ["1,234.50", " 12 ", "-0.5", "45%", "1e3", "AB", "12 APPLES"]


def test_fill_keeps_the_form_of_the_column():
    as_numbers = pd.DataFrame({"n": pd.Series([1, None, 4], dtype=object)})            # as read from Excel
    as_text = pd.DataFrame({"n": ["1", None, "4"]})                                     # as read from a CSV
    assert list(run("fill-missing", as_numbers, columns=["n"], method="constant", value="7").df["n"]) == [1, 7, 4]
    assert list(run("fill-missing", as_text, columns=["n"], method="constant", value="7").df["n"]) == ["1", "7", "4"]
    assert list(run("fill-missing", as_numbers, columns=["n"], method="mean").df["n"]) == [1, 2.5, 4]
    assert list(run("fill-missing", as_text, columns=["n"], method="mean").df["n"]) == ["1", "2.5", "4"]


def test_text_cleaner_leaves_numbers_and_blanks_alone_and_needs_a_rule():
    result = run("text-cleaner", MESSY, columns=["name"], case="upper")
    assert result.df["name"].iloc[5] == 12 and pd.isna(result.df["name"].iloc[3])
    with pytest.raises(ToolError, match="at least one thing"):
        run("text-cleaner", MESSY, columns=["name"], trim=False, collapse_spaces=False, remove_nonprinting=False)
    with pytest.raises(ToolError, match="at least one column"):
        run("text-cleaner", MESSY, columns=[])
    assert MESSY["name"].iloc[0] == "  Acme   Corp "                    # the input is not changed in place


# ------------------------------------------------------------------------------------------------ Fill Missing
GAPS = pd.DataFrame({"a": [1.0, None, 3.0, None, 5.0], "b": ["x", "  ", None, "y", "x"], "c": [None, None, 7, 8, 9]})


@pytest.mark.parametrize("method,expected_a", [
    ("previous", [1.0, 1.0, 3.0, 3.0, 5.0]), ("next", [1.0, 3.0, 3.0, 5.0, 5.0]),
    ("mean", [1.0, 3.0, 3.0, 3.0, 5.0]), ("median", [1.0, 3.0, 3.0, 3.0, 5.0]),
])
def test_fill_methods_match_pandas(method, expected_a):
    result = run("fill-missing", GAPS, columns=["a"], method=method)
    assert list(result.df["a"]) == expected_a
    assert dict(result.summary)["Cells filled"] == 2 and result.df["b"].equals(GAPS["b"])


def test_fill_constant_mode_blank_text_and_the_flag_column():
    result = run("fill-missing", GAPS, columns=["a", "b"], method="constant", value="0", flag=True)
    assert list(result.df["a"]) == [1.0, 0, 3.0, 0, 5.0]                 # a typed number stays a number in a numeric column
    assert list(result.df["b"]) == ["x", "0", "0", "y", "x"]            # whitespace-only text counts as blank
    assert list(result.df["filled_columns"]) == [None, "a, b", "b", "a", None]
    assert dict(result.summary)["Cells filled"] == 4
    mode = run("fill-missing", GAPS, columns=["b"], method="mode")
    assert list(mode.df["b"]) == ["x", "x", "x", "y", "x"]


def test_fill_edges_nothing_above_wrong_type_and_dropping_rows():
    up = run("fill-missing", GAPS, columns=["c"], method="previous")
    assert pd.isna(up.df["c"].iloc[0]) and dict(up.summary)["Still blank (nothing to fill from)"] == 2
    with pytest.raises(ToolError, match="not a numeric column"):
        run("fill-missing", GAPS, columns=["b"], method="mean")
    with pytest.raises(ToolError, match="Enter the value"):
        run("fill-missing", GAPS, columns=["a"], method="constant", value="")
    dropped = run("fill-missing", GAPS, columns=["a", "b"], method="drop_rows")
    assert list(dropped.df.index) == [0, 4] and len(dropped.extra_sheets[0][1]) == 3
    assert len(dropped.df) + len(dropped.extra_sheets[0][1]) == len(GAPS)


# ------------------------------------------------------------------------------------------- Remove Duplicates
DUPES = pd.DataFrame({
    "id": [1, 2, 3, 4, 5, 6],
    "name": ["Acme", "acme ", "Beta", "Acme", None, None],
    "note": ["a", None, "c", None, None, "f"],
})


def test_remove_duplicates_reconciles_and_agrees_with_pandas():
    for keep in ("first", "last"):
        result = run("remove-duplicates", DUPES, columns=["name"], keep=keep)
        expected = DUPES[~DUPES.fillna({"name": ""}).duplicated(subset=["name"], keep=keep)]
        assert list(result.df["id"]) == list(expected["id"])
        removed = result.extra_sheets[0][1]
        assert len(result.df) + len(removed) == len(DUPES)
        assert sorted(list(result.df["id"]) + list(removed["id"])) == list(DUPES["id"])


def test_remove_duplicates_loose_matching_most_complete_and_whole_rows():
    loose = run("remove-duplicates", DUPES, columns=["name"], ignore_case_and_spaces=True)
    assert list(loose.df["id"]) == [1, 3, 5]                            # "acme " joins "Acme"; the two blanks are equal
    best = run("remove-duplicates", DUPES, columns=["name"], keep="most_complete", ignore_case_and_spaces=True)
    assert list(best.df["id"]) == [1, 3, 6]                             # 1 has no blanks; 6 has fewer than 5; order kept
    whole = run("remove-duplicates", pd.concat([DUPES, DUPES.iloc[[2]]]), columns=[])
    assert len(whole.df) == 6 and dict(whole.summary)["Duplicate rows removed"] == 1


def test_remove_duplicates_agrees_with_the_duplicate_finder(client, golden_csv=None):
    """Same file, same columns: the finder's "IDs to remove" are exactly the rows this tool removes."""
    df = pd.read_excel(io.BytesIO(client.get("/generate-test-file?num_rows=400").data))
    result = run("remove-duplicates", df, columns=["Vendor_Name", "Department"])
    expected_removed = df[df.duplicated(subset=["Vendor_Name", "Department"], keep="first")]
    assert list(result.extra_sheets[0][1].index) == list(expected_removed.index)


# ----------------------------------------------------------------------------------------- Sort, Rank & Sample
SORT = pd.DataFrame({"g": ["b", "a", "B", None, "a"], "n": ["10", "9", "100", "1", None], "k": [1, 2, 3, 4, 5]})


def test_sort_is_stable_numeric_aware_and_puts_blanks_last():
    by_text = run("sort-and-sample", SORT, sort1="g")
    assert list(by_text.df["k"]) == [2, 5, 1, 3, 4]                     # a, a, b, B (case ignored, ties in file order), blank
    by_number = run("sort-and-sample", SORT, sort1="n", direction1="descending")
    assert list(by_number.df["n"])[:4] == ["100", "10", "9", "1"] and by_number.df["n"].iloc[4] is None
    two = run("sort-and-sample", SORT, sort1="g", sort2="k", direction2="descending")
    assert list(two.df["k"]) == [5, 2, 3, 1, 4]


def test_rank_top_n_every_nth_and_a_repeatable_sample():
    ranked = run("sort-and-sample", SORT, sort1="k", direction1="descending", rank=True)
    assert list(ranked.df["rank"]) == [1, 2, 3, 4, 5] and list(ranked.df["k"]) == [5, 4, 3, 2, 1]
    grouped = run("sort-and-sample", SORT, sort1="k", rank=True, rank_within="g")
    assert list(grouped.df["rank"]) == [1, 1, 1, 1, 2]
    assert list(run("sort-and-sample", SORT, sort1="k", keep="first", count=2).df["k"]) == [1, 2]
    assert list(run("sort-and-sample", SORT, keep="every_nth", count=2).df["k"]) == [1, 3, 5]
    big = pd.DataFrame({"k": range(1000)})
    one = run("sort-and-sample", big, keep="sample", count=50, seed=7)
    two = run("sort-and-sample", big, keep="sample", count=50, seed=7)
    other = run("sort-and-sample", big, keep="sample", count=50, seed=8)
    assert list(one.df["k"]) == list(two.df["k"]) != list(other.df["k"])
    assert len(one.df) == 50 and list(one.df["k"]) == sorted(one.df["k"]) and dict(one.summary)["Sample seed"] == 7
    with pytest.raises(ToolError, match="whole number"):
        run("sort-and-sample", SORT, keep="first")
    with pytest.raises(ToolError, match="chosen twice"):
        run("sort-and-sample", SORT, sort1="g", sort2="g")


# ------------------------------------------------------------------------------------- options and the preview
def test_options_are_checked_against_the_tool_and_the_file():
    tool = TOOLS["fill-missing"]
    with pytest.raises(ToolError, match='"nope" is not in this file'):
        parse_options(tool, {"columns": ["nope"]}, GAPS)
    with pytest.raises(ToolError, match="must be one of"):
        parse_options(tool, {"columns": ["a"], "method": "magic"}, GAPS)
    with pytest.raises(ToolError, match="must be a number"):
        parse_options(TOOLS["sort-and-sample"], {"keep": "first", "count": "many"}, SORT)
    numbered = pd.DataFrame({1: [1], 2: [2]})                           # column names that are not text
    assert parse_options(TOOLS["text-cleaner"], {"columns": ["1"]}, numbered)["columns"] == [1]
    step = tools.step_record(tool, parse_options(tool, {"columns": ["a"], "value": "secret"}, GAPS))
    assert step == {"tool": "fill-missing", "options": {"columns": ["a"], "method": "constant", "flag": False,
                                                        "value": {"literal": True, "length": 6}}}


def test_describe_change_counts_cells_rows_and_samples():
    after = MESSY.copy()
    after.loc[0, "name"] = "Acme Corp"
    after.loc[2, "city"] = "Paris"
    change = describe_change(MESSY, after)
    assert (change["cells_changed"], change["rows_changed"]) == (2, 2)
    assert change["sample"][0] == {"row": 2, "cells": [{"column": "name", "before": "  Acme   Corp ", "after": "Acme Corp"}]}
    assert describe_change(MESSY, MESSY.copy())["cells_changed"] == 0     # blanks equal blanks
    shorter = describe_change(MESSY, MESSY.iloc[:3])
    assert shorter["rows_after"] == 3 and shorter["cells_changed"] is None
    assert describe_change(MESSY, MESSY.iloc[::-1])["reordered"] is True


# ---------------------------------------------------------------------------------------------------- the routes
def upload(df):
    buffer = io.BytesIO()
    df.to_csv(buffer, index=False)
    buffer.seek(0)
    return buffer


@pytest.mark.parametrize("slug", sorted(slug for slug, tool in TOOLS.items() if tool.page))
def test_every_tool_has_a_page_in_the_hub_with_a_help_panel(client, slug):
    page = client.get(f"/{slug}")
    html = page.get_data(as_text=True)
    assert page.status_code == 200 and 'id="toolConfig"' in html and "js/tool.js" in html and 'id="helpModal"' in html
    assert f"('/{slug}', '{TOOLS[slug].name}'" in open("templates/_tools.html", encoding="utf-8").read()
    assert f'href="/{slug}" aria-current="page"' in html


def test_preview_then_run_writes_a_file_with_the_log_and_the_step(client):
    data = {"file": (upload(MESSY), "messy.csv"), "options": json.dumps({"columns": ["name", "city"], "case": "upper"})}
    preview = client.post("/text-cleaner/preview", data=data, content_type="multipart/form-data").get_json()
    assert preview["success"] and preview["change"]["cells_changed"] == 8 and preview["change"]["sample"]
    data = {"file": (upload(MESSY), "messy.csv"), "options": json.dumps({"columns": ["name", "city"], "case": "upper"})}
    done = client.post("/text-cleaner", data=data, content_type="multipart/form-data").get_json()
    assert done["success"] and dict(done["summary"])["Cells changed"] == 8 and done["preview"]["total_rows"] == 6
    workbook = openpyxl.load_workbook(io.BytesIO(client.get(done["download_url"]).data))
    assert workbook.worksheets[0]["A2"].value == "ACME CORP"
    log = {row[0].value: row[1].value for row in workbook["_DataDragon_Log"].iter_rows()}
    assert log["Tool"] == "Text Cleaner" and json.loads(log["Step"])["tool"] == "text-cleaner"
    assert "ACME" not in log["Parameters"] and "Acme" not in log["Step"]          # structure, not cell values
    cached = client.get("/get-cached-files").get_json()["files"]
    assert any(f["source_tool"] == "Text Cleaner" for f in cached)                 # the next tool can pick it up


def test_remove_duplicates_route_writes_both_sheets_and_bad_options_are_a_400(client):
    data = {"file": (upload(DUPES), "d.csv"), "options": json.dumps({"columns": ["name"]})}
    done = client.post("/remove-duplicates", data=data, content_type="multipart/form-data").get_json()
    workbook = openpyxl.load_workbook(io.BytesIO(client.get(done["download_url"]).data))
    assert workbook.sheetnames[:2] == ["Kept rows", "Removed rows"] and done["extra_sheets"] == ["Removed rows"]
    bad = client.post("/fill-missing", data={"file": (upload(GAPS), "g.csv"), "options": json.dumps({"columns": ["zzz"]})},
                      content_type="multipart/form-data")
    assert bad.status_code == 400 and "zzz" in bad.get_json()["error"]
    broken = client.post("/fill-missing", data={"file": (upload(GAPS), "g.csv"), "options": "{not json"},
                         content_type="multipart/form-data")
    assert broken.status_code == 400
