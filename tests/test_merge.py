"""Merge: blank keys never match, statistics are consistent, duplicates are reported."""
import csv
import os
import random
import shutil
from collections import Counter
from queue import Queue

import openpyxl
import pytest

import datadragon
from helpers import make_csv, output_path

JOINS = ("inner", "left", "right", "outer")


def write_rows(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(header)
        w.writerows([["" if v is None else v for v in r] for r in rows])
    return str(path)


_counter = iter(range(10**9))


def run_merge(left, right, join="inner", dup="keep_all", left_key="id", right_key="id", lcols=None, rcols=None):
    q = Queue()
    session = f"merge_test_{next(_counter)}"
    # the tool deletes its inputs when it finishes, so give it private copies
    left_copy = os.path.join(os.path.dirname(left), f"{session}_left.csv")
    right_copy = os.path.join(os.path.dirname(right), f"{session}_right.csv")
    shutil.copy(left, left_copy)
    shutil.copy(right, right_copy)
    datadragon.merge_files_async(left_copy, right_copy, left_key, right_key, join, lcols, rcols, dup, q, session)
    msgs = []
    while not q.empty():
        msgs.append(q.get())
    return msgs[-1]


def merged_rows(final):
    sheet = openpyxl.load_workbook(output_path(final))["Merged Data"]
    return [[c.value for c in row] for row in sheet.iter_rows(min_row=2)]


def reference_join(left, right, how):
    """Obviously-correct nested loops. Rows are (key, value); a None key is blank and never matches."""
    out, used_right, matched = Counter(), set(), 0
    left_with_partner = 0
    for lk, lv in left:
        partners = [i for i, (rk, _) in enumerate(right) if lk is not None and rk is not None and lk == rk]
        if partners:
            left_with_partner += 1
            for i in partners:
                out[(lk, lv, right[i][1])] += 1
                used_right.add(i)
                matched += 1
        elif how in ("left", "outer"):
            out[(lk, lv, None)] += 1
    if how in ("right", "outer"):
        for i, (rk, rv) in enumerate(right):
            if i not in used_right:
                out[(rk, None, rv)] += 1
    return out, matched, len(left) - left_with_partner, len(right) - len(used_right)


@pytest.mark.parametrize("seed", range(40))
def test_matches_reference_join_for_every_join_type(tmp_path, seed):
    rng = random.Random(seed)
    keys = ["1", "2", "3", "x", None]  # None = blank; repeated draws create duplicate keys on both sides
    left = [(rng.choice(keys), f"L{i}") for i in range(rng.randint(0, 9))]
    right = [(rng.choice(keys), f"R{i}") for i in range(rng.randint(0, 9))]
    lpath = write_rows(tmp_path / "l.csv", ["id", "lv"], left)
    rpath = write_rows(tmp_path / "r.csv", ["id", "rv"], right)
    for how in JOINS:
        final = run_merge(lpath, rpath, how)
        assert final["stage"] == "done", (seed, how, final)
        expected, matched, unmatched_left, unmatched_right = reference_join(left, right, how)
        actual = Counter((r[0], r[1], r[2]) for r in merged_rows(final))
        assert actual == expected, (seed, how)
        s = final["summary"]
        assert s["merged_rows"] == sum(expected.values())
        assert s["matched"] == matched
        assert s["unmatched_left"] == unmatched_left and s["unmatched_right"] == unmatched_right
        assert min(s["matched"], s["unmatched_left"], s["unmatched_right"]) >= 0


def test_blank_keys_stay_blank_and_leak_no_placeholders(tmp_path):
    left = write_rows(tmp_path / "l.csv", ["id", "lv"], [[None, "a"], [None, "b"], ["1", "c"]])
    right = write_rows(tmp_path / "r.csv", ["id", "rv"], [[None, "x"], ["1", "y"]])
    final = run_merge(left, right, "outer")
    rows = merged_rows(final)
    assert Counter(map(tuple, rows)) == Counter([("1", "c", "y"), (None, "a", None), (None, "b", None),
                                                 (None, None, "x")])
    assert "BlankKey" not in str(rows)


def test_different_key_names_restore_blanks_in_both_key_columns(tmp_path):
    left = write_rows(tmp_path / "l.csv", ["a", "lv"], [["1", "p"], [None, "q"]])
    right = write_rows(tmp_path / "r.csv", ["b", "rv"], [["1", "x"], [None, "y"]])
    final = run_merge(left, right, "outer", left_key="a", right_key="b")
    sheet = openpyxl.load_workbook(output_path(final))["Merged Data"]
    header = [c.value for c in sheet[1]]
    assert header[:2] == ["a", "lv"] and "b" in header
    data = [dict(zip(header, [c.value for c in row])) for row in sheet.iter_rows(min_row=2)]
    assert {"a": "1", "b": "1", "lv": "p", "rv": "x"} in data
    assert {"a": None, "b": None, "lv": "q", "rv": None} in data
    assert {"a": None, "b": None, "lv": None, "rv": "y"} in data


def test_repeated_keys_are_reported_as_row_multiplication(tmp_path):
    left = write_rows(tmp_path / "l.csv", ["id", "lv"], [["1", "a"], ["2", "b"]])
    right = write_rows(tmp_path / "r.csv", ["id", "rv"], [["1", "x"], ["1", "y"], ["1", "z"], ["2", "w"]])
    final = run_merge(left, right, "inner", "keep_all")
    assert final["summary"]["merged_rows"] == 4 and final["summary"]["multiplication_factor"] == 2.0
    assert "repeated" not in final["warning"] and "more than once in the right file" in final["warning"]
    assert "2 left rows with a partner produced 4 merged rows" in final["warning"]
    assert "'keep first'" in final["warning"]


def test_right_join_measures_multiplication_against_the_right_file(tmp_path):
    left = write_rows(tmp_path / "l.csv", ["id", "lv"], [["1", "a"], ["1", "b"]])
    right = write_rows(tmp_path / "r.csv", ["id", "rv"], [["1", "x"]])
    final = run_merge(left, right, "right", "keep_all")
    assert "more than once in the left file" in final["warning"]
    assert "1 right rows with a partner produced 2 merged rows" in final["warning"]


def test_many_to_one_lookup_is_not_flagged(tmp_path):
    orders = write_rows(tmp_path / "o.csv", ["cust", "order"], [["c1", "o1"], ["c1", "o2"], ["c2", "o3"], ["c1", "o4"]])
    customers = write_rows(tmp_path / "c.csv", ["cust", "name"], [["c1", "Ann"], ["c2", "Bo"]])
    final = run_merge(orders, customers, "left", "keep_all", left_key="cust", right_key="cust")
    assert final["summary"]["merged_rows"] == 4 and final["summary"]["multiplication_factor"] == 1.0
    assert "warning" not in final


def test_one_to_one_merge_has_no_warning(tmp_path):
    left = write_rows(tmp_path / "l.csv", ["id", "lv"], [["1", "a"], ["2", "b"], [None, "c"]])
    right = write_rows(tmp_path / "r.csv", ["id", "rv"], [["1", "x"], ["3", "y"], [None, "z"]])
    final = run_merge(left, right, "inner")
    assert "warning" not in final and final["summary"]["multiplication_factor"] == 1.0


def test_keep_first_drops_repeated_keys_but_keeps_every_blank_key_row(tmp_path):
    left = write_rows(tmp_path / "l.csv", ["id", "lv"], [["1", "a"], ["1", "b"], [None, "c"], [None, "d"]])
    right = write_rows(tmp_path / "r.csv", ["id", "rv"], [["1", "x"]])
    final = run_merge(left, right, "left", "keep_first")
    assert Counter(map(tuple, merged_rows(final))) == Counter([("1", "a", "x"), (None, "c", None), (None, "d", None)])
    assert final["summary"]["left_rows"] == 3


def test_error_mode_ignores_blank_keys_but_catches_real_duplicates(tmp_path):
    blanks = write_rows(tmp_path / "l.csv", ["id", "lv"], [["1", "a"], [None, "b"], [None, "c"]])
    right = write_rows(tmp_path / "r.csv", ["id", "rv"], [["1", "x"], [None, "y"], [None, "z"]])
    assert run_merge(blanks, right, "inner", "error")["stage"] == "done"
    dupes = write_rows(tmp_path / "d.csv", ["id", "lv"], [["1", "a"], ["1", "b"]])
    final = run_merge(dupes, right, "inner", "error")
    assert final["stage"] == "error" and "Duplicate keys found" in final["message"]


def test_selected_columns_still_always_include_the_key(tmp_path):
    left = write_rows(tmp_path / "l.csv", ["id", "keep", "drop"], [["1", "k", "d"]])
    right = write_rows(tmp_path / "r.csv", ["id", "rv"], [["1", "x"]])
    final = run_merge(left, right, "inner", lcols=["keep"], rcols=None)
    sheet = openpyxl.load_workbook(output_path(final))["Merged Data"]
    assert [c.value for c in sheet[1]] == ["id", "keep", "rv"]
