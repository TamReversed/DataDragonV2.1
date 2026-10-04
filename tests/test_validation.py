"""Validation rules: one mask per rule, spreadsheet row numbers, failure counts, honest truncation, and agreement
with a plain row-by-row reference."""
import json
import random

import openpyxl

import datadragon
from helpers import make_csv, output_path, post_job


def run(client, tmp_path, text, rules, name="v.csv"):
    path = make_csv(tmp_path / name, text)
    final, _ = post_job(client, "/validate-data", {"file": path}, {"validation_rules": json.dumps(rules)})
    return final


def sheet(final, name):
    wb = openpyxl.load_workbook(output_path(final))
    return [list(r) for r in wb[name].iter_rows(values_only=True)]


def test_row_numbers_are_spreadsheet_rows_and_failures_are_counted(client, tmp_path):
    final = run(client, tmp_path, "a,b\n1,x\n,y\n3,\n", [
        {"column": "a", "type": "required"}, {"column": "b", "type": "required"}])
    assert final["summary"]["invalid_rows"] == 2
    assert final["summary"]["errors_total"] == 2
    details = sheet(final, "Error Details")
    assert [r[0] for r in details[1:]] == [3, 4]        # header is row 1, so data row 1 is row 2


def test_one_row_failing_two_rules_counts_both(client, tmp_path):
    final = run(client, tmp_path, "a\nabcdef\n", [
        {"column": "a", "type": "numeric"}, {"column": "a", "type": "length", "value": {"min": None, "max": 3}}])
    assert final["summary"]["invalid_rows"] == 1 and final["summary"]["errors_total"] == 2


def test_range_rejects_non_numbers_and_accepts_string_bounds(client, tmp_path):
    final = run(client, tmp_path, "q\n5\nabc\n2\n\n", [{"column": "q", "type": "range", "value": {"min": "1", "max": "3"}}])
    assert final["summary"]["invalid_rows"] == 2           # 5 too big, 'abc' not a number; blank is not checked


def test_bad_bound_gives_a_clear_error(client, tmp_path):
    final = run(client, tmp_path, "q\n5\n", [{"column": "q", "type": "range", "value": {"min": "low"}}])
    assert final["stage"] == "error" and "not a number" in final["message"]


def test_list_compares_as_text_and_blank_is_not_checked(client, tmp_path):
    final = run(client, tmp_path, "c\n1\n2\n\nx\n", [{"column": "c", "type": "list", "value": [1, "x"]}])
    assert final["summary"]["invalid_rows"] == 1


def test_details_are_truncated_with_a_note(client, tmp_path, monkeypatch):
    monkeypatch.setattr(datadragon, "VALIDATION_DETAIL_LIMIT", 3)
    final = run(client, tmp_path, "a\n" + "\n".join(["x"] * 8) + "\n", [{"column": "a", "type": "numeric"}])
    assert final["summary"]["invalid_rows"] == 8 and final["summary"]["details_truncated"] is True
    assert len(sheet(final, "Error Details")) == 1 + 3
    assert any("first 3" in str(r[1]) for r in sheet(final, "Validation Summary"))


def test_matches_a_row_by_row_reference(client, tmp_path):
    rng = random.Random(7)
    words = ["12345", "123456", "ab", "", "7", "-3", "3.5", "N", "S", "x y"]
    rows = [[rng.choice(words) for _ in range(3)] for _ in range(200)]
    text = "a,b,c\n" + "\n".join(",".join(r) for r in rows) + "\n"
    rules = [
        {"column": "a", "type": "pattern", "value": r"\d{5}"},
        {"column": "b", "type": "range", "value": {"min": 0, "max": 5}},
        {"column": "c", "type": "list", "value": ["N", "S"]},
        {"column": "a", "type": "length", "value": {"min": 2, "max": 5}},
    ]

    def reference(row):
        a, b, c = row
        bad = 0
        if a != "" and not (a.isdigit() and len(a) == 5):
            bad += 1
        if b != "":
            try:
                if not 0 <= float(b) <= 5:
                    bad += 1
            except ValueError:
                bad += 1
        if c != "" and c not in ("N", "S"):
            bad += 1
        if a != "" and not 2 <= len(a) <= 5:
            bad += 1
        return bad

    expected = [reference(r) for r in rows]
    final = run(client, tmp_path, text, rules)
    assert final["summary"]["invalid_rows"] == sum(1 for e in expected if e)
    assert final["summary"]["errors_total"] == sum(expected)


def test_big_valid_sheet_is_omitted_with_a_note(client, tmp_path, monkeypatch):
    monkeypatch.setattr(datadragon, "VALIDATION_VALID_SHEET_LIMIT", 2)
    final = run(client, tmp_path, "a\n1\n2\n3\nx\n", [{"column": "a", "type": "numeric"}])
    wb = openpyxl.load_workbook(output_path(final))
    assert "Valid Records" not in wb.sheetnames and final["summary"]["valid_sheet_omitted"] is True
    assert any("omitted" in str(r[1]) for r in sheet(final, "Validation Summary"))
