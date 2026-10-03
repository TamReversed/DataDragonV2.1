"""Regression tests for the findings of the Phase 2 independent review."""
import json
import os
import time
from queue import Queue

import openpyxl
import pandas as pd

import datadragon
from helpers import make_csv, output_path, post_form, post_job
from test_column_ops import op, rows
from test_pipeline_flow import run_stage, start


def sheet_rows(path, name):
    return [[c.value for c in r] for r in openpyxl.load_workbook(path)[name].iter_rows()]


# --- XSS checker ------------------------------------------------------------------------------
def test_checker_no_longer_trusts_bare_join_results():
    import importlib.util
    spec = importlib.util.spec_from_file_location("chk", "scripts/check_innerhtml.py")
    chk = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(chk)
    assert not chk.is_safe("state.selectedKey.join(' + ')", "")
    assert chk.is_safe("escapeHtml(state.selectedKey.join(' + '))", "")
    assert not chk.is_safe("data.columns", "")


def test_pipeline_summary_escapes_the_selected_key():
    html = open("templates/data_readiness_pipeline.html", encoding="utf-8").read()
    assert "${escapeHtml(state.selectedKey.join(' + ')" in html


# --- pivot ------------------------------------------------------------------------------------
def test_pivot_with_infinite_values_still_finishes(client, tmp_path):
    path = make_csv(tmp_path / "i.csv", "g,n\na,inf\nb,1\n")
    final, _ = post_job(client, "/generate-pivot", {"file": path}, {"rows[]": ["g"], "values[]": ["n"], "aggfunc": "sum"})
    assert final["stage"] == "done", final


def test_pivot_keeps_date_row_labels_as_dates(client, tmp_path):
    from datetime import datetime
    path = tmp_path / "d.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["d", "n"])
    ws.append([datetime(2020, 1, 1), 5])
    ws.append([datetime(2020, 1, 2), 7])
    wb.save(path)
    final, _ = post_job(client, "/generate-pivot", {"file": str(path)}, {"rows[]": ["d"], "values[]": ["n"], "aggfunc": "sum"})
    assert final["stage"] == "done", final
    cell = openpyxl.load_workbook(output_path(final))["Pivot Table"]["A2"]
    assert cell.is_date and cell.value == datetime(2020, 1, 1)


# --- CSV exports ------------------------------------------------------------------------------
def test_csv_headers_and_category_columns_are_neutralised():
    df = pd.DataFrame({"=1+1": ["a", "b"], "cat": pd.Categorical(["=2+2", "x"]), "s": pd.array(["@x", "y"], dtype="string"),
                       "n": [-5, 3]})
    out = datadragon.sanitize_csv(df)
    assert list(out.columns) == ["'=1+1", "cat", "s", "n"]
    assert list(out["cat"]) == ["'=2+2", "x"] and list(out["s"]) == ["'@x", "y"] and list(out["n"]) == [-5, 3]


# --- Unique-ID export --------------------------------------------------------------------------
def test_unique_id_export_does_not_overwrite_user_columns(client, tmp_path):
    import zipfile, io
    path = make_csv(tmp_path / "u.csv", "Unique_ID,_is_duplicate,k\nq,x,1\nr,y,2\n")
    final, _ = post_job(client, "/find-unique-identifier", {"file": path}, {"selected_columns[]": ["k"]})
    assert final["stage"] == "done", final
    z = zipfile.ZipFile(output_path(final))
    data = [n for n in z.namelist() if n.endswith("_data.xlsx")][0]
    ws = openpyxl.load_workbook(io.BytesIO(z.read(data)))["Data with Unique IDs"]
    table = [[c.value for c in r] for r in ws.iter_rows()]
    assert table[0] == ["Unique_ID", "_is_duplicate", "k", "Unique_ID_", "_is_duplicate_"]
    assert [r[0] for r in table[1:]] == ["q", "r"] and [r[1] for r in table[1:]] == ["x", "y"]


# --- cache / queues / uploads --------------------------------------------------------------------
def test_eviction_from_the_cache_keeps_the_download_alive(client, tmp_path):
    path = make_csv(tmp_path / "f.csv", "a\n1\n2\n")
    urls = []
    for _ in range(11):
        resp, payload = post_form(client, "/row-filter", {"file": path},
                                  {"conditions": json.dumps([{"column": "a", "operator": "is_not_empty"}])})
        assert resp.status_code == 200, payload
        urls.append(payload["download_url"])
    assert all(client.get(u).status_code == 200 for u in urls)


def test_a_job_that_keeps_reporting_is_not_swept_as_abandoned():
    q = Queue()
    datadragon.register_job("busy_job", q, owner="owner-x")
    datadragon.progress_queue_created["busy_job"] = time.time() - 11 * 60
    q.put({"stage": "working"})                      # activity refreshes the channel
    datadragon.cleanup_progress_queues()
    assert "busy_job" in datadragon.progress_queues
    datadragon.progress_queue_created["busy_job"] = time.time() - 11 * 60   # silent for 11 minutes
    datadragon.cleanup_progress_queues()
    assert "busy_job" not in datadragon.progress_queues


def test_pipeline_upload_is_discarded_after_loading(client, tmp_path):
    before = set(os.listdir(datadragon.app.config["UPLOAD_FOLDER"]))
    start(client, make_csv(tmp_path / "p.csv", "a\n1\n"))
    assert set(os.listdir(datadragon.app.config["UPLOAD_FOLDER"])) == before


def test_allow_null_keys_text_false_means_false(client, tmp_path):
    sid = start(client, make_csv(tmp_path / "k.csv", "id,v\n1,a\n2,\n"))
    run_stage(client, sid, "analyze")
    run_stage(client, sid, "keys", selected_columns=["id", "v"], allow_null_keys="false")
    assert datadragon.pipeline_sessions[sid].stage_data[3]["excluded_null_columns"] == ["v"]


def test_no_wasted_full_data_search_when_duplicates_exist(client, tmp_path):
    sid = start(client, make_csv(tmp_path / "d.csv", "id\n1\n1\n2\n"))
    run_stage(client, sid, "analyze")
    run_stage(client, sid, "keys", selected_columns=["id"])
    res = datadragon.pipeline_sessions[sid].stage_data[3]
    assert res["minimal_combinations"] == [] and res["minimal_combinations_after_dedup"] == [["id"]]


# --- Column Operations input validation -----------------------------------------------------------
def test_column_ops_reject_wrong_json_types(client, tmp_path):
    for operation, field, value in [("rename", "renames", ["a"]), ("rename", "renames", {"a": None}),
                                    ("rename", "renames", {"a": 5}), ("reorder", "column_order", [["a"]]),
                                    ("reorder", "column_order", "ab"), ("delete", "columns_to_delete", [["a"]]),
                                    ("merge", "columns_to_merge", [["a"], ["b"]])]:
        resp, payload = op(client, tmp_path, "a,b\n1,2\n", operation, **{field: json.dumps(value),
                                                                         "new_column_name": "m"})
        assert resp.status_code == 400, (operation, value, payload)


def test_column_ops_on_a_header_only_file(client, tmp_path):
    resp, payload = op(client, tmp_path, "a,b\n", "split", column_to_split="a", delimiter=",")
    assert resp.status_code == 200 and rows(payload)[0] == ["a_part1", "b"]
    resp, payload = op(client, tmp_path, "a,b\n", "merge", columns_to_merge=json.dumps(["a", "b"]),
                       separator="-", new_column_name="m")
    assert resp.status_code == 200, payload
