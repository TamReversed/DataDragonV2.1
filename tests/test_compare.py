"""Compare: rows are matched on real keys (nth occurrence with nth), blank keys are set aside, values are
compared across file types, and the cost stays near-linear."""
import os
import random
import shutil
import time
from collections import Counter
from queue import Queue

import openpyxl
import pytest

import datadragon
from helpers import make_xlsx, output_path
from test_merge import write_rows

_counter = iter(range(10**9))


def run_compare(file1, file2, keys, compare_columns=None):
    q = Queue()
    session = f"compare_test_{next(_counter)}"
    copies = []
    for tag, path in (("one", file1), ("two", file2)):  # the tool deletes its inputs when it finishes
        copy = os.path.join(os.path.dirname(path), f"{session}_{tag}{os.path.splitext(path)[1]}")
        shutil.copy(path, copy)
        copies.append(copy)
    datadragon.compare_files_async(copies[0], copies[1], keys, compare_columns, q, session)
    msgs = []
    while not q.empty():
        msgs.append(q.get())
    return msgs[-1]


def sheet(final, name):
    wb = openpyxl.load_workbook(output_path(final))
    if name not in wb.sheetnames:
        return []
    return [[c.value for c in row] for row in wb[name].iter_rows(min_row=2)]


def reference(rows1, rows2):
    """rows are (key, v1, v2); a None key is blank = no key. Pair the nth occurrence of a key with the nth."""
    def occurrences(rows):
        seen, out = Counter(), []
        for i, (k, *vals) in enumerate(rows):
            if k is None:
                continue
            out.append(((k, seen[k]), i, tuple(vals)))
            seen[k] += 1
        return out
    a, b = occurrences(rows1), occurrences(rows2)
    b_by_id = {ident: (i, vals) for ident, i, vals in b}
    a_ids = {ident for ident, _, _ in a}
    unchanged, changed, removed, added = [], [], [], []
    for ident, i, vals in a:
        if ident in b_by_id:
            (changed if vals != b_by_id[ident][1] else unchanged).append((ident[0],) + vals)
        else:
            removed.append((ident[0],) + vals)
    for ident, i, vals in b:
        if ident not in a_ids:
            added.append((ident[0],) + vals)
    return unchanged, changed, removed, added


@pytest.mark.parametrize("seed", range(40))
def test_matches_reference(tmp_path, seed):
    rng = random.Random(seed)
    keys, vals = ["1", "2", "3", "x", None], ["a", "b", None]
    rows1 = [(rng.choice(keys), rng.choice(vals), rng.choice(vals)) for _ in range(rng.randint(0, 10))]
    rows2 = [(rng.choice(keys), rng.choice(vals), rng.choice(vals)) for _ in range(rng.randint(0, 10))]
    f1 = write_rows(tmp_path / "a.csv", ["k", "v1", "v2"], rows1)
    f2 = write_rows(tmp_path / "b.csv", ["k", "v1", "v2"], rows2)
    final = run_compare(f1, f2, ["k"])
    assert final["stage"] == "done", final
    unchanged, changed, removed, added = reference(rows1, rows2)
    s = final["summary"]
    assert (s["common"], s["unchanged"], s["changed"], s["removed"], s["added"]) == (
        len(unchanged) + len(changed), len(unchanged), len(changed), len(removed), len(added))
    assert Counter(map(tuple, sheet(final, "Unchanged Rows"))) == Counter(unchanged)
    assert Counter(map(tuple, sheet(final, "Removed Rows"))) == Counter(removed)
    assert Counter(map(tuple, sheet(final, "Added Rows"))) == Counter(added)
    assert s["rows_without_key_file1"] == sum(1 for r in rows1 if r[0] is None)
    assert s["rows_without_key_file2"] == sum(1 for r in rows2 if r[0] is None)


def test_output_order_follows_the_source_files(tmp_path):
    f1 = write_rows(tmp_path / "a.csv", ["k", "v"], [["c", "1"], ["a", "1"], ["b", "1"], ["gone2", "1"], ["gone1", "1"]])
    f2 = write_rows(tmp_path / "b.csv", ["k", "v"], [["new2", "1"], ["b", "1"], ["a", "1"], ["c", "1"], ["new1", "1"]])
    final = run_compare(f1, f2, ["k"])
    assert [r[0] for r in sheet(final, "Unchanged Rows")] == ["c", "a", "b"]       # file 1 order
    assert [r[0] for r in sheet(final, "Removed Rows")] == ["gone2", "gone1"]      # file 1 order
    assert [r[0] for r in sheet(final, "Added Rows")] == ["new2", "new1"]          # file 2 order


def test_changed_rows_keep_the_old_new_layout(tmp_path):
    f1 = write_rows(tmp_path / "a.csv", ["id", "name", "qty", "note"], [["1", "x", "5", "same"]])
    f2 = write_rows(tmp_path / "b.csv", ["id", "name", "qty", "note"], [["1", "x", "6", "same"]])
    final = run_compare(f1, f2, ["id"])
    wb = openpyxl.load_workbook(output_path(final))
    header = [c.value for c in wb["Changed Rows"][1]]
    assert header == ["id", "name", "qty (Old)", "qty (New)", "note"]
    assert [c.value for c in wb["Changed Rows"][2]] == ["1", "x", "5", "6", "same"]


def test_repeated_keys_are_paired_in_order_and_reported(tmp_path):
    f1 = write_rows(tmp_path / "a.csv", ["k", "v"], [["1", "a"], ["1", "b"], ["1", "c"]])
    f2 = write_rows(tmp_path / "b.csv", ["k", "v"], [["1", "a"], ["1", "CHANGED"]])
    final = run_compare(f1, f2, ["k"])
    s = final["summary"]
    assert (s["common"], s["unchanged"], s["changed"], s["removed"], s["added"]) == (2, 1, 1, 1, 0)
    assert s["repeated_key_rows_file1"] == 2 and s["repeated_key_rows_file2"] == 1
    assert "File 1 has 2 row(s) whose key repeats" in final["warning"]
    assert [r[0] for r in sheet(final, "Removed Rows")] == ["1"] and sheet(final, "Removed Rows")[0][1] == "c"


def test_rows_without_a_key_are_set_aside_not_called_added_or_removed(tmp_path):
    f1 = write_rows(tmp_path / "a.csv", ["k", "v"], [["1", "a"], [None, "b"]])
    f2 = write_rows(tmp_path / "b.csv", ["k", "v"], [["1", "a"], [None, "b"], [None, "c"]])
    final = run_compare(f1, f2, ["k"])
    s = final["summary"]
    assert (s["common"], s["added"], s["removed"]) == (1, 0, 0)
    assert (s["rows_without_key_file1"], s["rows_without_key_file2"]) == (1, 2)
    assert "1 row(s) in file 1 and 2 row(s) in file 2 have no key value" in final["warning"]
    assert [r[:3] for r in sheet(final, "Rows Without Key")] == [["File 1", None, "b"], ["File 2", None, "b"],
                                                                 ["File 2", None, "c"]]


def test_composite_key_with_one_blank_part_is_still_a_key(tmp_path):
    f1 = write_rows(tmp_path / "a.csv", ["k1", "k2", "v"], [["1", "x", "a"], [None, "y", "b"], [None, None, "z"]])
    f2 = write_rows(tmp_path / "b.csv", ["k1", "k2", "v"], [["1", "x", "a"], [None, "y", "CHANGED"], [None, None, "z"]])
    final = run_compare(f1, f2, ["k1", "k2"])
    s = final["summary"]
    assert (s["common"], s["unchanged"], s["changed"], s["added"], s["removed"]) == (2, 1, 1, 0, 0)
    assert s["rows_without_key_file1"] == 1 and s["rows_without_key_file2"] == 1  # both parts blank: no key


def test_values_are_compared_across_file_types(tmp_path):
    import datetime as dt
    xl = make_xlsx(tmp_path / "a.xlsx", ["k", "qty", "price", "when", "flag"],
                   [["1", 5, 2.5, dt.datetime(2024, 1, 5), True], ["2", 7, 3.0, dt.datetime(2024, 2, 1), False]])
    cv = write_rows(tmp_path / "b.csv", ["k", "qty", "price", "when", "flag"],
                    [["1", "5", "2.5", "2024-01-05", "True"], ["2", "8", "3", "2024-02-01", "False"]])
    final = run_compare(xl, cv, ["k"])
    s = final["summary"]
    assert (s["unchanged"], s["changed"]) == (1, 1)
    wb = openpyxl.load_workbook(output_path(final))
    header = [c.value for c in wb["Changed Rows"][1]]
    # only qty differs; price (2.5 vs '2.5', 3.0 vs '3'), the date and the boolean match across Excel and CSV
    assert header == ["k", "qty (Old)", "qty (New)", "price", "when", "flag"]
    assert [c.value for c in wb["Changed Rows"][2]] == ["2", 7, "8", 3, dt.datetime(2024, 2, 1), False]


def test_blank_equals_blank_and_blank_differs_from_a_value(tmp_path):
    f1 = write_rows(tmp_path / "a.csv", ["k", "v"], [["1", None], ["2", None], ["3", "x"]])
    f2 = write_rows(tmp_path / "b.csv", ["k", "v"], [["1", None], ["2", "now set"], ["3", None]])
    s = run_compare(f1, f2, ["k"])["summary"]
    assert (s["unchanged"], s["changed"]) == (1, 2)


def test_compare_only_the_selected_columns(tmp_path):
    f1 = write_rows(tmp_path / "a.csv", ["k", "v", "ignored"], [["1", "a", "x"]])
    f2 = write_rows(tmp_path / "b.csv", ["k", "v", "ignored"], [["1", "a", "DIFFERENT"]])
    assert run_compare(f1, f2, ["k"], ["v"])["summary"]["changed"] == 0
    assert run_compare(f1, f2, ["k"])["summary"]["changed"] == 1


def test_missing_key_column_is_an_error(tmp_path):
    f1 = write_rows(tmp_path / "a.csv", ["k", "v"], [["1", "a"]])
    f2 = write_rows(tmp_path / "b.csv", ["other", "v"], [["1", "a"]])
    final = run_compare(f1, f2, ["k"])
    assert final["stage"] == "error" and "not found in file 2" in final["message"]


def test_scales_near_linearly(tmp_path):
    """The old implementation scanned the whole frame once per key (100k rows took 335 s)."""
    n = 30000
    f1 = write_rows(tmp_path / "a.csv", ["id", "amount", "status"], [[f"R{i}", i, "ok"] for i in range(n)])
    f2 = write_rows(tmp_path / "b.csv", ["id", "amount", "status"],
                    [[f"R{i}", i + (1 if i % 10 == 0 else 0), "ok"] for i in range(n)])
    started = time.time()
    final = run_compare(f1, f2, ["id"])
    elapsed = time.time() - started
    assert final["summary"]["changed"] == n // 10 and final["summary"]["unchanged"] == n - n // 10
    assert elapsed < 8, f"compare of {n} rows took {elapsed:.1f}s"
