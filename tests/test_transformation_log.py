"""T3.11: every output carries a record of what produced it."""
import io
import json
import zipfile

import openpyxl
import pytest

import datadragon
from helpers import fixture_path, make_csv, output_path, post_form, post_job

KEYS = ["Tool", "Version", "Timestamp (UTC)", "Rows in", "Rows out", "Parameters", "Warnings"]


def log_of(path_or_bytes):
    source = io.BytesIO(path_or_bytes) if isinstance(path_or_bytes, bytes) else path_or_bytes
    wb = openpyxl.load_workbook(source)
    assert wb.sheetnames[-1] == datadragon.LOG_SHEET            # the log is always the last sheet
    return {row[0].value: row[1].value for row in list(wb[datadragon.LOG_SHEET].iter_rows())[1:]}


def check_shape(log, tool):
    assert list(log) == KEYS and log["Tool"] == tool
    assert log["Version"].startswith("DataDragon ") and "(" in log["Version"]
    assert log["Timestamp (UTC)"][:2] == "20" and len(log["Timestamp (UTC)"]) == 19
    json.loads(log["Parameters"])


def test_find_replace_output_has_the_log(client, tmp_path):
    path = make_csv(tmp_path / "f.csv", "k,v\n1,foo\n2,bar\n3,foo\n")
    resp, payload = post_form(client, "/find-replace", {"file": path},
                              {"find_text": "foo", "replace_text": "baz", "column": "v"})
    log = log_of(output_path(payload))
    check_shape(log, "Find & Replace")
    params = json.loads(log["Parameters"])
    assert (log["Rows in"], log["Rows out"]) == ("3", "3") and params["find_text"] == "foo" and params["replacements_made"] == 2


def test_merge_output_has_the_log_after_its_data_sheets(client):
    final, _ = post_job(client, "/merge-data",
                        {"left_file": fixture_path("merge_left.csv"), "right_file": fixture_path("merge_right.csv")},
                        {"left_key": "id", "right_key": "id", "join_type": "inner", "duplicate_handling": "keep_all"})
    wb = openpyxl.load_workbook(output_path(final))
    assert wb.sheetnames == ["Merged Data", "Join Summary", datadragon.LOG_SHEET]
    log = log_of(output_path(final))
    check_shape(log, "Data Merge")
    assert json.loads(log["Parameters"])["join_type"] == "inner"


def test_pivot_output_has_the_log(client, tmp_path):
    path = make_csv(tmp_path / "p.csv", "g,n\na,1\nb,2\na,3\n")
    final, _ = post_job(client, "/generate-pivot", {"file": path}, {"rows[]": ["g"], "values[]": ["n"], "aggfunc": "sum"})
    log = log_of(output_path(final))
    check_shape(log, "Pivot Table Generator")
    assert json.loads(log["Parameters"])["aggfunc"] == "sum" and log["Rows in"] == "3"


def test_validation_and_transpose_and_schema_comparison_have_it(client, tmp_path):
    path = make_csv(tmp_path / "v.csv", "a\n1\nx\n")
    final, _ = post_job(client, "/validate-data", {"file": path},
                        {"validation_rules": json.dumps([{"column": "a", "type": "numeric"}])})
    check_shape(log_of(output_path(final)), "Data Validation")
    final, _ = post_job(client, "/transpose-data", {"file": path})
    check_shape(log_of(output_path(final)), "Transpose")
    other = make_csv(tmp_path / "o.csv", "a,b\n1,2\n")
    final, _ = post_job(client, "/compare-columns", {"file1": path, "file2": other})
    check_shape(log_of(output_path(final)), "Schema Comparison")


def test_the_log_records_warnings(client, tmp_path, monkeypatch):
    monkeypatch.setattr(datadragon, "EXCEL_MAX_ROWS", 1000)
    path = make_csv(tmp_path / "w.csv", "a\n1\n")
    datadragon.note_warning("job_warn_1", "something to know")
    assert datadragon.make_log("X", 1, 1, {}, "job_warn_1")["Warnings"] == "something to know"


def test_secret_looking_parameters_are_dropped_and_values_are_not_logged():
    log = datadragon.make_log("T", 1, 1, {"api_key": "k", "Password": "p", "token": "t", "column": "ok",
                                          "nested": {"client_secret": "s", "keep": 1}})
    params = json.loads(log["Parameters"])
    assert params == {"column": "ok", "nested": {"keep": 1}}


def test_csv_anonymizer_output_gets_a_log_sidecar_and_a_link(client, tmp_path):
    path = make_csv(tmp_path / "s.csv", "Name,Email\nAnn,a@x.org\nBob,b@x.org\n")
    final, _ = post_job(client, "/scrub-data", {"file": path}, {"columns[]": ["Name", "Email"], "relationship_preserve": "false"})
    assert final["stage"] == "done" and final["output_filename"].endswith(".csv")
    assert final["log_url"].endswith(".log.json")
    sidecar = client.get(final["log_url"])
    assert sidecar.status_code == 200
    log = json.loads(sidecar.get_data(as_text=True))
    assert log["Tool"] == "Data Anonymizer" and "Ann" not in json.dumps(log) and "a@x.org" not in json.dumps(log)


def test_oversize_single_table_falls_back_to_csv_with_a_sidecar(client, tmp_path, monkeypatch):
    monkeypatch.setattr(datadragon, "EXCEL_MAX_ROWS", 10)
    path = make_csv(tmp_path / "big.csv", "a\n" + "\n".join(str(i) for i in range(30)) + "\n")
    resp, payload = post_form(client, "/column-operations", {"file": path},
                              {"operation": "duplicate", "source_column": "a", "new_column_name": "b"})
    assert payload["filename"].endswith(".csv") and payload["log_url"].endswith(".log.json")
    assert json.loads(client.get(payload["log_url"]).get_data(as_text=True))["Tool"] == "Column Operations"


def test_splitter_zip_has_the_record_but_the_chunks_stay_clean(client):
    final, _ = post_job(client, "/upload", {"file": fixture_path("golden.xlsx")}, {"chunk_size": "15", "base_filename": "g"})
    z = zipfile.ZipFile(output_path(final))
    for name in z.namelist():
        if name.endswith(".xlsx"):
            assert openpyxl.load_workbook(io.BytesIO(z.read(name))).sheetnames == ["Sheet1"]
    log = json.loads(z.read("_DataDragon_Log.json"))
    assert log["Tool"] == "File Splitter" and json.loads(log["Parameters"])["files"] == 3


def test_a_result_with_a_log_sheet_is_not_reported_as_multi_sheet_when_reused(client, tmp_path):
    """Chaining tools: the second tool must not warn "has 2 sheets" about our own log sheet."""
    path = make_csv(tmp_path / "c.csv", "a,b\n1,2\n")
    resp, first = post_form(client, "/column-operations", {"file": path},
                            {"operation": "duplicate", "source_column": "a", "new_column_name": "c"})
    cache_id = client.get("/get-cached-files").get_json()["files"][0]["cache_id"]
    resp, second = post_form(client, "/column-operations", None,
                             {"cache_id": cache_id, "operation": "rename", "renames": json.dumps({"a": "z"})})
    assert resp.status_code == 200 and "warning" not in second and not second.get("sheet_names")


def test_reading_a_result_ignores_the_log_sheet(tmp_path):
    path = tmp_path / "r.xlsx"
    datadragon.write_excel(__import__("pandas").DataFrame({"a": [1, 2]}), str(path), log=datadragon.make_log("T", 2, 2))
    df = datadragon.read_data_file(str(path))
    assert list(df.columns) == ["a"] and len(df) == 2
    assert datadragon.sheet_names_of(str(path)) == ["Sheet1"]
