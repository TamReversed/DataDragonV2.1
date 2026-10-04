"""Phase B of the toolset plan: Unpivot, Group & Summarise, delimiter detection, Open a File, Append Files."""
import io
import json

import openpyxl
import pandas as pd
import pytest

import datadragon
from datadragon_tools import TOOLS, ToolError, append_tables, parse_options


def run(slug, df, **raw):
    tool = TOOLS[slug]
    return tool.run(df, parse_options(tool, raw, df))


def upload(text, name="f.csv"):
    return (io.BytesIO(text.encode("utf-8") if isinstance(text, str) else text), name)


# ------------------------------------------------------------------------------------------------------ Unpivot
WIDE = pd.DataFrame({"account": ["a", "b", "c"], "region": ["n", "s", "n"],
                     "jan": ["1", None, "5"], "feb": ["2", "4", None], "mar": ["3", "", "7"]})


def test_unpivot_then_pivot_returns_the_original_table():
    long = run("unpivot", WIDE, id_columns=["account", "region"], name_column="month", value_column="amount").df
    assert len(long) == 9 and list(long.columns) == ["account", "region", "month", "amount"]
    assert list(long["month"][:3]) == ["jan", "feb", "mar"] and list(long["account"][:4]) == ["a", "a", "a", "b"]   # row by row
    back = long.pivot(index=["account", "region"], columns="month", values="amount").reset_index()
    back = back[["account", "region", "jan", "feb", "mar"]]
    back.columns.name = None
    pd.testing.assert_frame_equal(back.astype(object).where(back.notna(), None), WIDE.astype(object).where(WIDE.notna(), None))


def test_unpivot_options_and_errors():
    dropped = run("unpivot", WIDE, id_columns=["account"], value_columns=["jan", "feb", "mar"], drop_blank=True)
    assert len(dropped.df) == 6 and dict(dropped.summary)["Blank values left out"] == 3     # None, None and ""
    assert "region" not in dropped.df.columns                                               # only what was asked for
    with pytest.raises(ToolError, match="both as a column to keep"):
        run("unpivot", WIDE, id_columns=["account"], value_columns=["account", "jan"])
    with pytest.raises(ToolError, match="different names"):
        run("unpivot", WIDE, id_columns=["account"], name_column="x", value_column="x")
    with pytest.raises(ToolError, match='already called "account"'):
        run("unpivot", WIDE, id_columns=["account"], name_column="account")
    with pytest.raises(ToolError, match="no column left"):
        run("unpivot", WIDE[["account"]], id_columns=["account"])


# -------------------------------------------------------------------------------------------- Group & Summarise
SALES = pd.DataFrame({
    "dept": ["x", "y", "x", None, "y", " ", "x"],
    "who": ["Bob", "amy", "Cy", "Dee", "Eve", "Fay", None],
    "amount": ["10", "20", "5", "7", None, "3", "2.5"],
})


def test_group_matches_pandas_groupby_including_blank_groups():
    result = run("group-and-summarise", SALES, group_columns=["dept"], value_columns=["amount"],
                 sum=True, mean=True, median=True, min=True, max=True, count=True, distinct=True).df
    numbers = SALES.assign(dept=SALES["dept"].where(SALES["dept"].str.strip() != "", None),
                           amount=pd.to_numeric(SALES["amount"]))
    expected = numbers.groupby("dept", sort=False, dropna=False)["amount"].agg(["sum", "mean", "median", "min", "max", "count", "nunique"])
    assert list(result["dept"].where(result["dept"].notna(), None)) == ["x", "y", None]     # first appearance; blanks together
    assert list(result["rows"]) == [3, 2, 2]
    for ours, theirs in (("sum", "sum"), ("mean", "mean"), ("median", "median"), ("min", "min"), ("max", "max"),
                         ("count", "count"), ("distinct", "nunique")):
        assert [float(v) for v in result[f"amount_{ours}"]] == pytest.approx(list(expected[theirs].astype(float))), ours


def test_group_text_functions_first_last_and_errors():
    result = run("group-and-summarise", SALES, group_columns=["dept"], value_columns=["who"], sum=False,
                 min=True, max=True, first=True, last=True, count=True, row_count=False).df
    assert list(result.columns) == ["dept", "who_min", "who_max", "who_count", "who_first", "who_last"]
    assert result.iloc[0].tolist() == ["x", "Bob", "Cy", 2, "Bob", "Cy"]                    # the blank name is left out
    assert result.iloc[1].tolist()[1:3] == ["amy", "Eve"]                                   # compared without case
    only_rows = run("group-and-summarise", SALES, group_columns=["dept"], sum=False).df
    assert list(only_rows.columns) == ["dept", "rows"]
    with pytest.raises(ToolError, match='"who" has values that are not numbers'):
        run("group-and-summarise", SALES, group_columns=["dept"], value_columns=["who"], sum=True)
    with pytest.raises(ToolError, match="at least one thing"):
        run("group-and-summarise", SALES, group_columns=["dept"], sum=False, row_count=False)
    with pytest.raises(ToolError, match="both to group by and to summarise"):
        run("group-and-summarise", SALES, group_columns=["dept"], value_columns=["dept"], count=True, sum=False)


# ------------------------------------------------------------------------------------------ delimiter detection
@pytest.mark.parametrize("text,expected", [
    ("a,b,c\n1,2,3\n", ","), ("a;b;c\n1,5;2;3\n", ";"), ("a\tb\tc\n1\t2\t3\n", "\t"), ("a|b\n1|2\n", "|"),
    ('"x;y",b\n1,2\n', ","), ('"x,y";b\n1;2\n', ";"), ("single\n1\n", ","), ("\n\na;b\n1;2\n", ";"), ("", ","),
])
def test_the_delimiter_is_read_from_the_first_line_and_a_comma_always_wins(tmp_path, text, expected):
    path = tmp_path / "f.csv"
    path.write_text(text, encoding="utf-8")
    assert datadragon.csv_delimiter(str(path)) == expected


def test_a_semicolon_file_opens_in_an_ordinary_tool(client):
    text = "name;amount;note\nacme;1,5;ok\nbeta;2,5;\n"
    columns = client.post("/get-columns", data={"file": upload(text)}, content_type="multipart/form-data").get_json()
    assert [c["name"] for c in columns["columns"]] == ["name", "amount", "note"]
    done = client.post("/fill-missing", data={"file": upload(text), "options": json.dumps({"columns": ["note"], "value": "-"})},
                       content_type="multipart/form-data").get_json()
    assert done["success"] and done["preview"]["rows"][0] == {"name": "acme", "amount": "1,5", "note": "ok"}
    assert done["preview"]["rows"][1]["note"] == "-"


# ------------------------------------------------------------------------------------------------- Open a File
def workbook_bytes():
    book = openpyxl.Workbook()
    book.active.title = "Cover"
    book.active.append(["Quarterly export"])
    data = book.create_sheet("Data")
    for row in (["ACME LTD - EXPORT", None, None, None], [None] * 4, ["generated 2026-10-01", None, None, None],
                ["id", "amount", None, "amount"], ["a1", 10, None, 1], [None] * 4, ["a2", 20, None, 2]):
        data.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def test_a_second_sheet_with_junk_rows_is_inspected_and_opened(client):
    inspected = client.post("/open-file/inspect", data={"file": upload(workbook_bytes(), "q.xlsx"), "sheet": "Data"},
                            content_type="multipart/form-data").get_json()
    assert inspected["sheets"] == ["Cover", "Data"] and inspected["sheet"] == "Data" and inspected["delimiter"] is None
    assert inspected["rows"][3] == ["id", "amount", None, "amount"] and inspected["suggested_header_row"] == 4
    opened = client.post("/open-file", data={"file": upload(workbook_bytes(), "q.xlsx"), "sheet": "Data", "header_row": "4"},
                         content_type="multipart/form-data").get_json()
    assert opened["success"] and (opened["rows"], opened["columns"]) == (2, 3)             # the empty row and column are gone
    assert opened["preview"]["columns"] == ["id", "amount", "amount_2"]                     # a repeated name is made unique
    assert opened["preview"]["rows"] == [{"id": "a1", "amount": "10", "amount_2": "1"}, {"id": "a2", "amount": "20", "amount_2": "2"}]
    cached = client.get("/get-cached-files").get_json()["files"]
    newest = max(cached, key=lambda f: f["timestamp"])
    assert newest["source_tool"] == "Open a File" and newest["rows"] == 2
    # ... and every tool can now use it: here Sort, with a recipe that starts from this table
    sorted_ = client.post("/sort-and-sample", data={"cache_id": newest["cache_id"],
                                                    "options": json.dumps({"sort1": "amount", "direction1": "descending"})},
                          content_type="multipart/form-data").get_json()
    assert sorted_["preview"]["rows"][0]["id"] == "a2" and sorted_["recipe_steps"] == ["Sort, Rank & Sample"]


def test_open_file_text_options_no_header_and_bad_settings(client):
    text = "report\n\nid|name\n1|a\n2|b\n"
    inspected = client.post("/open-file/inspect", data={"file": upload(text, "r.txt")}, content_type="multipart/form-data").get_json()
    assert inspected["delimiter"] == "comma" and inspected["sheets"] == []                  # the first line has no delimiter at all
    piped = client.post("/open-file/inspect", data={"file": upload(text, "r.txt"), "delimiter": "pipe"},
                        content_type="multipart/form-data").get_json()
    assert piped["rows"][2] == ["id", "name"] and piped["suggested_header_row"] == 3
    opened = client.post("/open-file", data={"file": upload(text, "r.txt"), "delimiter": "pipe", "header_row": "3"},
                         content_type="multipart/form-data").get_json()
    assert opened["preview"]["rows"] == [{"id": "1", "name": "a"}, {"id": "2", "name": "b"}]
    bare = client.post("/open-file", data={"file": upload("1,2\n3,4\n"), "header_row": "0"}, content_type="multipart/form-data").get_json()
    assert bare["preview"]["columns"] == ["Column 1", "Column 2"] and bare["rows"] == 2
    for bad in ({"header_row": "99"}, {"header_row": "x"}):
        response = client.post("/open-file", data={"file": upload("a,b\n1,2\n"), **bad}, content_type="multipart/form-data")
        assert response.status_code == 400 and response.get_json()["error"]
    wrong = client.post("/open-file/inspect", data={"file": upload("x", "x.pdf")}, content_type="multipart/form-data")
    assert wrong.status_code == 400


def test_table_from_grid_keeps_a_named_empty_column():
    grid = pd.DataFrame([["id", "note", None], ["1", None, None], ["2", None, None]])
    table = datadragon.table_from_grid(grid, 1)
    assert list(table.columns) == ["id", "note"] and len(table) == 2                        # "note" is empty but named


# ------------------------------------------------------------------------------------------------ Append Files
JAN = pd.DataFrame({"Vendor": ["a", "b"], "Amount": ["1", "2"]})
FEB = pd.DataFrame({"amount ": ["3"], "Vendor Name": ["c"], "Note": ["late"]})
MAR = pd.DataFrame({"Amount": ["4", "5"], "Vendor": ["d", "e"]})


def test_append_matches_by_name_and_the_row_count_is_the_sum():
    stacked, report = append_tables([("jan.csv", JAN), ("mar.csv", MAR)])
    assert list(stacked.columns) == ["source_file", "Vendor", "Amount"] and len(stacked) == len(JAN) + len(MAR)
    assert list(stacked["Vendor"]) == ["a", "b", "d", "e"] and list(stacked["source_file"]) == ["jan.csv"] * 2 + ["mar.csv"] * 2
    assert report["partial_columns"] == {} and report["rows_per_file"] == [("jan.csv", 2), ("mar.csv", 2)]


def test_append_shows_mismatches_and_maps_them():
    stacked, report = append_tables([("jan.csv", JAN), ("feb.csv", FEB)], source_column=False)
    assert list(stacked.columns) == ["Vendor", "Amount", "amount ", "Vendor Name", "Note"] and len(stacked) == 3
    assert report["partial_columns"]["Vendor"] == ["feb.csv"] and report["partial_columns"]["Note"] == ["jan.csv"]
    mapped, report = append_tables([("jan.csv", JAN), ("feb.csv", FEB)], renames={1: {"Vendor Name": "Vendor"}}, loose=True)
    assert list(mapped.columns) == ["source_file", "Vendor", "Amount", "Note"]              # "amount " joins "Amount" loosely
    assert list(mapped["Vendor"]) == ["a", "b", "c"] and list(mapped["Amount"]) == ["1", "2", "3"]
    assert report["partial_columns"] == {"Note": ["jan.csv"]} and pd.isna(mapped["Note"].iloc[0])


def test_append_errors():
    with pytest.raises(ToolError, match="at least two"):
        append_tables([("jan.csv", JAN)])
    with pytest.raises(ToolError, match='"Nope" is not a column of feb.csv'):
        append_tables([("jan.csv", JAN), ("feb.csv", FEB)], renames={1: {"Nope": "Vendor"}})
    with pytest.raises(ToolError, match='more than one column ends up as "Note"'):
        append_tables([("jan.csv", JAN), ("feb.csv", FEB)], renames={1: {"Vendor Name": "Note"}})


def test_append_route_writes_the_file_and_refuses_one_file(client):
    files = {"file0": upload(JAN.to_csv(index=False), "jan.csv"), "file1": upload(MAR.to_csv(index=False), "mar.csv"),
             "file2": upload(FEB.to_csv(index=False), "feb.csv")}
    done = client.post("/append-files", data={**files, "renames": json.dumps({"2": {"Vendor Name": "Vendor"}}), "loose": "true"},
                       content_type="multipart/form-data").get_json()
    assert done["success"] and done["rows"] == 5 and done["rows_per_file"] == [["jan.csv", 2], ["mar.csv", 2], ["feb.csv", 1]]
    assert done["partial_columns"] == {"Note": ["jan.csv", "mar.csv"]}
    sheet = openpyxl.load_workbook(io.BytesIO(client.get(done["download_url"]).data)).worksheets[0]
    assert [c.value for c in sheet[1]] == ["source_file", "Vendor", "Amount", "Note"] and sheet.max_row == 6
    newest = max(client.get("/get-cached-files").get_json()["files"], key=lambda f: f["timestamp"])
    assert newest["source_tool"] == "Append Files" and newest["recipe_steps"] is None        # several inputs: not replayable
    one = client.post("/append-files", data={"file0": upload("a\n1\n")}, content_type="multipart/form-data")
    assert one.status_code == 400 and "at least two" in one.get_json()["error"]


@pytest.mark.parametrize("path", ["/open-file", "/append-files", "/unpivot", "/group-and-summarise"])
def test_the_new_pages_are_in_the_hub(client, path):
    html = client.get(path).get_data(as_text=True)
    assert f'href="{path}" aria-current="page"' in html and 'id="helpModal"' in html
