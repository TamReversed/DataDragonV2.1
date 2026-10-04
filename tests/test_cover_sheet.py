"""The cover sheet of Excel outputs: first sheet, branded, skipped when the file is used as input, and optional.

The rest of the suite runs with the cover switched off (tests/conftest.py), so these tests switch it on.
"""
import io
import json

import openpyxl
import pandas as pd
import pytest

import datadragon
from datadragon_cover import COVER_SHEET
from helpers import fixture_path, post_job


@pytest.fixture
def cover_on(monkeypatch):
    monkeypatch.setattr(datadragon, "COVER_SHEET_DEFAULT", True)


def csv(df):
    buffer = io.BytesIO()
    df.to_csv(buffer, index=False)
    buffer.seek(0)
    return buffer


TABLE = pd.DataFrame({"id": ["a", "b", "b", "c"], "name": [" Acme ", "Beta", "Beta", None], "amount": ["1", "2", "2", "4"]})


def workbook_of(client, payload):
    return openpyxl.load_workbook(io.BytesIO(client.get(payload["download_url"]).data))


def run_tool(client, url, **form):
    response = client.post(url, data={"file": (csv(TABLE), "t.csv"), **form}, content_type="multipart/form-data")
    assert response.status_code == 200, response.get_json()
    return response.get_json()


def cover_text(sheet):
    return [[cell for cell in row if cell is not None] for row in sheet.iter_rows(values_only=True) if any(c is not None for c in row)]


def test_the_cover_is_the_first_sheet_and_says_what_the_file_is(client, cover_on):
    done = run_tool(client, "/remove-duplicates", options=json.dumps({"columns": ["id"]}))
    book = workbook_of(client, done)
    assert book.sheetnames == [COVER_SHEET, "Kept rows", "Removed rows", "_DataDragon_Log"] and book.active.title == COVER_SHEET
    text = cover_text(book[COVER_SHEET])
    assert ["Remove Duplicates"] in text and ["Rows in", 4] in text and ["Rows out", 3] in text
    assert ["Kept rows", 3, 3, "data"] in text and ["Removed rows", 1, 3] in text
    links = {cell.value: cell.hyperlink.location for row in book[COVER_SHEET].iter_rows() for cell in row if cell.hyperlink}
    assert links == {"Kept rows": "'Kept rows'!A1", "Removed rows": "'Removed rows'!A1", "_DataDragon_Log": "'_DataDragon_Log'!A1"}
    assert len(book[COVER_SHEET]._images) == 2                                    # the mark and the banner
    flat = json.dumps(text)
    assert "Acme" not in flat and "Beta" not in flat                              # no cell values from the data


def test_the_data_is_the_same_with_and_without_the_cover(client, monkeypatch):
    def data_sheets(enabled):
        monkeypatch.setattr(datadragon, "COVER_SHEET_DEFAULT", enabled)
        book = workbook_of(client, run_tool(client, "/fill-missing", options=json.dumps({"columns": ["name"], "value": "?"})))
        return {s.title: [[c.value for c in row] for row in s.iter_rows()] for s in book.worksheets
                if s.title not in (COVER_SHEET, "_DataDragon_Log")}
    assert data_sheets(True) == data_sheets(False)


def test_a_result_with_a_cover_is_read_correctly_by_the_next_tool(client, cover_on):
    run_tool(client, "/text-cleaner", options=json.dumps({"columns": ["name"]}))
    newest = max(client.get("/get-cached-files").get_json()["files"], key=lambda f: f["timestamp"])
    columns = client.post("/get-columns", data={"cache_id": newest["cache_id"]}, content_type="multipart/form-data").get_json()
    assert [c["name"] for c in columns["columns"]] == ["id", "name", "amount"]
    preview = client.post("/preview-data", data={"cache_id": newest["cache_id"]}, content_type="multipart/form-data").get_json()
    assert preview["total_rows"] == 4 and preview["rows"][0]["name"] == "Acme"
    again = client.post("/sort-and-sample", data={"cache_id": newest["cache_id"], "options": json.dumps({"sort1": "amount", "direction1": "descending"})},
                        content_type="multipart/form-data").get_json()
    assert again["preview"]["rows"][0]["id"] == "c" and again["recipe_steps"] == ["Text Cleaner", "Sort, Rank & Sample"]
    path = datadragon.file_cache[newest["cache_id"]]["path"]
    assert datadragon.read_headers(path) == ["id", "name", "amount"] and datadragon.count_data_rows(path) == 4
    assert datadragon.sheet_names_of(path) == ["Sheet1"]


def test_the_cover_can_be_switched_off_by_the_browser_and_by_the_deployment(client, monkeypatch):
    monkeypatch.setattr(datadragon, "COVER_SHEET_DEFAULT", True)
    client.set_cookie("dd_cover", "0")
    assert COVER_SHEET not in workbook_of(client, run_tool(client, "/text-cleaner", options=json.dumps({"columns": ["name"]}))).sheetnames
    client.delete_cookie("dd_cover")
    assert COVER_SHEET in workbook_of(client, run_tool(client, "/text-cleaner", options=json.dumps({"columns": ["name"]}))).sheetnames
    monkeypatch.setattr(datadragon, "COVER_SHEET_DEFAULT", False)
    assert COVER_SHEET not in workbook_of(client, run_tool(client, "/text-cleaner", options=json.dumps({"columns": ["name"]}))).sheetnames


def test_background_jobs_and_the_styled_pivot_get_the_cover_too(client, cover_on):
    golden = fixture_path("golden.xlsx")
    final, _ = post_job(client, "/generate-pivot", files={"file": golden},
                        data={"rows[]": "Region", "values[]": "Amount", "aggfunc": "sum"})
    assert final.get("stage") == "done", final
    book = openpyxl.load_workbook(io.BytesIO(client.get(final["download_url"]).data))
    assert book.sheetnames[0] == COVER_SHEET and ["Pivot Table Generator"] in cover_text(book[COVER_SHEET])
    text = cover_text(book[COVER_SHEET])
    listed = [row[0] for row in text[text.index(["sheet", "rows", "columns", "what it is"]) + 1:] if row[0] in book.sheetnames]
    assert listed == [name for name in book.sheetnames if name != COVER_SHEET]     # every other sheet is listed, in order


def test_a_background_job_respects_the_browser_that_started_it(client, cover_on):
    client.set_cookie("dd_cover", "0")
    final, _ = post_job(client, "/generate-pivot", files={"file": fixture_path("golden.xlsx")},
                        data={"rows[]": "Region", "values[]": "Amount", "aggfunc": "sum"})
    book = openpyxl.load_workbook(io.BytesIO(client.get(final["download_url"]).data))
    assert COVER_SHEET not in book.sheetnames


def test_split_files_and_csv_outputs_never_get_a_cover(client, cover_on):
    final, _ = post_job(client, "/upload", files={"file": fixture_path("golden.xlsx")}, data={"chunk_size": "15"})
    assert final.get("stage") == "done", final
    import zipfile
    archive = zipfile.ZipFile(io.BytesIO(client.get(final["download_url"]).data))
    for name in archive.namelist():
        if name.endswith(".xlsx"):
            assert COVER_SHEET not in openpyxl.load_workbook(io.BytesIO(archive.read(name))).sheetnames, name


def test_the_sidebar_has_the_switch(client):
    html = client.get("/").get_data(as_text=True)
    assert 'id="coverToggle"' in html and "dd_cover=0" in open("static/js/shell.js", encoding="utf-8").read()
