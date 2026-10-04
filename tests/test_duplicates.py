"""Duplicate finder: rows are duplicates when their VALUES match in every selected column."""
import os
import random
import shutil
import time
from collections import OrderedDict
from queue import Queue

import openpyxl
import pytest

import datadragon
from helpers import make_xlsx, output_path
from test_merge import write_rows

_counter = iter(range(10**9))


def run_duplicates(path, id_column, columns, treat_blank_as_value=True):
    q = Queue()
    session = f"dups_test_{next(_counter)}"
    copy = os.path.join(os.path.dirname(path), f"{session}{os.path.splitext(path)[1]}")
    shutil.copy(path, copy)  # the tool deletes its input when it finishes
    datadragon.find_duplicates_async(copy, id_column, columns, q, session, treat_blank_as_value)
    msgs = []
    while not q.empty():
        msgs.append(q.get())
    return msgs[-1]


def groups_of(final):
    sheet = openpyxl.load_workbook(output_path(final))["Duplicates"]
    out = []
    for row in sheet.iter_rows(min_row=2):
        preview, count, keep, all_ids, to_remove = [c.value for c in row]
        out.append((preview, count, keep, all_ids, to_remove))
    return out


def reference(rows, treat_blank_as_value):
    """rows: (id, a, b); None = blank. Groups in order of first appearance, ids in file order."""
    groups = OrderedDict()
    for rid, a, b in rows:
        key = (a, b)
        if not treat_blank_as_value and (a is None or b is None):
            continue
        groups.setdefault(key, []).append(rid)
    return [(key, ids) for key, ids in groups.items() if len(ids) > 1]


@pytest.mark.parametrize("seed", range(40))
@pytest.mark.parametrize("treat_blank", [True, False])
def test_matches_reference(tmp_path, seed, treat_blank):
    rng = random.Random(seed)
    alphabet = ["x", "y", "x|||y", "y|||x", "x||", "||y", "", None, "None", "nan", "0", "00"]
    rows = [(f"r{i}", rng.choice(alphabet), rng.choice(alphabet)) for i in range(rng.randint(0, 14))]
    rows = [(rid, a or None if a == "" else a, b or None if b == "" else b) for rid, a, b in rows]
    path = write_rows(tmp_path / "d.csv", ["id", "a", "b"], rows)
    final = run_duplicates(path, "id", ["a", "b"], treat_blank)
    assert final["stage"] == "done", final
    expected = reference(rows, treat_blank)
    actual = groups_of(final) if expected else []
    assert [(g[3]) for g in actual] == [", ".join(ids) for _, ids in expected]
    assert [g[1] for g in actual] == [len(ids) for _, ids in expected]
    assert [g[2] for g in actual] == [ids[0] for _, ids in expected]
    assert [g[4] for g in actual] == [", ".join(ids[1:]) for _, ids in expected]
    assert final["total_duplicates"] == len(expected)
    assert final["total_ids_to_remove"] == sum(len(ids) - 1 for _, ids in expected)


def test_values_containing_the_old_separator_are_not_duplicates(tmp_path):
    path = make_xlsx(tmp_path / "d.xlsx", ["id", "a", "b"], [[1, "x|||y", "z"], [2, "x", "y|||z"]], text_cols=("a", "b"))
    assert run_duplicates(path, "id", ["a", "b"])["total_duplicates"] == 0


def test_ids_to_remove_sheet_lists_every_repeat_in_file_order(tmp_path):
    path = write_rows(tmp_path / "d.csv", ["id", "k"], [["a", "1"], ["b", "1"], ["c", "2"], ["d", "1"], ["e", "2"]])
    final = run_duplicates(path, "id", ["k"])
    wb = openpyxl.load_workbook(output_path(final))
    assert [r[0].value for r in wb["IDs to Remove"].iter_rows(min_row=2)] == ["b", "d", "e"]
    assert groups_of(final)[0][:5] == ("k=1", 3, "a", "a, b, d", "b, d")


def test_blank_cells_are_shown_empty_not_as_nan(tmp_path):
    path = write_rows(tmp_path / "d.csv", ["id", "a", "b"], [["1", None, "x"], ["2", None, "x"]])
    assert groups_of(run_duplicates(path, "id", ["a", "b"]))[0][0] == "a=, b=x"


def test_blank_ids_are_not_written_as_nan(tmp_path):
    path = write_rows(tmp_path / "d.csv", ["id", "k"], [[None, "1"], ["b", "1"]])
    assert groups_of(run_duplicates(path, "id", ["k"]))[0][3] == ", b"


def test_the_blank_option_changes_the_answer(tmp_path):
    path = write_rows(tmp_path / "d.csv", ["id", "a"], [["1", None], ["2", None], ["3", "x"], ["4", "x"]])
    on = run_duplicates(path, "id", ["a"], treat_blank_as_value=True)
    off = run_duplicates(path, "id", ["a"], treat_blank_as_value=False)
    assert on["total_duplicates"] == 2 and on["treat_blank_as_value"] is True
    assert off["total_duplicates"] == 1 and off["treat_blank_as_value"] is False


def test_excel_numbers_and_text_stay_distinct_by_value(tmp_path):
    path = make_xlsx(tmp_path / "d.xlsx", ["id", "v"], [["a", 7], ["b", "7"], ["c", 7]], text_cols=())
    final = run_duplicates(path, "id", ["v"])
    assert groups_of(final)[0][3] == "a, c"  # number 7 and text '7' are different values


def test_route_option_reaches_the_tool(client, tmp_path):
    from helpers import post_job
    path = write_rows(tmp_path / "d.csv", ["id", "a"], [["1", None], ["2", None]])
    final, _ = post_job(client, "/find-duplicates", {"file": path},
                        {"id_column": "id", "duplicate_columns[]": ["a"], "treat_blank_as_value": "false"})
    assert final["total_duplicates"] == 0 and final["treat_blank_as_value"] is False
    final, _ = post_job(client, "/find-duplicates", {"file": path}, {"id_column": "id", "duplicate_columns[]": ["a"]})
    assert final["total_duplicates"] == 1 and final["treat_blank_as_value"] is True


def test_scales_with_many_groups(tmp_path):
    """The old code rescanned the whole frame once per group (100k rows / 23k groups took 40 s)."""
    n = 30000
    rows = [[f"R{i}", f"n{i % 7000}", f"d{i % 3}"] for i in range(n)]
    path = write_rows(tmp_path / "d.csv", ["id", "a", "b"], rows)
    started = time.time()
    final = run_duplicates(path, "id", ["a", "b"])
    elapsed = time.time() - started
    assert final["total_duplicates"] > 5000
    assert elapsed < 5, f"{n} rows took {elapsed:.1f}s"
