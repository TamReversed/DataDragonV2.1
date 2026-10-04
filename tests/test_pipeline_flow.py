"""Pipeline: stage prerequisites, invalidation when a stage is re-run, hand-off between stages, chained output."""
import os
import time
import zipfile
from queue import Queue

import openpyxl
import pytest

import datadragon
from helpers import drain_progress, fixture_path, make_csv, output_path, post_form, post_job, sse_messages
from test_merge import write_rows


def start(client, path=None):
    datadragon.rate_limit_store.clear()
    path = path or fixture_path("golden.csv")
    with open(path, "rb") as fh:
        return client.post("/pipeline/start", data={"file": (fh, os.path.basename(path))},
                           content_type="multipart/form-data").get_json()["session_id"]


def run_stage(client, sid, stage, **body):
    routes = {"analyze": ("post", "analyze"), "keys": ("post", "keys"), "execute": ("post", "execute")}
    method, path = routes[stage]
    resp = getattr(client, method)(f"/pipeline/{sid}/{path}", json=body or None)
    assert resp.status_code == 200, resp.get_data(as_text=True)
    return drain_progress(client, sid)


def full_flow(client):
    sid = start(client)
    run_stage(client, sid, "analyze")
    client.post(f"/pipeline/{sid}/gaps/triage", json={"Region": "acceptable"})
    run_stage(client, sid, "keys", selected_columns=["ID", "Zip", "Account"])
    client.post(f"/pipeline/{sid}/keys/confirm", json={"selected_key": ["ID"]})
    client.post(f"/pipeline/{sid}/transformations/select", json={"anonymization": {"enabled": True, "columns": ["Name"]}})
    msgs = run_stage(client, sid, "execute")
    assert msgs[-1]["stage"] == "done", msgs[-1]
    return sid, msgs[-1]


# ------------------------------------------------------------------ prerequisites
def test_execute_needs_the_analysis_first(client):
    sid = start(client)
    resp = client.post(f"/pipeline/{sid}/execute")
    assert resp.status_code == 409 and "stage 1" in resp.get_json()["error"]
    assert datadragon.progress_queues.get(sid) is None          # nothing was started


def test_execute_works_with_only_the_analysis_done(client):
    sid = start(client)
    run_stage(client, sid, "analyze")
    msgs = run_stage(client, sid, "execute")                   # no keys, no triage, no selections
    assert msgs[-1]["stage"] == "done", msgs[-1]
    assert os.path.exists(output_path(msgs[-1]))


def test_report_generation_tolerates_every_missing_stage(client):
    sid = start(client)
    run_stage(client, sid, "analyze")
    state = datadragon.pipeline_sessions[sid]
    for stage in (2, 3, 4):
        state.stage_data[stage] = None                          # present-but-None used to crash .get(n, {}).get(...)
    msgs = run_stage(client, sid, "execute")
    assert msgs[-1]["stage"] == "done"


# ------------------------------------------------------------------ invalidation
def stage_status(client, sid):
    return client.get(f"/pipeline/{sid}/state").get_json()


def test_rerunning_the_analysis_discards_later_stages_and_decisions(client):
    sid, _ = full_flow(client)
    state = datadragon.pipeline_sessions[sid]
    assert all(state.stage_data[n] is not None for n in (1, 2, 3, 4, 5))      # a complete run has every stage
    run_stage(client, sid, "analyze")
    assert state.stage_data[1] is not None
    assert [state.stage_data[n] for n in (2, 3, 4, 5)] == [None] * 4
    assert state.user_decisions == {2: {}, 3: {}, 4: {}}
    assert stage_status(client, sid)["state"]["stages_completed"] == [1]


def test_rerunning_key_discovery_discards_the_chosen_key_and_later_stages(client):
    sid, _ = full_flow(client)
    state = datadragon.pipeline_sessions[sid]
    triage_before = dict(state.user_decisions[2])
    run_stage(client, sid, "keys", selected_columns=["ID", "Zip"])
    assert state.stage_data[3] is not None
    assert state.user_decisions[3] == {} and state.user_decisions[4] == {}
    assert state.stage_data[4] is None and state.stage_data[5] is None
    assert state.user_decisions[2] == triage_before             # earlier stages are untouched
    assert state.stage_data[1] is not None


def test_new_transformation_choices_discard_the_old_execution_result(client):
    sid, _ = full_flow(client)
    state = datadragon.pipeline_sessions[sid]
    assert state.stage_data[5] is not None
    client.post(f"/pipeline/{sid}/transformations/select", json={})
    assert state.stage_data[5] is None and state.stage_data[3] is not None and state.stage_data[4] is not None


def test_the_report_after_rerunning_keys_does_not_show_the_stale_key(client):
    from helpers import pdf_text
    import re
    sid, _ = full_flow(client)
    run_stage(client, sid, "keys", selected_columns=["Zip", "Account"])
    client.post(f"/pipeline/{sid}/keys/confirm", json={"selected_key": ["Zip"]})
    msgs = run_stage(client, sid, "execute")
    zf = zipfile.ZipFile(output_path(msgs[-1]))
    text = re.sub(r"\s+", "", pdf_text(zf.read([n for n in zf.namelist() if n.endswith(".pdf")][0])))
    assert "SelectednaturalkeyZip" in text and "SelectednaturalkeyID" not in text      # the heading, then the key in its box


# ------------------------------------------------------------------ stream hand-off
def test_a_finished_stream_does_not_delete_the_next_stages_queue(client):
    client.get("/")
    with client.session_transaction() as sess:
        owner = sess["owner"]
    sid = "handoff_test_session"
    first, second = Queue(), Queue()
    first.put({"stage": "done", "percentage": 100, "message": "stage 1 finished"})
    datadragon.register_job(sid, first, owner=owner)
    response = client.get(f"/progress/{sid}", buffered=False)
    chunk = next(iter(response.response))
    assert b"stage 1 finished" in chunk
    datadragon.register_job(sid, second, owner=owner)           # the next stage starts before the old stream closes
    list(response.response)                                     # the old stream now finishes and cleans up
    assert datadragon.progress_queues.get(sid) is second        # ... and must leave the new queue alone
    second.put({"stage": "done", "percentage": 100, "message": "stage 2 finished"})
    assert "stage 2 finished" in client.get(f"/progress/{sid}").get_data(as_text=True)


# ------------------------------------------------------------------ chaining from the pipeline's result
def test_the_cached_pipeline_result_is_a_readable_xlsx_with_a_matching_name(client):
    sid, final = full_flow(client)
    files = client.get("/get-cached-files").get_json()["files"]
    (entry,) = [f for f in files if f["source_tool"] == "Data Readiness Pipeline"]
    assert entry["filename"].endswith("_data.xlsx")
    datadragon.rate_limit_store.clear()
    used = client.post("/use-cached-file", json={"cache_id": entry["cache_id"]})
    assert used.status_code == 200 and used.get_json()["preview"]["total_rows"] == 40
    body = client.get(f"/download-cached-file/{entry['cache_id']}").data
    assert body[:2] == b"PK"
    assert openpyxl.load_workbook(__import__("io").BytesIO(body)).active.max_row == 41
    assert output_path(final).endswith(".zip")                  # the download itself is still the zip package


def test_the_pipeline_zip_still_contains_report_and_data(client):
    _, final = full_flow(client)
    names = zipfile.ZipFile(output_path(final)).namelist()
    assert any(n.endswith("_report.pdf") for n in names) and any(n.endswith("_data.xlsx") for n in names)


# ------------------------------------------------------------------ calculated columns with a cached file
def make_cached(client):
    resp, payload = post_form(client, "/row-filter", {"file": fixture_path("golden.xlsx")},
                              {"conditions": '[{"column": "Name", "operator": "equals", "value": "Bob"}]'})
    assert resp.status_code == 200
    return client.get("/get-cached-files").get_json()["files"][0]["cache_id"]


def test_calculated_columns_can_use_a_cached_file(client):
    cache_id = make_cached(client)
    resp = client.post("/calculated-columns", data={"cache_id": cache_id, "formula": "[Account] * 2", "new_column_name": "Double"})
    assert resp.status_code == 200, resp.get_data(as_text=True)
    payload = resp.get_json()
    assert payload["new_column"] == "Double" and payload["rows"] == 10
    ws = openpyxl.load_workbook(os.path.join(datadragon.app.config["OUTPUT_FOLDER"], *payload["download_url"].split("/")[-2:])).active
    header = [c.value for c in ws[1]]
    assert header[-1] == "Double" and ws.cell(row=2, column=header.index("Account") + 1).value * 2 == ws.cell(row=2, column=len(header)).value


def test_calculated_columns_rejects_unknown_and_foreign_cache_ids(client):
    cache_id = make_cached(client)
    assert client.post("/calculated-columns", data={"cache_id": "nope", "formula": "1", "new_column_name": "x"}).status_code == 400
    other = datadragon.app.test_client()
    resp = other.post("/calculated-columns", data={"cache_id": cache_id, "formula": "1", "new_column_name": "x"})
    assert resp.status_code == 400 and "not found" in resp.get_json()["error"]


# ------------------------------------------------------------------ transpose
def run_transpose(client, tmp_path, rows):
    path = write_rows(tmp_path / "t.csv", ["a", "b"], [[str(i), "x"] for i in range(rows)])
    final, _ = post_job(client, "/transpose-data", {"file": path})
    return final


def test_transposing_more_rows_than_excel_has_columns_gives_a_clear_message(client, tmp_path):
    final = run_transpose(client, tmp_path, datadragon.EXCEL_MAX_COLUMNS)           # 16,384 rows -> 16,385 columns
    assert final["stage"] == "error" and "16,385 columns" in final["message"] and "16,383 rows" in final["message"]
    assert "download_url" not in final


def test_the_largest_transposable_file_works(client, tmp_path):
    final = run_transpose(client, tmp_path, datadragon.EXCEL_MAX_COLUMNS - 1)       # 16,383 rows -> 16,384 columns
    assert final["stage"] == "done", final
    ws = openpyxl.load_workbook(output_path(final), read_only=True).active
    first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    assert first[0] == "Original Column/Row" and len(first) == datadragon.EXCEL_MAX_COLUMNS


def test_transposed_first_column_is_named_by_rename_not_by_mutation(client, tmp_path):
    final = run_transpose(client, tmp_path, 3)
    ws = openpyxl.load_workbook(output_path(final)).active
    assert [c.value for c in ws[1]][0] == "Original Column/Row"
    assert [[c.value for c in r][0] for r in ws.iter_rows(min_row=2)] == ["a", "b"]


# ------------------------------------------------------------------ T3.7: preview and resumable sessions
def test_state_carries_the_first_twenty_rows_and_the_columns(client):
    sid = start(client)                                   # golden.csv: 40 rows
    payload = client.get(f"/pipeline/{sid}/state").get_json()
    assert len(payload["preview"]) == 20
    assert payload["columns"][:3] == ["ID", "Zip", "Account"]
    assert payload["preview"][0]["ID"] == "PR-00001" and payload["analysis"] is None
    assert payload["state"]["row_count"] == 40


def test_state_includes_the_analysis_once_it_exists(client):
    sid = start(client)
    run_stage(client, sid, "analyze")
    payload = client.get(f"/pipeline/{sid}/state").get_json()
    assert payload["analysis"]["overview"]["shape"]["rows"] == 40


def test_another_browser_cannot_resume_my_session(client):
    sid = start(client)
    other = datadragon.app.test_client()
    assert other.get(f"/pipeline/{sid}/state").status_code == 404


def test_the_page_puts_the_session_in_the_url_and_resumes_from_it():
    html = open("templates/data_readiness_pipeline.html", encoding="utf-8").read()
    assert "history.replaceState(null, '', location.pathname + '?s='" in html
    assert "new URLSearchParams(location.search).get('s')" in html and "resumeFromUrl();" in html


def test_the_right_panel_has_a_view_for_every_stage():
    """Each stage shows its own part of the analysis (the panel used to stay the same from stage 2 on)."""
    html = open("templates/data_readiness_pipeline.html", encoding="utf-8").read()
    assert "const STAGE_VIEWS = {" in html and "showStageView(num);" in html
    for view in ("gaps", "keys", "report"):
        assert f'.analysis-dashboard[data-view="{view}"]' in html, view
    for name in ("card-types", "card-uniqueness", "card-completeness", 'id="previewTitle"'):
        assert name in html, name
