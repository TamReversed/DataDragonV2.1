"""T3.9: accurate previews, filter precedence and validation, Excel-compatible ROUND, friendly 1/0, quoted names."""
import json

import openpyxl
import pandas as pd
import pytest

import datadragon
from datadragon_formula import FormulaError, evaluate_formula
from helpers import make_csv, post_form, post_job, sync_output
from test_known_bugs import sheets


# ------------------------------------------------------------------ preview row counts
def test_csv_count_treats_a_quoted_line_break_as_one_row(tmp_path):
    path = make_csv(tmp_path / "q.csv", 'id,note\n1,"first line\nsecond line"\n2,plain\n\n3,"x"\n')
    assert datadragon.count_data_rows(str(path)) == 3


def test_xlsx_count_uses_the_sheet_dimension_and_ignores_other_sheets(tmp_path):
    path = tmp_path / "w.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["a"])
    for i in range(7):
        ws.append([i])
    wb.create_sheet("other").append(["x", "y"])
    wb.save(path)
    assert datadragon.count_data_rows(str(path)) == 7


def test_preview_route_reports_the_true_total(client, tmp_path):
    path = make_csv(tmp_path / "q.csv", 'id,note\n1,"a\nb"\n2,c\n')
    resp, payload = post_form(client, "/preview-data", {"file": path})
    assert resp.status_code == 200 and payload["total_rows"] == 2 and len(payload["rows"]) == 2


def test_an_unreadable_file_gives_a_friendly_error_not_the_parser_text(client, tmp_path):
    path = tmp_path / "broken.xlsx"
    path.write_bytes(b"this is not a workbook")
    resp, payload = post_form(client, "/preview-data", {"file": str(path)})
    assert resp.status_code == 400 and "could not be read" in payload["error"] and "zip" not in payload["error"].lower()


# ------------------------------------------------------------------ row filter
def run_filter(client, tmp_path, text, conditions):
    path = make_csv(tmp_path / "f.csv", text)
    resp, payload = post_form(client, "/row-filter", {"file": path},
                              {"conditions": json.dumps(conditions), "preview_only": "true"})
    return resp, payload


def test_and_binds_tighter_than_or(client, tmp_path):
    # a=1 OR (b=1 AND c=1): rows 1 and 4 match. Left to right it would be (a=1 OR b=1) AND c=1: only row 4.
    text = "a,b,c\n1,0,0\n0,1,0\n0,0,1\n0,1,1\n"
    resp, payload = run_filter(client, tmp_path, text, [
        {"column": "a", "operator": "equals", "value": "1"},
        {"column": "b", "operator": "equals", "value": "1", "logic": "OR"},
        {"column": "c", "operator": "equals", "value": "1", "logic": "AND"}])
    assert payload["matching_rows"] == 2


def test_the_first_condition_logic_is_ignored_and_or_chains_work(client, tmp_path):
    text = "a\n1\n2\n3\n"
    resp, payload = run_filter(client, tmp_path, text, [
        {"column": "a", "operator": "equals", "value": "1", "logic": "AND"},
        {"column": "a", "operator": "equals", "value": "3", "logic": "OR"}])
    assert payload["matching_rows"] == 2


def test_a_number_condition_with_text_is_a_400(client, tmp_path):
    resp, payload = run_filter(client, tmp_path, "n\n1\n2\n", [{"column": "n", "operator": "greater_than", "value": "abc"}])
    assert resp.status_code == 400 and '"abc" is not a number' in payload["error"]


def test_numeric_comparison_still_works_and_skips_text_cells(client, tmp_path):
    resp, payload = run_filter(client, tmp_path, "n\n5\nabc\n2\n\n", [{"column": "n", "operator": "greater_than", "value": "3"}])
    assert payload["matching_rows"] == 1


def test_blank_cells_do_not_match_contains(client, tmp_path):
    # the old str(NaN) == 'nan' made every blank cell match "an"
    resp, payload = run_filter(client, tmp_path, "t\nbanana\n\nplum\n", [{"column": "t", "operator": "contains", "value": "an"}])
    assert payload["matching_rows"] == 1


# ------------------------------------------------------------------ formulas
@pytest.mark.parametrize("formula,expected", [
    ("ROUND([v])", [3.0, 4.0, -3.0, 0.0]),
    ("ROUND([v], 0)", [3.0, 4.0, -3.0, 0.0]),
])
def test_round_is_half_up_like_excel(formula, expected):
    df = pd.DataFrame({"v": [2.5, 3.5, -2.5, 0.4]})
    assert evaluate_formula(df, formula).tolist() == expected


def test_round_with_decimals_uses_the_written_decimal_value():
    df = pd.DataFrame({"v": [2.675, 1.005, 1.25]})
    assert evaluate_formula(df, "ROUND([v], 2)").tolist() == [2.68, 1.01, 1.25]
    assert evaluate_formula(df, "ROUND([v], 1)").tolist() == [2.7, 1.0, 1.3]


def test_round_to_tens_and_integers_and_blanks():
    df = pd.DataFrame({"v": [15.0, 25.0, None], "i": [15, 25, 35]})
    assert evaluate_formula(df, "ROUND([v], -1)").dropna().tolist() == [20.0, 30.0]
    assert evaluate_formula(df, "ROUND([i])").tolist() == [15, 25, 35]


def test_division_by_zero_is_a_friendly_formula_error():
    with pytest.raises(FormulaError, match="Division by zero"):
        evaluate_formula(pd.DataFrame({"a": [1]}), "1/0")


def test_a_column_name_with_a_quote_works_next_to_a_quoted_string():
    df = pd.DataFrame({'We"ird': [1, 2], "It's": ["a", "b"]})
    assert evaluate_formula(df, 'CONCAT([We"ird], "-", [It\'s], \'!\')').tolist() == ["1-a!", "2-b!"]


# ------------------------------------------------------------------ column comparison headers
def test_compare_columns_reads_headers_of_cp1252_csv_and_quoted_names(client, tmp_path):
    a = tmp_path / "a.csv"
    a.write_bytes('Stra\xdfe,"Name, full",id\n1,2,3\n'.encode("cp1252"))
    b = make_csv(tmp_path / "b.csv", "id,Extra\n1,2\n")
    final, _ = post_job(client, "/compare-columns", {"file1": str(a), "file2": b})
    assert final["stage"] == "done", final
    names = {row["Column Headers"] for row in final["comparison_data"]}
    assert names == {"Straße", "Name, full", "id", "Extra"}
    assert final["common_columns"] == 1


def test_read_headers_reads_row_one_of_the_first_sheet(tmp_path):
    path = tmp_path / "h.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["a", None, " b "])
    ws.append([1, 2, 3])
    wb.create_sheet("second").append(["zzz"])
    wb.save(path)
    assert datadragon.read_headers(str(path)) == ["a", "b"]


def test_read_headers_of_an_empty_csv_is_a_clear_error(tmp_path):
    path = tmp_path / "e.csv"
    path.write_text("")
    with pytest.raises(datadragon.UserError, match="empty"):
        datadragon.read_headers(str(path), "e.csv")


def test_top_values_break_ties_by_first_appearance():
    """Equally frequent values are listed in file order, the same on every machine (the default sort is not stable)."""
    import pandas as pd
    import datadragon
    values = [f"v{i:02d}" for i in range(30)]                       # every value once: all ties
    df = pd.DataFrame({"name": values + ["v17", "v17", "v03"]})     # v17 three times, v03 twice
    result = datadragon.analyze_dataframe(df)
    assert list(result["columns"]["name"]["top_values"]) == ["v17", "v03", "v00", "v01", "v02", "v04", "v05", "v06", "v07", "v08"]
