"""Regression tests for the findings of the Phase 3 independent review."""
import json
import logging
import math
import os
import shutil
import subprocess
import zipfile
from queue import Queue

import openpyxl
import pandas as pd
import pytest

import datadragon
from datadragon_formula import FormulaError, evaluate_formula
from helpers import make_csv, output_path, post_form, post_job, sse_messages

NODE = shutil.which("node")


# 1: NaN on the wire
def test_the_progress_stream_never_sends_a_bare_nan(client, tmp_path):
    path = make_csv(tmp_path / "t.csv", "a,b,c\n1,x,2024-01-01\n2,y,2024-01-02\n3,z,2024-01-03\n")   # < 4 values: kurtosis is NaN
    datadragon.rate_limit_store.clear()
    with open(path, "rb") as fh:
        sid = client.post("/analyze", data={"file": (fh, "t.csv")}, content_type="multipart/form-data").get_json()["session_id"]
    raw = client.get(f"/progress/{sid}").get_data(as_text=True)
    assert "NaN" not in raw and "Infinity" not in raw
    final = sse_messages(raw)[-1]
    assert final["stage"] == "done" and final["analysis"]["columns"]["a"]["statistics"]["kurtosis"] is None


# 2 + parse errors: the browser helper (needs node)
def run_node(script):
    return subprocess.run([NODE, "-e", script], capture_output=True, text=True, check=True, timeout=30).stdout


HARNESS = """
const calls = []; let sources = [];
global.document = { addEventListener() {}, documentElement: {} };
global.MutationObserver = class { observe() {} };
global.requestAnimationFrame = f => f();
global.EventSource = class { constructor(url) { this.url = url; calls.push('es ' + url); sources.push(this); } close() { this.closed = true; } };
global.fetch = (url, init) => { calls.push('fetch ' + url + ' ' + ((init || {}).method || 'GET')); return global.__respond(url, init); };
""" + open("static/js/common.js", encoding="utf-8").read().replace("const DD = (() => {", "global.DD = (() => {", 1) + "\n"


@pytest.mark.skipif(NODE is None, reason="node is not installed")
def test_cancel_pressed_while_the_post_is_in_flight_still_cancels_the_job():
    out = run_node(HARNESS + """
global.__respond = (url) => new Promise(resolve => setTimeout(() => resolve({ ok: true, status: 200,
    json: async () => ({ success: true, session_id: 'job1' }) }), 200));
const errors = [];
const job = DD.startJob('/slow', null, { onError: m => errors.push(m), onDone() {}, onProgress() {} });
setTimeout(() => job.cancel(), 50);
setTimeout(() => { console.log(JSON.stringify({ calls, errors })); }, 600);
""")
    result = json.loads(out)
    assert "fetch /jobs/job1/cancel POST" in result["calls"]
    assert not any(c.startswith("es ") for c in result["calls"]) and result["errors"] == []


@pytest.mark.skipif(NODE is None, reason="node is not installed")
def test_an_unreadable_progress_message_is_shown_as_an_error_not_swallowed():
    out = run_node(HARNESS + """
global.__respond = async () => ({ ok: true, status: 200, json: async () => ({ success: true, session_id: 'j2' }) });
const errors = [];
DD.startJob('/x', null, { onError: m => errors.push(m), onDone() {}, onProgress() {} });
setTimeout(() => { sources[0].onmessage({ data: '{"stage": "done", "k": NaN}' }); console.log(JSON.stringify({ errors })); }, 300);
""")
    assert "could not read" in json.loads(out)["errors"][0]


# 3: too large
def test_an_over_limit_upload_says_so(client, tmp_path, monkeypatch):
    monkeypatch.setitem(datadragon.app.config, "MAX_CONTENT_LENGTH", 500)
    path = make_csv(tmp_path / "big.csv", "a\n" + "x" * 2000 + "\n")
    resp, payload = post_form(client, "/preview-data", {"file": path})
    assert resp.status_code == 413 and "too large" in payload["error"]


# 4: no literals in the log
def test_row_filter_and_validation_logs_keep_structure_not_values(client, tmp_path):
    path = make_csv(tmp_path / "f.csv", "name\nAlice\nBob\n")
    resp, payload = post_form(client, "/row-filter", {"file": path}, {"conditions": json.dumps(
        [{"column": "name", "operator": "equals", "value": "Alice-secret-value"}])})
    wb = openpyxl.load_workbook(output_path(payload))
    text = " ".join(str(c.value) for row in wb[datadragon.LOG_SHEET].iter_rows() for c in row)
    assert "Alice-secret-value" not in text and "equals" in text
    final, _ = post_job(client, "/validate-data", {"file": path}, {"validation_rules": json.dumps(
        [{"column": "name", "type": "list", "value": ["Alice-secret-value"]}])})
    wb = openpyxl.load_workbook(output_path(final))
    assert "Alice-secret-value" not in " ".join(str(c.value) for row in wb[datadragon.LOG_SHEET].iter_rows() for c in row)


# 5: chaining does not grow names forever
def test_chained_results_keep_working_past_six_steps(client, tmp_path):
    path = make_csv(tmp_path / "start.csv", "a,b\n1,2\n")
    resp, payload = post_form(client, "/find-replace", {"file": path}, {"find_text": "1", "replace_text": "1"})
    assert resp.status_code == 200
    longest = 0
    for _ in range(8):
        cache_id = client.get("/get-cached-files").get_json()["files"][0]["cache_id"]
        resp, payload = post_form(client, "/find-replace", None, {"cache_id": cache_id, "find_text": "1", "replace_text": "1"})
        assert resp.status_code == 200, payload
        longest = max(longest, len(payload["filename"]))
    assert longest < 130


# 6: a cancel does not poison the next stage of the same session
def test_a_new_job_under_the_same_id_starts_clean():
    q = Queue()
    datadragon.cancelled_jobs.add("sess_reuse")
    datadragon.register_job("sess_reuse", q, owner="o")
    q.put({"stage": "working"})                                  # would raise JobCancelled if the old flag stayed
    assert "sess_reuse" not in datadragon.cancelled_jobs


# 7: xlsx row counts
def test_xlsx_count_ignores_trailing_formatted_rows_and_a_stale_dimension(tmp_path):
    path = tmp_path / "t.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["a"])
    ws.append([1])
    ws.append([2])
    for r in range(4, 300):
        ws.cell(row=r, column=1).font = openpyxl.styles.Font(bold=True)          # styled but empty
    wb.save(path)
    assert datadragon.count_data_rows(str(path)) == 2 == len(datadragon.read_data_file(str(path)))


def test_xlsx_with_a_wrong_dimension_is_still_counted(tmp_path):
    path = tmp_path / "d.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in (["a"], [1], [2]):
        ws.append(row)
    wb.save(path)
    with zipfile.ZipFile(path) as src:
        files = {n: src.read(n) for n in src.namelist()}
    files["xl/worksheets/sheet1.xml"] = files["xl/worksheets/sheet1.xml"].replace(b'<dimension ref="A1:A3"/>', b'<dimension ref="A1"/>')
    fixed = tmp_path / "e.xlsx"
    with zipfile.ZipFile(fixed, "w") as out:
        for name, data in files.items():
            out.writestr(name, data)
    assert datadragon.count_data_rows(str(fixed)) == 2


# 10: formula edges
def test_round_edges():
    df = pd.DataFrame({"v": [-2.5], "big": [10 ** 18 + 1], "q": [3]})
    with pytest.raises(FormulaError):
        evaluate_formula(df, "ROUND([v], 400)")
    assert evaluate_formula(df, "ROUND([big], -3)").tolist() == [10 ** 18]
    assert str(evaluate_formula(df, "ROUND(-0.4, 0) + [q] * 0").tolist()[0]) in ("0.0", "0")
    assert math.copysign(1, evaluate_formula(pd.DataFrame({"v": [-0.4]}), "ROUND([v], 0)").tolist()[0]) == 1.0
    with pytest.raises(FormulaError, match="Division by zero"):
        evaluate_formula(df, "[q] / 0")


# 12: JSON logging
def test_log_lines_are_real_json_even_with_quotes_and_line_breaks(caplog):
    from datadragon_logging import JsonFormatter
    record = logging.LogRecord("datadragon", logging.INFO, __file__, 1, 'a "quoted"\nfake {"level": "ERROR"} line', None, None)
    record.job = "j1"
    line = JsonFormatter().format(record)
    assert "\n" not in line and json.loads(line)["message"] == 'a "quoted"\nfake {"level": "ERROR"} line'


# 9: focus ring wins
def test_the_focus_ring_cannot_be_switched_off_by_component_rules():
    css = open("static/css/a11y.css", encoding="utf-8").read()
    assert "outline: 2px solid var(--dd-focus, #17605C) !important" in css
