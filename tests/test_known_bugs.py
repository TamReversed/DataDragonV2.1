"""One strict-xfail test per audited defect, asserting the CORRECT behaviour.

Each fix commit removes the matching marker (strict=True turns an unremoved marker into a failure).
Finding IDs refer to REVIEW_PLAN.md section 4.
"""
import io
import json
import os
import time
import zipfile

import openpyxl
import pandas as pd
import pytest

import datadragon
from helpers import (drain_progress, fixture_path, make_csv, make_xlsx, output_path, post_form,
                     post_job, sync_output)

GOLDEN_XLSX = fixture_path("golden.xlsx")


def bug(finding):
    return pytest.mark.xfail(strict=True, reason=finding)


def sheets(path):
    wb = openpyxl.load_workbook(path)
    return {ws.title: [[c.value for c in row] for row in ws.iter_rows()] for ws in wb.worksheets}


def types(path):
    wb = openpyxl.load_workbook(path)
    return {ws.title: [[c.data_type for c in row] for row in ws.iter_rows()] for ws in wb.worksheets}


def numbers(rows):
    return [v for r in rows for v in r if isinstance(v, (int, float)) and not isinstance(v, bool)]


# ---------------------------------------------------------------- A: data correctness

def test_read_preserves_leading_zeros(client):
    final, _ = post_job(client, "/upload", {"file": GOLDEN_XLSX}, {"chunk_size": "100", "base_filename": "z"})
    assert final["stage"] == "done", final
    zf = zipfile.ZipFile(output_path(final))
    ws = openpyxl.load_workbook(io.BytesIO(zf.read(zf.namelist()[0]))).active
    assert ws["B2"].value == "01000" and ws["B2"].data_type == "s"


def _merge(client, join="inner", dup="keep_all"):
    final, _ = post_job(client, "/merge-data",
                        {"left_file": fixture_path("merge_left.csv"), "right_file": fixture_path("merge_right.csv")},
                        {"left_key": "id", "right_key": "id", "join_type": join, "duplicate_handling": dup})
    assert final["stage"] == "done", final
    return final


def test_merge_null_keys_do_not_match(client):
    final = _merge(client)
    merged = sheets(output_path(final))["Merged Data"]
    assert len(merged) - 1 == 3  # 1<->1 and 2<->2 (x2); blank keys must not join


def test_merge_stats_non_negative(client):
    summary = _merge(client)["summary"]
    assert summary["unmatched_left"] >= 0 and summary["unmatched_right"] >= 0


def test_compare_int_float_composite_keys(client, tmp_path):
    a = make_csv(tmp_path / "a.csv", "k1,k2,v\n1,x,a\n2,y,b\n")
    b = make_csv(tmp_path / "b.csv", "k1,k2,v\n1,x,a\n2,y,b\n,z,c\n")  # blank makes k1 float in file 2
    final, _ = post_job(client, "/compare-data", {"file1": a, "file2": b}, {"key_columns[]": ["k1", "k2"]})
    s = final["summary"]
    assert (s["common"], s["added"], s["removed"]) == (2, 1, 0)


def test_compare_null_key_not_added_and_removed(client):
    final, _ = post_job(client, "/compare-data",
                        {"file1": fixture_path("compare_null_a.csv"), "file2": fixture_path("compare_null_b.csv")},
                        {"key_columns[]": ["k"]})
    s = final["summary"]
    assert s["added"] == 0 and s["removed"] == 0


def test_compare_duplicate_keys_reported(client):
    final, _ = post_job(client, "/compare-data",
                        {"file1": fixture_path("compare_dup_a.csv"), "file2": fixture_path("compare_dup_b.csv")},
                        {"key_columns[]": ["k1"]})
    assert final["summary"]["changed"] >= 1


def test_duplicates_no_separator_collision(client, tmp_path):
    path = make_xlsx(tmp_path / "d.xlsx", ["id", "a", "b"], [[1, "x|||y", "z"], [2, "x", "y|||z"]],
                     text_cols=("a", "b"))
    final, _ = post_job(client, "/find-duplicates", {"file": path},
                        {"id_column": "id", "duplicate_columns[]": ["a", "b"]})
    assert final["stage"] == "done", final
    assert final["total_duplicates"] == 0


def _pivot(client, path, data):
    final, _ = post_job(client, "/generate-pivot", {"file": path}, data)
    return final


def test_pivot_keeps_blank_keys(client, tmp_path):
    path = make_csv(tmp_path / "p.csv", "Region,amt\nN,10\nS,5\n,100\n")
    final = _pivot(client, path, {"rows[]": ["Region"], "values[]": ["amt"], "aggfunc": "sum"})
    assert final["stage"] == "done", final
    assert 115 in numbers(sheets(output_path(final))["Pivot Table"])


def test_pivot_mean_empty_is_blank(client, tmp_path):
    path = make_csv(tmp_path / "p.csv", "Region,Flag,amt\nN,a,10\nS,b,20\n")
    final = _pivot(client, path, {"rows[]": ["Region"], "columns[]": ["Flag"], "values[]": ["amt"],
                                  "aggfunc": "mean"})
    assert final["stage"] == "done", final
    table = sheets(output_path(final))["Pivot Table"]
    assert 0 not in numbers([r[1:] for r in table[1:]])  # empty combinations must stay blank, not 0


def test_pivot_filters_work(client, tmp_path):
    path = make_csv(tmp_path / "p.csv", "Region,amt\nN,10\nS,20\n")
    final = _pivot(client, path, {"rows[]": ["Region"], "values[]": ["amt"], "aggfunc": "sum",
                                  "filter_columns[]": ["Region"], "filter_values[]": ["N"]})
    assert final["stage"] == "done", final
    nums = numbers(sheets(output_path(final))["Pivot Table"])
    assert 10 in nums and 20 not in nums


def test_scrub_relationship_never_leaks(client, tmp_path):
    path = make_xlsx(tmp_path / "s.xlsx", ["Name", "Email"],
                     [["Alice", "a@x.com"], ["Bob", None], ["Carol", "c@x.com"]], text_cols=("Name", "Email"))
    final, _ = post_job(client, "/scrub-data", {"file": path},
                        {"columns[]": ["Name", "Email"], "relationship_preserve": "true", "export_mapping": "false"})
    assert final["stage"] == "done", final
    values = {v for rows in sheets(output_path(final)).values() for r in rows for v in r}
    assert not values & {"Alice", "Bob", "Carol", "a@x.com", "c@x.com"}


def test_normalize_euro_decimal(client, tmp_path):
    path = make_csv(tmp_path / "n.csv", "v\n\"2,5\"\n\"1.234,56\"\n")
    final, _ = post_job(client, "/normalize-columns", {"file": path},
                        {"column_types": json.dumps({"v": "float"}), "decimal_separator": ","})
    assert final["stage"] == "done", final
    rows = sheets(output_path(final))["Normalized Data"]
    assert [r[0] for r in rows[1:]] == [2.5, 1234.56]


def test_normalize_never_overwrites_with_nat(client, tmp_path):
    # With an explicit order, a value that is not a date keeps its original text (it used to become a blank NaT
    # cell, and a mixed ISO/slash column lost its ISO rows).
    path = make_csv(tmp_path / "n.csv", "d\n01/02/2024\n2024-03-05\nnot a date\n")
    final, _ = post_job(client, "/normalize-columns", {"file": path},
                        {"column_types": json.dumps({"d": "date"}), "date_order": "MDY"})
    assert final["stage"] == "done", final
    values = [r[0] for r in sheets(output_path(final))["Normalized Data"][1:]]
    import datetime as dt
    assert values == [dt.datetime(2024, 1, 2), dt.datetime(2024, 3, 5), "not a date"]
    assert final["summary"]["total_errors"] == 1 and "not a date" in final["warning"]


def test_validation_range_string_bounds(client, tmp_path):
    path = make_csv(tmp_path / "v.csv", "qty\n5\n2\n")
    rules = [{"column": "qty", "type": "range", "value": {"min": "1", "max": "3"}}]
    final, _ = post_job(client, "/validate-data", {"file": path}, {"validation_rules": json.dumps(rules)})
    assert final["stage"] == "done", final
    assert final["summary"]["invalid_rows"] == 1  # 5 is above the maximum of 3


def test_validation_pattern_fullmatch(client, tmp_path):
    path = make_csv(tmp_path / "v.csv", "zip\n123456789\n12345\n")
    rules = [{"column": "zip", "type": "pattern", "value": r"\d{5}"}]
    final, _ = post_job(client, "/validate-data", {"file": path}, {"validation_rules": json.dumps(rules)})
    assert final["stage"] == "done", final
    assert final["summary"]["invalid_rows"] == 1  # 9 digits is not a 5-digit zip


def test_semantic_id_not_postal(dd):
    df = pd.DataFrame({"ID": [f"PR-{i:05d}" for i in range(50)], "n": range(50)})
    cols = dd.analyze_dataframe(df)["columns"]
    assert cols["ID"]["detected_type"] != "Postal Code"
    assert all(c["type_confidence"] <= 100 for c in cols.values())


def test_analyze_bool_column(dd):
    dd.analyze_dataframe(pd.DataFrame({"a": [True, False, True], "b": [1, 2, 3]}))


def test_findreplace_wholecell_alternation(client, tmp_path):
    path = make_csv(tmp_path / "f.csv", "v\nA\nAxx\nB\n")
    resp, payload = post_form(client, "/find-replace", {"file": path},
                              {"find_text": "A|B", "replace_text": "X", "column": "v", "use_regex": "true",
                               "match_whole_cell": "true", "case_sensitive": "true"})
    assert resp.status_code == 200, payload
    assert [r[0] for r in sheets(sync_output(payload))["Sheet1"][1:]] == ["X", "Axx", "X"]


def test_blank_not_written_as_nan(client, tmp_path):
    path = make_csv(tmp_path / "f.csv", "k,v\n1,foo\n2,\n3,foo\n")  # row 2 has a truly empty cell
    resp, payload = post_form(client, "/find-replace", {"file": path},
                              {"find_text": "foo", "replace_text": "bar", "column": "v"})
    assert resp.status_code == 200, payload
    assert [r[1] for r in sheets(sync_output(payload))["Sheet1"][1:]] == ["bar", None, "bar"]


# ---------------------------------------------------------------- G: security

def test_calc_formula_cannot_reach_os(client):
    resp, _ = post_form(client, "/calculated-columns", {"file": GOLDEN_XLSX},
                        {"formula": "pd.io.common.os.getcwd()", "new_column_name": "x", "preview_only": "true"})
    assert os.getcwd() not in resp.get_data(as_text=True)
    assert resp.status_code in (400, 422)


def _run_cached(client, tag):
    resp, payload = post_form(client, "/row-filter", {"file": GOLDEN_XLSX},
                              {"conditions": json.dumps([{"column": "Name", "operator": "equals", "value": "Bob"}])})
    assert resp.status_code == 200, payload
    return payload


def test_cache_isolated_between_clients(client):
    _run_cached(client, "A")
    cache_id = client.get("/get-cached-files").get_json()["files"][0]["cache_id"]
    other = datadragon.app.test_client()
    assert other.get("/get-cached-files").get_json()["files"] == []
    assert other.get(f"/download-cached-file/{cache_id}").status_code == 404


def test_download_requires_owner(client):
    payload = _run_cached(client, "A")
    other = datadragon.app.test_client()
    assert other.get(payload["download_url"]).status_code in (403, 404)


def test_export_formula_injection_neutralised(client):
    resp, payload = post_form(client, "/row-filter", {"file": GOLDEN_XLSX},
                              {"conditions": json.dumps([{"column": "Notes", "operator": "is_not_empty"}])})
    assert resp.status_code == 200, payload
    rows = sheets(sync_output(payload))["Sheet1"]
    kinds = types(sync_output(payload))["Sheet1"]
    notes_col = rows[0].index("Notes")
    evil = [i for i, r in enumerate(rows) if isinstance(r[notes_col], str) and r[notes_col].startswith("=")]
    assert evil and all(kinds[i][notes_col] == "s" for i in evil)  # must stay text, never a live formula


def test_regex_redos_bounded(client, tmp_path):
    # (a|aa)+$ on a long non-matching cell backtracks catastrophically; it must be stopped, not run for minutes
    path = make_csv(tmp_path / "r.csv", "v\n" + "a" * 40 + "!\n")
    started = time.time()
    resp, payload = post_form(client, "/find-replace", {"file": path},
                              {"find_text": "(a|aa)+$", "replace_text": "x", "column": "v", "use_regex": "true"})
    assert time.time() - started < 3
    assert resp.status_code == 422 and "too complex" in payload["error"]


# ---------------------------------------------------------------- H: reliability

def test_concurrent_split_isolated(client, tmp_path, monkeypatch):
    import datetime as real

    class Frozen(real.datetime):
        @classmethod
        def now(cls, tz=None):
            return real.datetime(2026, 1, 1, 12, 0, 0)

    monkeypatch.setattr(datadragon, "datetime", Frozen)
    a = make_xlsx(tmp_path / "a.xlsx", ["owner", "i"], [["ALICE", i] for i in range(30)])
    b = make_xlsx(tmp_path / "b.xlsx", ["owner", "i"], [["BOB", i] for i in range(50)])
    datadragon.rate_limit_store.clear()
    sids = []
    for path in (a, b):
        with open(path, "rb") as fh:
            resp = client.post("/upload", data={"file": (fh, os.path.basename(path)), "chunk_size": "10"},
                               content_type="multipart/form-data")
        sids.append(resp.get_json()["session_id"])
    finals = [next(m for m in drain_progress(client, sid) if m.get("stage") in ("done", "error")) for sid in sids]
    assert [f["stage"] for f in finals] == ["done", "done"], finals
    # Each job has its own directory, so even identical zip names cannot collide: every download
    # must contain exactly its own rows.
    assert finals[0]["download_url"] != finals[1]["download_url"]
    for final, owner, rows in zip(finals, ("ALICE", "BOB"), (30, 50)):
        zf = zipfile.ZipFile(output_path(final))
        frames = [pd.read_excel(io.BytesIO(zf.read(n))) for n in zf.namelist() if n.endswith('.xlsx')]
        combined = pd.concat(frames)
        assert len(combined) == rows and set(combined["owner"]) == {owner}


def test_calc_with_cache_id(client):
    _run_cached(client, "A")
    cache_id = client.get("/get-cached-files").get_json()["files"][0]["cache_id"]
    resp = client.post("/calculated-columns", data={"cache_id": cache_id, "formula": "[Account] * 2",
                                                    "new_column_name": "X"})
    assert resp.status_code == 200, resp.get_data(as_text=True)


def _start_pipeline(client, path):
    datadragon.rate_limit_store.clear()
    with open(path, "rb") as fh:
        resp = client.post("/pipeline/start", data={"file": (fh, os.path.basename(path))},
                           content_type="multipart/form-data")
    return resp.get_json()["session_id"]


def test_pipeline_execute_without_analyze(client):
    sid = _start_pipeline(client, fixture_path("compare_a.csv"))
    resp = client.post(f"/pipeline/{sid}/execute")
    assert resp.status_code == 409  # a prerequisite stage has not run


def test_pipeline_pdf_escapes_column_names(client):
    sid = _start_pipeline(client, fixture_path("pdf_inject.csv"))
    client.post(f"/pipeline/{sid}/analyze")
    drain_progress(client, sid)
    client.post(f"/pipeline/{sid}/keys", json={"selected_columns": ["id", "a<b"]})
    drain_progress(client, sid)
    client.post(f"/pipeline/{sid}/keys/confirm", json={"selected_key": ["a<b"]})
    client.post(f"/pipeline/{sid}/transformations/select", json={})
    client.post(f"/pipeline/{sid}/execute")
    final = next(m for m in drain_progress(client, sid) if m.get("stage") in ("done", "error"))
    assert final["stage"] == "done", final


def test_pipeline_key_uniqueness_on_full_data(client, tmp_path):
    path = make_csv(tmp_path / "k.csv", "id,v\n007,a\n007,a\n008,b\n")
    sid = _start_pipeline(client, path)
    client.post(f"/pipeline/{sid}/analyze")
    drain_progress(client, sid)
    client.post(f"/pipeline/{sid}/keys", json={"selected_columns": ["id", "v"]})
    drain_progress(client, sid)
    found = datadragon.pipeline_sessions[sid].stage_data[3]["minimal_combinations"]
    assert ["id"] not in found and ["v"] not in found  # duplicate rows exist, so neither is unique


def test_xls_upload_reads(client, tmp_path):
    xlwt = pytest.importorskip("xlwt", reason="xlwt is not a dependency; cannot build an .xls fixture")
    wb = xlwt.Workbook()
    ws = wb.add_sheet("S")
    ws.write(0, 0, "a")
    ws.write(1, 0, 1)
    path = tmp_path / "t.xls"
    wb.save(str(path))
    final, _ = post_job(client, "/upload", {"file": str(path)}, {"chunk_size": "10"})
    assert final["stage"] == "done"
