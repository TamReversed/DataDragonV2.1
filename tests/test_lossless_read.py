"""read_data_file keeps cell values exactly; inference is opt-in; multi-sheet workbooks are flagged."""
import datetime as dt
import json

import numpy as np
import openpyxl
import pandas as pd
import pytest

from helpers import fixture_path, make_csv, make_xlsx, output_path, post_form, post_job, sync_output

GOLDEN_XLSX = fixture_path("golden.xlsx")
GOLDEN_CSV = fixture_path("golden.csv")


# ------------------------------------------------------------------ reader

def test_xlsx_text_cells_keep_leading_zeros_and_exact_digits(dd):
    df = dd.read_data_file(GOLDEN_XLSX)
    assert df["Zip"].iloc[0] == "01000" and df["Zip"].iloc[2] == "01014"
    assert df["BigId"].iloc[0] == "9007199254740993"  # beyond float precision


def test_xlsx_values_keep_their_python_types(dd):
    df = dd.read_data_file(GOLDEN_XLSX)
    assert isinstance(df["Account"].iloc[0], int) and not isinstance(df["Account"].iloc[0], bool)
    assert pd.isna(df["Account"].iloc[4])                       # blank stays blank, ints are not turned into floats
    assert df["Amount"].iloc[0] == -300 and isinstance(df["Amount"].iloc[1], float)
    assert df["Date"].iloc[0] == dt.datetime(2024, 1, 1)
    assert df["Flag"].iloc[0] is True or df["Flag"].iloc[0] == True  # noqa: E712
    assert pd.isna(df["Region"].iloc[2])


def test_csv_lossless_is_all_text_and_blanks_are_nan(dd):
    df = dd.read_data_file(GOLDEN_CSV)
    assert all(df[c].dropna().map(type).eq(str).all() for c in df.columns)
    assert df["Zip"].iloc[0] == "01000" and df["Account"].iloc[0] == "1000"
    assert pd.isna(df["Account"].iloc[4]) and pd.isna(df["Region"].iloc[2])


def test_csv_words_that_pandas_treats_as_missing_stay_text(dd, tmp_path):
    path = make_csv(tmp_path / "w.csv", "code,v\nNA,1\nnull,2\nNone,3\nN/A,4\n,5\n")
    df = dd.read_data_file(path)
    assert df["code"].tolist()[:4] == ["NA", "null", "None", "N/A"]  # Namibia, a literal 'null' ... are data
    assert pd.isna(df["code"].iloc[4])                               # only a truly empty cell is blank


def test_infer_mode_uses_pandas_types(dd):
    df = dd.read_data_file(GOLDEN_CSV, mode="infer")
    assert df["Account"].dtype.kind == "f" and df["Zip"].dtype.kind == "i"
    xl = dd.read_data_file(GOLDEN_XLSX, mode="infer")
    assert xl["Zip"].dtype.kind == "i"  # the old behaviour, still available on request


def test_sheet_selection(dd):
    assert len(dd.read_data_file(GOLDEN_XLSX)) == 40
    assert len(dd.read_data_file(GOLDEN_XLSX, sheet_name="Second")) == 5


def test_cp1252_csv_decodes_curly_quotes_and_euro(dd):
    df = dd.read_data_file(fixture_path("golden_cp1252.csv"))
    assert df["Text"].iloc[0] == "“Quoted”" and df["Price"].iloc[0] == "€5"


def test_utf8_bom_is_not_part_of_the_first_header(dd, tmp_path):
    path = tmp_path / "bom.csv"
    path.write_bytes("﻿ID,name\n1,café\n".encode("utf-8"))
    df = dd.read_data_file(str(path))
    assert list(df.columns) == ["ID", "name"] and df["name"].iloc[0] == "café"


def test_bytes_undefined_in_cp1252_fall_back_to_latin1(dd, tmp_path):
    path = tmp_path / "odd.csv"
    path.write_bytes(b"a\n\x81x\n")  # 0x81 is undefined in cp1252
    assert dd.read_data_file(str(path))["a"].iloc[0] == "\x81x"


def test_real_parse_errors_are_not_hidden_by_the_encoding_fallback(dd, tmp_path):
    path = make_csv(tmp_path / "bad.csv", "a,b\n1,2\n1,2,3,4,5\n")
    with pytest.raises(pd.errors.ParserError):
        dd.read_data_file(path)


def test_unsupported_extension_is_rejected(dd, tmp_path):
    with pytest.raises(ValueError, match="Unsupported file format"):
        dd.read_data_file(str(tmp_path / "x.txt"))


def test_inferred_copy_converts_fully_numeric_columns_only(dd):
    df = pd.DataFrame({"n": ["1", "2", None], "mixed": ["1", "x", "3"], "d": [dt.datetime(2024, 1, 1)] * 3,
                       "b": [True, False, True], "t": ["a", "b", "c"]}, dtype=object)
    out = dd.inferred_copy(df)
    assert out["n"].dtype.kind == "f" and out["mixed"].dtype == object and out["t"].dtype == object
    assert str(out["d"].dtype).startswith("datetime64") and out["b"].dtype == bool
    assert df["n"].tolist() == ["1", "2", None]  # the original is untouched


# ------------------------------------------------------------------ multi-sheet notice

def two_sheet(path):
    wb = openpyxl.Workbook()
    wb.active.title = "First"
    wb.active.append(["a"]); wb.active.append([1])
    wb.create_sheet("Other").append(["b"])
    wb.save(path)
    return str(path)


def test_async_tools_report_extra_sheets(client):
    final, _ = post_job(client, "/upload", {"file": GOLDEN_XLSX}, {"chunk_size": "100"})
    assert final["sheet_names"] == ["Main", "Second"]
    assert "golden.xlsx has 2 sheets (Main, Second)" in final["warning"] and "'Main'" in final["warning"]


def test_single_sheet_workbooks_and_csv_have_no_warning(client, tmp_path):
    single = make_xlsx(tmp_path / "one.xlsx", ["a"], [[1]])
    for path in (single, GOLDEN_CSV):
        ext = "xlsx" if str(path).endswith("xlsx") else "csv"
        final, _ = post_job(client, "/find-duplicates", {"file": path}, {"id_column": "a" if ext == "xlsx" else "ID",
                                                                        "duplicate_columns[]": ["a" if ext == "xlsx" else "Zip"]})
        assert "warning" not in final and "sheet_names" not in final


def test_sync_tools_report_extra_sheets(client):
    resp, payload = post_form(client, "/row-filter", {"file": GOLDEN_XLSX},
                              {"conditions": json.dumps([{"column": "Name", "operator": "equals", "value": "Bob"}])})
    assert resp.status_code == 200 and payload["sheet_names"] == ["Main", "Second"]
    assert "only the first sheet" in payload["warning"]


def test_two_file_tools_name_each_workbook(client, tmp_path):
    a, b = two_sheet(tmp_path / "left_book.xlsx"), two_sheet(tmp_path / "right_book.xlsx")
    final, _ = post_job(client, "/compare-columns", {"file1": a, "file2": b})
    assert "left_book.xlsx has 2 sheets" in final["warning"] and "right_book.xlsx has 2 sheets" in final["warning"]


def test_warning_names_the_uploaded_file_not_internal_ids(client):
    final, _ = post_job(client, "/upload", {"file": GOLDEN_XLSX}, {"chunk_size": "100"})
    assert final["warning"].startswith("golden.xlsx has")


# ------------------------------------------------------------------ pivot

def test_pivot_ignores_non_numbers_and_says_so(client, tmp_path):
    path = make_csv(tmp_path / "p.csv", "g,amt\na,10\na,abc\nb,5\n")
    final, _ = post_job(client, "/generate-pivot", {"file": path},
                        {"rows[]": ["g"], "values[]": ["amt"], "aggfunc": "sum"})
    assert final["stage"] == "done", final
    assert final["ignored_non_numeric"] == {"amt": 1} and "1 in amt" in final["warning"]
    pivot = openpyxl.load_workbook(output_path(final))["Pivot Table"]
    assert 15 in [c.value for row in pivot.iter_rows() for c in row]


def test_pivot_count_works_on_text_values(client, tmp_path):
    path = make_csv(tmp_path / "p.csv", "g,label\na,x\na,y\nb,z\n")
    final, _ = post_job(client, "/generate-pivot", {"file": path},
                        {"rows[]": ["g"], "values[]": ["label"], "aggfunc": "count"})
    assert final["stage"] == "done", final
    assert "ignored_non_numeric" not in final


# ------------------------------------------------------------------ formulas on text numbers

def test_formulas_do_arithmetic_on_csv_numbers_but_output_keeps_original_text(client, tmp_path):
    path = make_csv(tmp_path / "c.csv", "id,qty\n007,5\n008,6\n")
    resp, payload = post_form(client, "/calculated-columns", {"file": path},
                              {"formula": "[qty] * 2", "new_column_name": "double"})
    assert resp.status_code == 200, payload
    rows = [[c.value for c in r] for r in openpyxl.load_workbook(sync_output(payload)).active.iter_rows(min_row=2)]
    assert rows == [["007", "5", 10], ["008", "6", 12]]  # 10, not '55'; ids keep their zeros


# ------------------------------------------------------------------ keys across file types

def test_merge_matches_excel_numbers_with_csv_text(client, tmp_path):
    left = make_xlsx(tmp_path / "l.xlsx", ["id", "l"], [[1, "a"], [2, "b"], [3, "c"]])
    right = make_csv(tmp_path / "r.csv", "id,r\n1,x\n2,y\n")
    final, _ = post_job(client, "/merge-data", {"left_file": left, "right_file": right},
                        {"left_key": "id", "right_key": "id", "join_type": "inner", "duplicate_handling": "keep_first"})
    assert final["stage"] == "done", final
    assert final["summary"]["merged_rows"] == 2


def test_merge_keeps_leading_zero_keys_distinct(client, tmp_path):
    left = make_csv(tmp_path / "l.csv", "id,l\n007,a\n7,b\n")
    right = make_csv(tmp_path / "r.csv", "id,r\n007,x\n")
    final, _ = post_job(client, "/merge-data", {"left_file": left, "right_file": right},
                        {"left_key": "id", "right_key": "id", "join_type": "inner", "duplicate_handling": "keep_first"})
    assert final["summary"]["merged_rows"] == 1  # '007' matches '007' only (7 is a different key)


def test_compare_matches_excel_numbers_with_csv_text(client, tmp_path):
    a = make_xlsx(tmp_path / "a.xlsx", ["k", "v"], [[1, "p"], [2, "q"]])
    b = make_csv(tmp_path / "b.csv", "k,v\n1,p\n2,CHANGED\n")
    final, _ = post_job(client, "/compare-data", {"file1": a, "file2": b}, {"key_columns[]": ["k"]})
    s = final["summary"]
    assert (s["common"], s["added"], s["removed"], s["changed"]) == (2, 0, 0, 1)


def test_compare_output_order_follows_file_one(client, tmp_path):
    keys = ["k%02d" % i for i in range(30)]
    a = make_csv(tmp_path / "a.csv", "k,v\n" + "".join(f"{k},1\n" for k in keys))
    b = make_csv(tmp_path / "b.csv", "k,v\n" + "".join(f"{k},1\n" for k in reversed(keys)))
    final, _ = post_job(client, "/compare-data", {"file1": a, "file2": b}, {"key_columns[]": ["k"]})
    unchanged = openpyxl.load_workbook(output_path(final))["Unchanged Rows"]
    assert [r[0].value for r in unchanged.iter_rows(min_row=2)] == keys
