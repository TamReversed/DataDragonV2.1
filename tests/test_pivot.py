"""Pivot: blank labels are kept, empty cells are honest, filters work, totals match a brute-force pivot."""
import os
import random
import shutil
from collections import OrderedDict
from queue import Queue

import openpyxl
import pytest

import datadragon
from helpers import output_path
from test_merge import write_rows

_counter = iter(range(10**9))
BLANK = "(blank)"


def run_pivot(path, rows, columns, values, aggfunc, filters=None):
    q = Queue()
    session = f"pivot_test_{next(_counter)}"
    copy = os.path.join(os.path.dirname(path), f"{session}{os.path.splitext(path)[1]}")
    shutil.copy(path, copy)  # the tool deletes its input when it finishes
    datadragon.generate_pivot_async(copy, rows, columns, values, aggfunc, filters, q, session)
    msgs = []
    while not q.empty():
        msgs.append(q.get())
    return msgs[-1]


def table_of(final):
    ws = openpyxl.load_workbook(output_path(final))["Pivot Table"]
    return [[c.value for c in row] for row in ws.iter_rows()]


def label(v):
    return BLANK if v is None else v


def number(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def aggregate(rows, value, agg):
    """Brute force. rows: dicts of raw strings/None. Mirrors the tool: sum/mean/min/max use numbers only,
    count/nunique use any non-blank value. Empty -> 0 for sum/count/nunique, None otherwise."""
    raw = [r[value] for r in rows if r[value] is not None]
    if agg == "count":
        return len(raw)
    if agg == "nunique":
        return len(set(raw))
    nums = [number(v) for v in raw if number(v) is not None]
    if not nums:
        return 0 if agg == "sum" else None
    return {"sum": sum(nums), "mean": sum(nums) / len(nums), "min": min(nums), "max": max(nums)}[agg]


def expected_table(data, row_dims, col_dim, values, agg):
    """-> {row_key_tuple_or_'Total': {column_name: value}}"""
    row_keys = list(OrderedDict.fromkeys(tuple(label(r[d]) for d in row_dims) for r in data))
    col_labels = list(OrderedDict.fromkeys(label(r[col_dim]) for r in data)) if col_dim else []
    out = {}
    for key in row_keys + ["Total"]:
        subset = data if key == "Total" else [r for r in data if tuple(label(r[d]) for d in row_dims) == key]
        cells = {}
        for value in values:
            if col_dim:
                for c in col_labels:
                    cells[f"{value}_{c}"] = aggregate([r for r in subset if label(r[col_dim]) == c], value, agg)
                cells[f"{value}_Total"] = aggregate(subset, value, agg)
            else:
                cells[value] = aggregate(subset, value, agg)
        out[key] = cells
    return out


def actual_table(final, n_row_dims):
    t = table_of(final)
    header = t[0]
    out = {}
    for row in t[1:]:
        key = "Total" if row[0] == "Total" else tuple(row[:n_row_dims])
        out[key] = dict(zip(header[n_row_dims:], row[n_row_dims:]))
    return out


@pytest.mark.parametrize("seed", range(20))
@pytest.mark.parametrize("agg", ["sum", "count", "mean", "min", "max", "nunique"])
def test_matches_brute_force(tmp_path, seed, agg):
    rng = random.Random(seed)
    n_row_dims, has_col, n_values = rng.choice([1, 2]), rng.random() < 0.6, rng.choice([1, 2])
    row_dims, col_dim, values = ["g1", "g2"][:n_row_dims], ("c" if has_col else None), ["v1", "v2"][:n_values]
    labels = ["a", "b", "c1", None]
    numbers = ["1", "2.5", "10", "-3", "abc", None]
    data = [{"g1": rng.choice(labels), "g2": rng.choice(labels), "c": rng.choice(["x", "y", None]),
             "v1": rng.choice(numbers), "v2": rng.choice(numbers)} for _ in range(rng.randint(3, 25))]
    path = write_rows(tmp_path / "p.csv", ["g1", "g2", "c", "v1", "v2"],
                      [[r["g1"], r["g2"], r["c"], r["v1"], r["v2"]] for r in data])
    final = run_pivot(path, row_dims, [col_dim] if col_dim else [], values, agg)
    assert final["stage"] == "done", final
    expected = expected_table(data, row_dims, col_dim, values, agg)
    actual = actual_table(final, n_row_dims)
    assert set(actual) == set(expected), (seed, agg)
    for key, cells in expected.items():
        for col, want in cells.items():
            got = actual[key][col]
            if want is None:
                assert got is None, (seed, agg, key, col, got)
            else:
                assert got == pytest.approx(want), (seed, agg, key, col)


def test_every_row_counts_in_the_grand_total_even_with_blank_labels(tmp_path):
    path = write_rows(tmp_path / "p.csv", ["g", "v"], [["a", "10"], ["b", "5"], [None, "100"]])
    final = run_pivot(path, ["g"], [], ["v"], "sum")
    t = table_of(final)
    assert t[0] == ["g", "v"] and [BLANK, 100] in [[r[0], r[1]] for r in t[1:]] and t[-1] == ["Total", 115]


def test_empty_combinations_stay_blank_for_mean_but_are_zero_for_sums(tmp_path):
    path = write_rows(tmp_path / "p.csv", ["g", "c", "v"], [["a", "x", "10"], ["b", "y", "20"]])
    mean = table_of(run_pivot(path, ["g"], ["c"], ["v"], "mean"))
    assert mean[0] == ["g", "v_x", "v_y", "v_Total"]
    assert mean[1] == ["a", 10, None, 10] and mean[2] == ["b", None, 20, 20]
    total = table_of(run_pivot(path, ["g"], ["c"], ["v"], "sum"))
    assert total[1] == ["a", 10, 0, 10] and total[2] == ["b", 0, 20, 20]


def test_a_blank_value_does_not_remove_the_row_from_other_values_totals(tmp_path):
    path = write_rows(tmp_path / "p.csv", ["g", "v1", "v2"], [["a", "10", None], ["a", "5", "7"]])
    t = table_of(run_pivot(path, ["g"], [], ["v1", "v2"], "sum"))
    assert t[-1] == ["Total", 15, 7]


def test_filters_apply_to_text_and_numbers(tmp_path):
    # Excel numeric cells 2023 and 2023.0 and the text cell '2023' all match the filter value "2023"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["year", "g", "v"])
    for row in ([2023, "a", 10], [2024, "a", 20], [2023.0, "b", 5], ["2023", "c", 1]):
        ws.append(row)
    ws["A5"].data_type = "s"  # a genuine text cell
    path = str(tmp_path / "p.xlsx")
    wb.save(path)
    final = run_pivot(path, ["g"], [], ["v"], "sum", {"year": "2023"})
    assert final["stage"] == "done", final
    assert {r[0]: r[1] for r in table_of(final)[1:]} == {"a": 10, "b": 5, "c": 1, "Total": 16}
    assert final["summary"]["source_rows"] == 3


def test_csv_text_is_compared_as_written(tmp_path):
    path = write_rows(tmp_path / "p.csv", ["year", "g", "v"], [["2023", "a", "10"], ["2023.0", "b", "5"]])
    final = run_pivot(path, ["g"], [], ["v"], "sum", {"year": "2023"})
    assert {r[0]: r[1] for r in table_of(final)[1:]} == {"a": 10, "Total": 10}  # '2023.0' is different text


def test_unknown_filter_column_is_a_clear_error(tmp_path):
    path = write_rows(tmp_path / "p.csv", ["g", "v"], [["a", "1"]])
    final = run_pivot(path, ["g"], [], ["v"], "sum", {"nope": "x"})
    assert final["stage"] == "error" and "nope" in final["message"]


def test_source_data_sheet_keeps_the_original_values(tmp_path):
    path = write_rows(tmp_path / "p.csv", ["g", "v"], [[None, "007"], ["a", "abc"]])
    final = run_pivot(path, ["g"], [], ["v"], "sum")
    source = [[c.value for c in r] for r in openpyxl.load_workbook(output_path(final))["Source Data"].iter_rows(min_row=2)]
    assert source == [[None, "007"], ["a", "abc"]]  # not '(blank)' and not converted to numbers
    assert final["ignored_non_numeric"] == {"v": 1}


def number_formats(final, min_row=2):
    ws = openpyxl.load_workbook(output_path(final))["Pivot Table"]
    return {c.number_format for row in ws.iter_rows(min_row=min_row, min_col=2) for c in row if c.value is not None}


def test_counts_use_a_whole_number_format_and_sums_two_decimals(tmp_path):
    path = write_rows(tmp_path / "p.csv", ["g", "v"], [["a", "1"], ["a", "2"]])
    assert number_formats(run_pivot(path, ["g"], [], ["v"], "count")) == {"#,##0"}
    assert number_formats(run_pivot(path, ["g"], [], ["v"], "sum")) == {"#,##0.00"}


def test_total_styling_follows_position_not_the_label(tmp_path):
    # A category literally called "Total ..." is ordinary data; only the margin row/column are totals.
    path = write_rows(tmp_path / "p.csv", ["g", "c", "v"], [["Total sales", "x", "1"], ["b", "x", "2"]])
    final = run_pivot(path, ["g"], ["c"], ["v"], "sum")
    ws = openpyxl.load_workbook(output_path(final))["Pivot Table"]
    assert [ws.cell(row=r, column=1).value for r in (2, 3, 4)] == ["Total sales", "b", "Total"]
    fill = lambda r, c: ws.cell(row=r, column=c).fill.start_color.rgb  # noqa: E731
    total_style = fill(4, 1)
    assert fill(2, 1) == fill(3, 1) != total_style              # "Total sales" is styled like 'b', not as a margin
    assert fill(4, 2) == fill(2, 3) == fill(3, 3) == total_style  # margin row cells and margin column cells
    assert total_style not in (fill(2, 2), fill(3, 2))          # ordinary data cells (zebra-striped) differ
