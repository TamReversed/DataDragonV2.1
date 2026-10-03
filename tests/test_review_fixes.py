"""Defects found by the independent Phase 1 review."""
import json
import os
import time

import openpyxl
import pytest

import datadragon
from helpers import fixture_path, make_csv, output_path, post_form, post_job
from test_compare import run_compare, sheet
from test_merge import merged_rows, run_merge, write_rows
from test_normalizer import column, run_normalize
from test_pivot import run_pivot, table_of


# ---- same-origin check behind a rewriting proxy
def test_allowed_origin_hosts_let_a_proxied_public_host_through(client, monkeypatch):
    data = {"formula": "1", "new_column_name": "x", "preview_only": "true"}
    headers = {"Origin": "https://app.example.com"}
    assert client.post("/calculated-columns", data=data, headers=headers).status_code == 403
    monkeypatch.setattr(datadragon, "ALLOWED_ORIGIN_HOSTS", {"app.example.com"})
    assert client.post("/calculated-columns", data=data, headers=headers).status_code == 400  # reached the route
    other = {"Origin": "https://evil.example"}
    assert client.post("/calculated-columns", data=data, headers=other).status_code == 403    # still blocked


# ---- failed jobs do not leave empty job directories behind
def test_a_failed_split_leaves_no_empty_output_directory(client, tmp_path):
    root = datadragon.app.config["OUTPUT_FOLDER"]
    before = set(os.listdir(root))
    bad = tmp_path / "bad.xlsx"
    bad.write_bytes(b"this is not a workbook")
    final, _ = post_job(client, "/upload", {"file": str(bad)}, {"chunk_size": "10"})
    assert final["stage"] == "error"
    assert set(os.listdir(root)) == before


# ---- pivot
def test_pivot_filter_on_a_blank_only_column_does_not_crash(tmp_path):
    path = write_rows(tmp_path / "p.csv", ["g", "blank_col", "v"], [["a", None, "1"], ["b", None, "2"]])
    final = run_pivot(path, ["g"], [], ["v"], "sum", {"blank_col": "x"})
    assert "accessor" not in str(final.get("message", ""))


def test_a_real_total_label_does_not_break_the_margin_row(tmp_path):
    path = write_rows(tmp_path / "p.csv", ["g", "v"], [["Total", "1"], ["b", "2"]])
    final = run_pivot(path, ["g"], [], ["v"], "sum")
    assert final["stage"] == "done", final
    assert table_of(final) == [["g", "v"], ["Total", 1], ["b", 2], ["Grand Total", 3]]


# ---- internal column names never collide with the user's columns
def test_compare_works_when_a_user_column_is_called_occ(tmp_path):
    f1 = write_rows(tmp_path / "a.csv", ["_occ", "v"], [["1", "a"], ["1", "b"]])
    f2 = write_rows(tmp_path / "b.csv", ["_occ", "v"], [["1", "a"], ["1", "CHANGED"]])
    final = run_compare(f1, f2, ["_occ"])
    assert final["stage"] == "done", final
    assert (final["summary"]["unchanged"], final["summary"]["changed"]) == (1, 1)


def test_compare_keeps_a_user_column_called_file_on_the_keyless_sheet(tmp_path):
    f1 = write_rows(tmp_path / "a.csv", ["k", "File"], [["1", "mine"], [None, "orphan"]])
    f2 = write_rows(tmp_path / "b.csv", ["k", "File"], [["1", "mine"]])
    final = run_compare(f1, f2, ["k"])
    ws = openpyxl.load_workbook(output_path(final))["Rows Without Key"]
    assert [c.value for c in ws[1]] == ["File_", "k", "File"]
    assert [c.value for c in ws[2]] == ["File 1", None, "orphan"]


def test_merge_works_when_a_user_column_is_called_like_the_indicator(tmp_path):
    left = write_rows(tmp_path / "l.csv", ["id", "__dd_merge__"], [["1", "keep me"]])
    right = write_rows(tmp_path / "r.csv", ["id", "rv"], [["1", "x"]])
    final = run_merge(left, right, "inner")
    assert final["stage"] == "done", final
    assert merged_rows(final) == [["1", "keep me", "x"]]


# ---- normalizer
def test_overflowing_numbers_are_not_converted_to_infinity(tmp_path):
    path = write_rows(tmp_path / "n.csv", ["v"], [["1e999"], ["1e3"]])
    final = run_normalize(path, {"v": "float"})
    assert column(final, "v") == ["1e999", 1000.0] and final["summary"]["total_errors"] == 1


# ---- the test-file generator no longer shares a name within one second
def test_test_file_downloads_work_back_to_back(client):
    for _ in range(2):
        assert client.get("/generate-test-file?num_rows=20").status_code == 200
