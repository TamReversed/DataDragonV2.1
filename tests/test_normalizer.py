"""Column normalizer: strict locale-aware numbers, explicit date order, nothing destroyed, formats on the right columns."""
import datetime as dt
import json
import os
import shutil
from queue import Queue

import numpy as np
import openpyxl
import pandas as pd
import pytest

import datadragon
from datadragon import detect_date_order, normalize_series, parse_date, parse_localized_number
from helpers import make_xlsx, output_path, post_form
from test_merge import write_rows

_counter = iter(range(10**9))


def run_normalize(path, types, trim=False, decimal=".", order="auto"):
    q = Queue()
    session = f"norm_test_{next(_counter)}"
    copy = os.path.join(os.path.dirname(path), f"{session}{os.path.splitext(path)[1]}")
    shutil.copy(path, copy)
    datadragon.normalize_columns_async(copy, types, trim, q, session, decimal, order)
    msgs = []
    while not q.empty():
        msgs.append(q.get())
    return msgs[-1]


def column(final, name):
    ws = openpyxl.load_workbook(output_path(final))["Normalized Data"]
    header = [c.value for c in ws[1]]
    i = header.index(name)
    return [row[i].value for row in ws.iter_rows(min_row=2)]


# ------------------------------------------------------------------ numbers
@pytest.mark.parametrize("text,expected", [
    ("1234.56", 1234.56), ("1,234.56", 1234.56), ("1,234,567", 1234567), ("-5", -5), ("+5", 5),
    ("(1,234.50)", -1234.5), ("$1,234", 1234), ("€ 1 234.5", 1234.5), ("1'234.5", 1234.5), (".5", 0.5), ("1e3", 1000.0),
    ("007", 7), ("9007199254740993", 9007199254740993), ("0.1", 0.1),
    # not numbers in dot-decimal notation
    ("2,5", None), ("1,23", None), ("1,2345", None), ("12,34,567", None), ("1.234,56", None), ("5.", None),
    ("1 2", None), ("abc", None), ("12abc", None), ("", None), ("  ", None), ("--5", None), ("$", None), ("50%", None),
])
def test_dot_decimal_numbers(text, expected):
    assert parse_localized_number(text, ".") == expected


@pytest.mark.parametrize("text,expected", [
    ("2,5", 2.5), ("1.234,56", 1234.56), ("1.234", 1234), ("1234,5", 1234.5), ("1 234,5", 1234.5), ("-0,5", -0.5),
    ("1,234.56", None), ("1.2.3", None), ("2,5,5", None),
])
def test_comma_decimal_numbers(text, expected):
    assert parse_localized_number(text, ",") == expected


def test_native_values_pass_through_and_booleans_are_not_numbers():
    assert parse_localized_number(5) == 5 and isinstance(parse_localized_number(5), int)
    assert parse_localized_number(2.5) == 2.5
    assert parse_localized_number(np.int64(7)) == 7 and isinstance(parse_localized_number(np.int64(7)), int)
    assert parse_localized_number(np.float64(1.5)) == 1.5
    assert parse_localized_number(True) is None and parse_localized_number(np.bool_(False)) is None
    assert parse_localized_number(None) is None and parse_localized_number(float("nan")) is None
    assert parse_localized_number(dt.datetime(2024, 1, 1)) is None


# ------------------------------------------------------------------ dates
@pytest.mark.parametrize("text,order,expected", [
    ("2024-03-05", "MDY", dt.datetime(2024, 3, 5)), ("2024-03-05", "DMY", dt.datetime(2024, 3, 5)),
    ("2024-03-05 14:30:15", "DMY", dt.datetime(2024, 3, 5, 14, 30, 15)), ("2024-03-05T14:30", "MDY", dt.datetime(2024, 3, 5, 14, 30)),
    ("03/04/2024", "MDY", dt.datetime(2024, 3, 4)), ("03/04/2024", "DMY", dt.datetime(2024, 4, 3)),
    ("03.04.2024", "DMY", dt.datetime(2024, 4, 3)), ("3-4-2024", "MDY", dt.datetime(2024, 3, 4)),
    ("2024/03/04", "MDY", dt.datetime(2024, 3, 4)), ("2024/03/04", "YMD", dt.datetime(2024, 3, 4)),
    ("03/04/24", "MDY", dt.datetime(2024, 3, 4)), ("03/04/70", "MDY", dt.datetime(1970, 3, 4)),
    ("13/02/2024", "DMY", dt.datetime(2024, 2, 13)), ("13/02/2024", "MDY", None),   # no month 13
    ("02/30/2024", "MDY", None), ("2024-02-30", "MDY", None), ("not a date", "MDY", None), ("", "MDY", None),
    ("5 Mar 2024", "MDY", dt.datetime(2024, 3, 5)), ("March 5, 2024", "DMY", dt.datetime(2024, 3, 5)),
    ("03/04/2024 10:05", "MDY", dt.datetime(2024, 3, 4, 10, 5)),
    (20240305, "MDY", None), (3.5, "MDY", None),
])
def test_parse_date(text, order, expected):
    assert parse_date(text, order) == expected


def test_parse_date_keeps_native_dates():
    assert parse_date(dt.datetime(2024, 1, 5, 3), "MDY") == dt.datetime(2024, 1, 5, 3)
    assert parse_date(pd.Timestamp("2024-01-05"), "MDY") == dt.datetime(2024, 1, 5)


@pytest.mark.parametrize("values,expected", [
    (["13/02/2024", "01/03/2024"], "DMY"), (["02/13/2024", "01/03/2024"], "MDY"),
    (["2024-01-05", "5 Mar 2024"], None), (["05/05/2024"], None), ([None, 3, "abc"], None),
    (["20/02/2024", "02/20/2024"], "CONFLICT"), (["01/02/2024", "2024-03-05"], "AMBIGUOUS"),
])
def test_detect_date_order(values, expected):
    if expected == "CONFLICT":
        with pytest.raises(ValueError, match="mixes day-first and month-first"):
            detect_date_order(values)
    elif expected == "AMBIGUOUS":
        with pytest.raises(ValueError, match="ambiguous.*01/02/2024"):
            detect_date_order(values)
    else:
        assert detect_date_order(values) == expected


# ------------------------------------------------------------------ per type
def series(values):
    return pd.Series(values, dtype=object)


def test_integer_never_truncates_and_keeps_big_numbers_exact():
    out, ok, bad, examples = normalize_series(series(["3", "3.0", 3.0, "3.7", 3.7, "1,234", "9007199254740993", None, "x"]),
                                              "integer")
    assert out.tolist()[:3] == [3, 3, 3] and all(isinstance(v, int) for v in out.tolist()[:3])
    assert out[3] == "3.7" and out[4] == 3.7 and out[8] == "x"             # kept exactly as they were
    assert out[5] == 1234 and out[6] == 9007199254740993 and pd.isna(out[7])
    assert (ok, bad) == (5, 3) and examples == ["3.7", "3.7", "x"]


def test_float_and_currency_convert_text_numbers_and_keep_the_rest():
    for target in ("float", "decimal", "currency", "dollar"):
        out, ok, bad, _ = normalize_series(series(["$1,234.50", 7, "abc", None]), target)
        assert out.tolist()[:2] == [1234.5, 7.0] and out[2] == "abc" and pd.isna(out[3]) and (ok, bad) == (2, 1)


def test_comma_decimal_locale():
    out, _, bad, _ = normalize_series(series(["2,5", "1.234,56", "1,234.56"]), "float", decimal_separator=",")
    assert out.tolist() == [2.5, 1234.56, "1,234.56"] and bad == 1


def test_percentage_divides_only_when_there_is_a_percent_sign():
    out, ok, bad, examples = normalize_series(series(["50%", "12.5 %", 0.5, "0.25", 50, "abc%", None]), "percentage")
    assert out[:5].tolist() == [0.5, 0.125, 0.5, 0.25, 50.0]   # only the two '%' strings were divided by 100
    assert out[5] == "abc%" and pd.isna(out[6]) and (ok, bad, examples) == (5, 1, ["abc%"])


def test_percentage_in_comma_notation():
    out, _, bad, _ = normalize_series(series([" 12,5 %", "0,25", "50%"]), "percentage", decimal_separator=",")
    assert out.tolist() == [0.125, 0.25, 0.5] and bad == 0


def test_percentage_text_without_a_sign_is_a_plain_number():
    out, _, _, _ = normalize_series(series(["0.25", "7"]), "percentage")
    assert out.tolist() == [0.25, 7.0]   # never rescaled: 7 stays 7 (not 0.07)


def test_text_keeps_blanks_blank_and_trims_on_request():
    out, ok, _, _ = normalize_series(series(["  a ", 5, 2.5, None, True]), "text", trim_whitespace=True)
    assert out.tolist()[:3] == ["a", "5", "2.5"] and pd.isna(out[3]) and out[4] == "True" and ok == 4
    untrimmed, *_ = normalize_series(series(["  a "]), "text")
    assert untrimmed[0] == "  a "


def test_boolean_words_numbers_and_unknowns():
    out, ok, bad, _ = normalize_series(series(["Yes", "no", "T", "0", "1", 1, 0, True, "2.5", "maybe", "2,5", None]), "boolean")
    assert out.tolist()[:9] == [True, False, True, False, True, True, False, True, True]
    assert out[9] == "maybe" and out[10] == "2,5"       # '2,5' is not a number in dot notation, so it is left alone
    assert pd.isna(out[11]) and (ok, bad) == (9, 2)


def test_date_conversion_keeps_unparseable_values_and_native_dates():
    native = dt.datetime(2024, 2, 1, 8)
    out, ok, bad, _ = normalize_series(series(["03/04/2024", native, "garbage", None, "2024-05-06"]), "date", date_order="DMY")
    assert out.tolist()[:3] == [dt.datetime(2024, 4, 3), native, "garbage"] and pd.isna(out[3])
    assert out[4] == dt.datetime(2024, 5, 6) and (ok, bad) == (3, 1)


def test_keep_original_and_unknown_types_change_nothing():
    original = series(["a", 1, None])
    for target in ("keep_original", "no_such_type"):
        out, ok, bad, _ = normalize_series(original, target)
        assert out.equals(original) and (ok, bad) == (0, 0)


def test_the_input_series_is_not_modified():
    original = series(["2,5", "3"])
    normalize_series(original, "float", decimal_separator=",")
    assert original.tolist() == ["2,5", "3"]


# ------------------------------------------------------------------ the tool
def test_unconvertible_values_are_reported_with_a_hint_and_left_unchanged(tmp_path):
    path = write_rows(tmp_path / "n.csv", ["amount"], [["1,234.5"], ["2,5"], ["10"]])
    final = run_normalize(path, {"amount": "float"})
    assert final["stage"] == "done", final
    assert column(final, "amount") == [1234.5, "2,5", 10]
    assert final["summary"] == {"total_columns": 1, "total_errors": 1, "total_transformed": 2}
    assert "left exactly as they were" in final["warning"] and "'2,5'" in final["warning"]
    assert "comma as the decimal separator" in final["warning"]


def test_the_comma_hint_is_only_shown_when_it_could_help(tmp_path):
    path = write_rows(tmp_path / "n.csv", ["v"], [["abc"], ["1"]])
    assert "comma as the decimal separator" in run_normalize(path, {"v": "float"}, decimal=".")["warning"]
    path = write_rows(tmp_path / "n.csv", ["v"], [["abc"], ["1"]])
    comma = run_normalize(path, {"v": "float"}, decimal=",")["warning"]
    assert "'abc'" in comma and "comma as the decimal separator" not in comma


def test_clean_runs_have_no_warning(tmp_path):
    path = write_rows(tmp_path / "n.csv", ["amount"], [["1"], ["2"]])
    assert "warning" not in run_normalize(path, {"amount": "integer"})


def test_comma_locale_through_the_tool(tmp_path):
    path = write_rows(tmp_path / "n.csv", ["v"], [["2,5"], ["1.234,56"]])
    assert column(run_normalize(path, {"v": "float"}, decimal=","), "v") == [2.5, 1234.56]


def test_ambiguous_dates_stop_the_job_without_saving(tmp_path):
    path = write_rows(tmp_path / "n.csv", ["d"], [["01/02/2024"], ["2024-03-05"]])
    final = run_normalize(path, {"d": "date"})
    assert final["stage"] == "error" and "ambiguous" in final["message"] and "01/02/2024" in final["message"]
    assert "download_url" not in final


def test_the_column_itself_can_settle_the_date_order(tmp_path):
    path = write_rows(tmp_path / "n.csv", ["d"], [["13/02/2024"], ["01/03/2024"], ["2024-03-05"]])
    assert column(run_normalize(path, {"d": "date"}), "d") == [dt.datetime(2024, 2, 13), dt.datetime(2024, 3, 1),
                                                              dt.datetime(2024, 3, 5)]


def test_explicit_order_resolves_ambiguity(tmp_path):
    path = write_rows(tmp_path / "n.csv", ["d"], [["01/02/2024"]])
    assert column(run_normalize(path, {"d": "date"}, order="MDY"), "d") == [dt.datetime(2024, 1, 2)]
    path = write_rows(tmp_path / "n.csv", ["d"], [["01/02/2024"]])
    assert column(run_normalize(path, {"d": "date"}, order="DMY"), "d") == [dt.datetime(2024, 2, 1)]


def test_excel_date_cells_are_never_reinterpreted(tmp_path):
    path = make_xlsx(tmp_path / "n.xlsx", ["d"], [[dt.datetime(2024, 1, 2)], [dt.datetime(2024, 12, 1)]])
    assert column(run_normalize(path, {"d": "date"}), "d") == [dt.datetime(2024, 1, 2), dt.datetime(2024, 12, 1)]


def test_formats_land_on_the_columns_they_belong_to(tmp_path):
    # The old code numbered columns by their position in the request, so formats went to the wrong columns.
    path = write_rows(tmp_path / "n.csv", ["a", "b", "c", "d"], [["0.5", "x", "12.5", "2024-01-02"]])
    final = run_normalize(path, {"zzz": "float", "d": "date", "c": "currency", "a": "percentage"})
    ws = openpyxl.load_workbook(output_path(final))["Normalized Data"]
    formats = {ws.cell(row=1, column=i).value: ws.cell(row=2, column=i).number_format for i in range(1, 5)}
    assert formats["b"] == "General"
    assert "$" in formats["c"] and "%" in formats["a"] and formats["d"] == "mm/dd/yyyy"


def test_other_columns_pass_through_untouched(tmp_path):
    path = write_rows(tmp_path / "n.csv", ["id", "v"], [["007", "1"], ["008", None]])
    final = run_normalize(path, {"v": "integer"})
    assert column(final, "id") == ["007", "008"] and column(final, "v") == [1, None]


def test_route_validates_the_new_options(client, tmp_path):
    path = write_rows(tmp_path / "n.csv", ["v"], [["1"]])
    for field, value in (("decimal_separator", ";"), ("date_order", "XYZ")):
        resp, payload = post_form(client, "/normalize-columns", {"file": path},
                                  {"column_types": json.dumps({"v": "float"}), field: value})
        assert resp.status_code == 400 and field in payload["error"]
    resp, payload = post_form(client, "/normalize-columns", {"file": path},
                              {"column_types": json.dumps({"v": "float"}), "decimal_separator": ",", "date_order": "DMY"})
    assert resp.status_code == 200 and payload["success"]
