"""Phase B, part 2: merge on several keys and anti-joins, splitting by a column's values, pivot upgrades.

Each result is checked against the inputs (counts that must add up) or against a brute-force answer.
"""
import io
import json
import random
import zipfile
from collections import Counter

import openpyxl
import pytest

from helpers import make_xlsx, output_path, post_job
from test_merge import merged_rows, run_merge, write_rows
from test_pivot import run_pivot, table_of


# ---------------------------------------------------------------- merge: several keys, anti-joins

def keyed(tmp_path, name, header, rows):
    return write_rows(tmp_path / name, header, rows)


@pytest.mark.parametrize("seed", range(25))
def test_two_key_join_matches_nested_loops(tmp_path, seed):
    rng = random.Random(seed)
    regions, years = ["N", "S", None], ["2023", "2024", None]
    left = [(rng.choice(regions), rng.choice(years), f"L{i}") for i in range(rng.randint(0, 9))]
    right = [(rng.choice(regions), rng.choice(years), f"R{i}") for i in range(rng.randint(0, 9))]
    left_path = keyed(tmp_path, "l.csv", ["region", "year", "lv"], left)
    right_path = keyed(tmp_path, "r.csv", ["region", "year", "rv"], right)

    def partners(row, other):                 # a blank in any key column never matches
        return [o for o in other if None not in row[:2] and row[:2] == o[:2]]

    for how in ("inner", "left", "left_only", "right_only"):
        expected = Counter()
        if how in ("inner", "left"):
            for row in left:
                found = partners(row, right)
                for other in found:
                    expected[(row[2], other[2])] += 1
                if not found and how == "left":
                    expected[(row[2], None)] += 1
        elif how == "left_only":
            expected = Counter((row[2],) for row in left if not partners(row, right))
        else:
            expected = Counter((row[2],) for row in right if not partners(row, left))

        final = run_merge(left_path, right_path, join=how, left_key=["region", "year"], right_key=["region", "year"])
        assert final["stage"] == "done", final
        sheet = openpyxl.load_workbook(output_path(final))["Merged Data"]
        header = [c.value for c in sheet[1]]
        wanted = [name for name in ("lv", "rv") if name in header]
        got = Counter(tuple(row[header.index(name)] for name in wanted) for row in merged_rows(final))
        assert got == expected, (how, left, right)


def test_anti_joins_and_inner_join_account_for_every_row(tmp_path):
    left = [(str(i % 7), f"L{i}") for i in range(30)]
    right = [(str(k), f"R{k}") for k in (0, 1, 2, 9)]
    left_path = keyed(tmp_path, "l.csv", ["id", "lv"], left)
    right_path = keyed(tmp_path, "r.csv", ["id", "rv"], right)
    only_left = merged_rows(run_merge(left_path, right_path, join="left_only"))
    only_right = merged_rows(run_merge(left_path, right_path, join="right_only"))
    inner = merged_rows(run_merge(left_path, right_path, join="inner"))
    assert len(only_left) + len(inner) == len(left)            # right keys are unique here: one row per match
    assert [row[0] for row in only_right] == ["9"]
    assert {row[0] for row in only_left} == {"3", "4", "5", "6"}
    assert all(len(row) == 2 for row in only_left)             # an anti-join keeps only its own side's columns


def test_merge_route_takes_extra_key_pairs(client, tmp_path):
    left = keyed(tmp_path, "l.csv", ["region", "year", "lv"], [("N", "2023", "a"), ("N", "2024", "b"), ("S", "2023", "c")])
    right = keyed(tmp_path, "r.csv", ["area", "yr", "rv"], [("N", "2024", "x"), ("S", "2024", "y")])
    final, _ = post_job(client, "/merge-data", {"left_file": left, "right_file": right},
                        {"left_key": "region", "right_key": "area", "extra_left_keys[]": ["year"],
                         "extra_right_keys[]": ["yr"], "join_type": "inner"})
    assert final["stage"] == "done", final
    rows = merged_rows(final)
    assert len(rows) == 1 and "b" in rows[0] and "x" in rows[0]


def test_merge_rejects_uneven_or_unknown_keys(tmp_path):
    left = keyed(tmp_path, "l.csv", ["a", "b"], [("1", "2")])
    right = keyed(tmp_path, "r.csv", ["a", "b"], [("1", "2")])
    for left_key, right_key in ((["a", "b"], ["a"]), (["a", "nope"], ["a", "b"])):
        with pytest.raises(Exception) as caught:
            final = run_merge(left, right, left_key=left_key, right_key=right_key)
            assert final["stage"] == "error", final
            raise ValueError(final["message"])
        assert str(caught.value)


# ---------------------------------------------------------------- splitter: one file per value

def split_by(client, path, column, **extra):
    final, _ = post_job(client, "/upload", {"file": path}, {"chunk_size": "100", "split_column": column, **extra})
    return final


def test_split_by_column_makes_one_file_per_value_and_loses_no_row(client, tmp_path):
    rows = [("N", 1), ("S", 2), ("N", 3), (None, 4), ("E/W", 5), ("S", 6), ("  ", 7), ("e w", 8)]
    path = make_xlsx(tmp_path / "in.xlsx", ["Region", "n"], rows)
    final = split_by(client, path, "Region", base_filename="by")
    assert final["stage"] == "done", final
    archive = zipfile.ZipFile(output_path(final))
    names = [name for name in archive.namelist() if name.endswith(".xlsx")]
    assert names == ["by_N.xlsx", "by_S.xlsx", "by_blank.xlsx", "by_E_W.xlsx", "by_e_w_2.xlsx"]    # first-appearance order
    seen = []
    for name in names:
        sheet = openpyxl.load_workbook(io.BytesIO(archive.read(name))).active
        body = [[c.value for c in row] for row in sheet.iter_rows(min_row=2)]
        assert len({(row[0] or "").strip() for row in body}) == 1              # one value per file
        seen += [row[1] for row in body]
    assert sorted(seen) == [1, 2, 3, 4, 5, 6, 7, 8] and final["total_rows"] == 8 and final["num_files"] == 5
    record = json.loads(archive.read("_DataDragon_Log.json"))
    assert "split_by_column" in json.dumps(record)


def test_split_by_column_refuses_too_many_values_and_unknown_columns(client, tmp_path, monkeypatch):
    import datadragon
    monkeypatch.setattr(datadragon, "MAX_SPLIT_VALUES", 3)
    path = make_xlsx(tmp_path / "in.xlsx", ["k", "n"], [(str(i), i) for i in range(5)])
    final = split_by(client, path, "k")
    assert final["stage"] == "error" and "5 different values" in final["message"]
    final = split_by(client, path, "missing")
    assert final["stage"] == "error" and "missing" in final["message"]


# ---------------------------------------------------------------- pivot upgrades

SALES = [("N", "2024-01-15", "10"), ("N", "2024-02-01", "30"), ("S", "2024-04-20", "5"),
         ("S", "2023-12-31", "15"), ("N", "2024-05-05", "40"), ("S", "", "7"), ("N", "soon", "3")]


def sales(tmp_path):
    return write_rows(tmp_path / "sales.csv", ["region", "day", "amt"], SALES)


def pivot(path, rows, columns, values, aggfunc, **more):
    import datadragon
    import os
    import shutil
    from queue import Queue
    q = Queue()
    session = f"pivot_b2_{random.getrandbits(40)}"
    copy = os.path.join(os.path.dirname(path), f"{session}.csv")
    shutil.copy(path, copy)
    datadragon.generate_pivot_async(copy, rows, columns, values, aggfunc, None, q, session, **more)
    final = None
    while not q.empty():
        final = q.get()
    return final


def test_several_aggregations_equal_the_separate_pivots(tmp_path):
    path = sales(tmp_path)
    together = table_of(pivot(path, ["region"], [], ["amt"], "sum", extra_aggfuncs=["count", "mean", "max"]))
    assert together[0] == ["region", "sum_amt", "count_amt", "mean_amt", "max_amt"]
    for position, agg in enumerate(["sum", "count", "mean", "max"], 1):
        alone = table_of(run_pivot(path, ["region"], [], ["amt"], agg))
        assert [row[position] for row in together[1:]] == [row[1] for row in alone[1:]], agg
        assert [row[0] for row in together] == [row[0] for row in alone]


def test_several_aggregations_with_a_column_field(tmp_path):
    path = sales(tmp_path)
    together = table_of(pivot(path, ["region"], ["day"], ["amt"], "sum", extra_aggfuncs=["count"],
                              date_groups={"day": "year"}))
    by_sum = table_of(pivot(path, ["region"], ["day"], ["amt"], "sum", date_groups={"day": "year"}))
    by_count = table_of(pivot(path, ["region"], ["day"], ["amt"], "count", date_groups={"day": "year"}))
    width = len(by_sum[0]) - 1
    assert together[0] == ["region"] + [f"sum_{h}" for h in by_sum[0][1:]] + [f"count_{h}" for h in by_count[0][1:]]
    assert [row[1:1 + width] for row in together[1:]] == [row[1:] for row in by_sum[1:]]
    assert [row[1 + width:] for row in together[1:]] == [row[1:] for row in by_count[1:]]


@pytest.mark.parametrize("unit, expected", [
    ("year", {"2023": 15, "2024": 85, "(blank)": 7, "(not a date)": 3}),
    ("quarter", {"2023-Q4": 15, "2024-Q1": 40, "2024-Q2": 45, "(blank)": 7, "(not a date)": 3}),
    ("month", {"2023-12": 15, "2024-01": 10, "2024-02": 30, "2024-04": 5, "2024-05": 40, "(blank)": 7, "(not a date)": 3}),
])
def test_dates_group_into_periods_and_the_total_is_unchanged(tmp_path, unit, expected):
    final = pivot(sales(tmp_path), ["day"], [], ["amt"], "sum", date_groups={"day": unit})
    table = table_of(final)
    assert {row[0]: row[1] for row in table[1:-1]} == expected
    assert table[-1] == ["Total", 110]
    assert "not dates" in final["warning"] and "1 in day" in final["warning"]


def test_ambiguous_dates_are_refused_not_guessed(tmp_path):
    path = write_rows(tmp_path / "d.csv", ["day", "amt"], [("03/04/2024", "1"), ("05/06/2024", "2")])
    with pytest.raises(Exception) as caught:
        final = pivot(path, ["day"], [], ["amt"], "sum", date_groups={"day": "month"})
        assert final["stage"] == "error", final
        raise ValueError(final["message"])
    assert "ambiguous" in str(caught.value)


def test_percentages_of_row_column_and_grand_total(tmp_path):
    path = sales(tmp_path)
    plain = table_of(pivot(path, ["region"], ["day"], ["amt"], "sum", date_groups={"day": "year"}))
    header, body = plain[0], plain[1:]
    grand = body[-1][-1]
    assert grand == 110
    for show_as in ("row", "column", "total"):
        shares = table_of(pivot(path, ["region"], ["day"], ["amt"], "sum", date_groups={"day": "year"}, show_as=show_as))
        assert shares[0] == header
        for r, row in enumerate(body):
            for c in range(1, len(row)):
                base = {"row": row[-1], "column": body[-1][c], "total": grand}[show_as]
                expected = None if not base else row[c] / base
                got = shares[1 + r][c]
                assert (got is None and expected is None) or got == pytest.approx(expected), (show_as, r, c)
    by_row = table_of(pivot(path, ["region"], ["day"], ["amt"], "sum", date_groups={"day": "year"}, show_as="row"))
    assert all(sum(row[1:-1]) == pytest.approx(1) and row[-1] == pytest.approx(1) for row in by_row[1:])


def test_percentages_without_a_column_field(tmp_path):
    table = table_of(pivot(sales(tmp_path), ["region"], [], ["amt"], "sum", show_as="total"))
    assert {row[0]: round(row[1], 6) for row in table[1:]} == {"N": round(83 / 110, 6), "S": round(27 / 110, 6), "Total": 1}


@pytest.mark.parametrize("more, words", [
    ({"show_as": "row"}, "column field"),
    ({"show_as": "total", "extra_aggfuncs": ["mean"]}, "Sum and Count"),
    ({"extra_aggfuncs": ["mode"]}, "Unknown aggregation"),
    ({"date_groups": {"amt": "month"}}, "not a row or column field"),
    ({"date_groups": {"region": "week"}}, "year, quarter or month"),
])
def test_pivot_options_that_cannot_work_say_why(tmp_path, more, words):
    with pytest.raises(Exception) as caught:
        final = pivot(sales(tmp_path), ["region"], [], ["amt"], "sum", **more)
        assert final["stage"] == "error", final
        raise ValueError(final["message"])
    assert words in str(caught.value)


def test_pivot_route_passes_the_new_options(client, tmp_path):
    final, _ = post_job(client, "/generate-pivot", {"file": sales(tmp_path)},
                        {"rows[]": ["day"], "values[]": ["amt"], "aggfunc": "sum", "extra_aggfuncs[]": ["count"],
                         "date_group_columns[]": ["day"], "date_group_units[]": ["quarter"]})
    assert final["stage"] == "done", final
    table = table_of(final)
    assert table[0] == ["day", "sum_amt", "count_amt"] and table[-1] == ["Total", 110, 7]
    assert final["summary"]["aggfunc"] == "sum, count"
    summary = openpyxl.load_workbook(output_path(final))["Summary"]
    assert ["Dates Grouped", "day by quarter"] in [[c.value for c in row] for row in summary.iter_rows()]
