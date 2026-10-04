"""User regular expressions: bounded time, honest errors, and Find & Replace touching only what matches."""
import json
import time

import openpyxl
import pytest

import datadragon
import datadragon_regex as rx
from helpers import make_csv, make_xlsx, post_form, post_job, sync_output
from test_merge import write_rows

EVIL = r"(a|aa)+$"
EVIL_TEXT = "a" * 40 + "!"


# ------------------------------------------------------------------ the module
def test_compile_rejects_empty_long_and_invalid_patterns():
    with pytest.raises(rx.PatternError, match="empty"):
        rx.compile_pattern("")
    with pytest.raises(rx.PatternError, match="too long"):
        rx.compile_pattern("a" * (rx.MAX_PATTERN_LENGTH + 1))
    assert rx.compile_pattern("a" * rx.MAX_PATTERN_LENGTH)
    with pytest.raises(rx.PatternError, match="Invalid pattern"):
        rx.compile_pattern("(unclosed")
    with pytest.raises(rx.PatternError):
        rx.compile_pattern(None)


def test_compiled_patterns_are_cached_per_case_setting():
    assert rx.compile_pattern("x+") is rx.compile_pattern("x+")
    assert rx.compile_pattern("x+") is not rx.compile_pattern("x+", ignore_case=True)
    assert rx.compile_pattern("abc", ignore_case=True).fullmatch("ABC")
    assert not rx.compile_pattern("abc").fullmatch("ABC")


def test_a_catastrophic_pattern_is_stopped_quickly():
    started = time.time()
    with pytest.raises(rx.PatternTooComplex):
        rx.substitute(rx.compile_pattern(EVIL), "x", EVIL_TEXT, rx.Budget())
    assert time.time() - started < 1
    with pytest.raises(rx.PatternTooComplex):
        rx.fullmatch(rx.compile_pattern(EVIL), EVIL_TEXT, rx.Budget())


def test_the_total_budget_is_enforced():
    budget = rx.Budget(seconds=0)
    with pytest.raises(rx.PatternTooComplex, match="too long"):
        budget.timeout()
    assert 0 < rx.Budget().timeout() <= rx.CELL_TIMEOUT


def test_substitute_counts_matches_and_supports_templates_and_functions():
    compiled = rx.compile_pattern(r"(\d+)-(\d+)")
    assert rx.substitute(compiled, r"\2-\1", "10-20 and 3-4", rx.Budget()) == ("20-10 and 4-3", 2)
    assert rx.substitute(rx.compile_pattern("a"), lambda m: r"\1", "aaa", rx.Budget()) == (r"\1\1\1", 3)
    assert rx.substitute(rx.compile_pattern("zzz"), "x", "abc", rx.Budget()) == ("abc", 0)


def test_a_bad_replacement_template_is_a_pattern_error():
    with pytest.raises(rx.PatternError, match="Invalid replacement"):
        rx.substitute(rx.compile_pattern("a"), r"\1", "a", rx.Budget())


def test_fullmatch_means_the_whole_text():
    compiled = rx.compile_pattern("A|B")
    assert rx.fullmatch(compiled, "A", rx.Budget()) and not rx.fullmatch(compiled, "Axx", rx.Budget())
    assert not rx.fullmatch(rx.compile_pattern(r"\d{5}"), "123456789", rx.Budget())


# ------------------------------------------------------------------ Find & Replace route
def run_fr(client, path, **data):
    return post_form(client, "/find-replace", {"file": path}, data)


def read_column(payload, index=0):
    ws = openpyxl.load_workbook(sync_output(payload)).active
    return [row[index].value for row in ws.iter_rows(min_row=2)]


def test_numbers_that_do_not_match_stay_numbers_and_blanks_stay_blank(client, tmp_path):
    path = make_xlsx(tmp_path / "n.xlsx", ["v"], [[5], [7], [None], ["a5"]], text_cols=())
    resp, payload = run_fr(client, path, find_text="5", replace_text="X", column="v")
    assert resp.status_code == 200
    assert read_column(payload) == ["X", 7, None, "aX"]       # 7 is still a number, the blank is still blank
    assert payload["replacements_made"] == 2 and payload["rows_affected"] == 2


def test_replacements_made_counts_every_match(client, tmp_path):
    path = write_rows(tmp_path / "c.csv", ["v"], [["aaa"], ["ba"], ["zzz"]])
    resp, payload = run_fr(client, path, find_text="a", replace_text="-", column="v", case_sensitive="true")
    assert read_column(payload) == ["---", "b-", "zzz"] and payload["replacements_made"] == 4 and payload["rows_affected"] == 2


def test_literal_search_never_treats_the_replacement_as_a_template(client, tmp_path):
    path = write_rows(tmp_path / "l.csv", ["v"], [["a.b"], ["axb"]])
    for case in ("true", "false"):                              # the case-insensitive path used to treat \1 as a group
        resp, payload = run_fr(client, path, find_text="a.b", replace_text=r"\1", column="v", case_sensitive=case)
        assert resp.status_code == 200, payload
        assert read_column(payload) == [r"\1", "axb"]


def test_literal_search_escapes_regex_characters(client, tmp_path):
    path = write_rows(tmp_path / "l.csv", ["v"], [["(x)"], ["x"]])
    resp, payload = run_fr(client, path, find_text="(X)", replace_text="y", column="v", case_sensitive="false")
    assert read_column(payload) == ["y", "x"]


def test_regex_mode_supports_backreferences(client, tmp_path):
    path = write_rows(tmp_path / "r.csv", ["v"], [["10-20"], ["no"]])
    resp, payload = run_fr(client, path, find_text=r"(\d+)-(\d+)", replace_text=r"\2-\1", column="v", use_regex="true")
    assert read_column(payload) == ["20-10", "no"]


def test_whole_cell_matches_numbers_by_how_they_read(client, tmp_path):
    path = make_xlsx(tmp_path / "w.xlsx", ["v"], [[5], [5.0], [50], ["5"]], text_cols=())
    resp, payload = run_fr(client, path, find_text="5", replace_text="five", column="v", match_whole_cell="true")
    assert read_column(payload) == ["five", "five", 50, "five"]


def test_whole_cell_regex_is_a_full_match(client, tmp_path):
    path = write_rows(tmp_path / "w.csv", ["v"], [["12345"], ["123456789"], ["x12345"]])
    resp, payload = run_fr(client, path, find_text=r"\d{5}", replace_text="ZIP", column="v", use_regex="true",
                           match_whole_cell="true")
    assert read_column(payload) == ["ZIP", "123456789", "x12345"]


def test_all_columns_mode_visits_every_column(client, tmp_path):
    path = write_rows(tmp_path / "a.csv", ["a", "b"], [["x", "x"], ["y", None]])
    resp, payload = run_fr(client, path, find_text="x", replace_text="Z")
    ws = openpyxl.load_workbook(sync_output(payload)).active
    assert [[c.value for c in r] for r in ws.iter_rows(min_row=2)] == [["Z", "Z"], ["y", None]]


@pytest.mark.parametrize("find,message", [("(unclosed", "Invalid pattern"), ("a" * 501, "too long")])
def test_bad_patterns_are_a_400_with_a_message(client, tmp_path, find, message):
    path = write_rows(tmp_path / "b.csv", ["v"], [["x"]])
    resp, payload = run_fr(client, path, find_text=find, replace_text="y", column="v", use_regex="true")
    assert resp.status_code == 400 and message in payload["error"]


def test_a_bad_replacement_group_is_a_400(client, tmp_path):
    path = write_rows(tmp_path / "b.csv", ["v"], [["x"]])
    resp, payload = run_fr(client, path, find_text="x", replace_text=r"\5", column="v", use_regex="true")
    assert resp.status_code == 400 and "Invalid replacement" in payload["error"]


def test_a_catastrophic_pattern_is_a_422_and_leaves_no_output(client, tmp_path):
    path = write_rows(tmp_path / "e.csv", ["v"], [[EVIL_TEXT]])
    started = time.time()
    resp, payload = run_fr(client, path, find_text=EVIL, replace_text="x", column="v", use_regex="true")
    assert resp.status_code == 422 and "too complex" in payload["error"] and time.time() - started < 3
    assert "download_url" not in payload


def test_the_request_budget_ends_a_slow_run(client, tmp_path, monkeypatch):
    path = write_rows(tmp_path / "s.csv", ["v"], [["a"], ["b"]])
    original = rx.Budget
    monkeypatch.setattr(rx, "Budget", lambda: original(seconds=0))   # a request that has no time at all
    resp, payload = run_fr(client, path, find_text="a", replace_text="x", column="v")
    assert resp.status_code == 422 and "too long" in payload["error"]


# ------------------------------------------------------------------ validation 'pattern' rule
def test_validation_pattern_is_bounded_and_reports_clearly(client, tmp_path):
    path = write_rows(tmp_path / "v.csv", ["code"], [[EVIL_TEXT]])
    rules = [{"column": "code", "type": "pattern", "value": EVIL}]
    started = time.time()
    final, _ = post_job(client, "/validate-data", {"file": path}, {"validation_rules": json.dumps(rules)})
    assert final["stage"] == "error" and "too complex" in final["message"] and time.time() - started < 5


def test_validation_pattern_invalid_regex_is_reported(client, tmp_path):
    path = write_rows(tmp_path / "v.csv", ["code"], [["x"]])
    rules = [{"column": "code", "type": "pattern", "value": "(oops"}]
    final, _ = post_job(client, "/validate-data", {"file": path}, {"validation_rules": json.dumps(rules)})
    assert final["stage"] == "error" and "Invalid pattern" in final["message"]
