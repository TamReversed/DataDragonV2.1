"""Normalise tool output into stable JSON for characterization (golden) tests."""
import datetime as dt
import io
import json
import os
import re
import zipfile

import openpyxl

EXPECTED = os.path.join(os.path.dirname(os.path.abspath(__file__)), "expected")
VOLATILE_KEY = re.compile(r"(url|path|session|timestamp|elapsed|duration|created|_at$|filename|output_file)", re.I)
TS = re.compile(r"\d{8}_\d{6}")
EPOCH = re.compile(r"\b1[5-9]\d{8}\b")


def _clean_str(s):
    return EPOCH.sub("<EPOCH>", TS.sub("<TS>", s))


def scrub(o):
    if isinstance(o, dict):
        return {str(k): scrub(v) for k, v in o.items() if not VOLATILE_KEY.search(str(k))}
    if isinstance(o, (list, tuple)):
        return [scrub(v) for v in o]
    if isinstance(o, str):
        return _clean_str(o)
    if isinstance(o, float):
        return round(o, 6)
    return o


def _cell(c):
    v = c.value
    if isinstance(v, (dt.datetime, dt.date)):
        v = v.isoformat()
    return f"{c.data_type}:{v!r}"


LOG_SHEET = "_DataDragon_Log"
LOG_VOLATILE = {"Timestamp (UTC)": "<TIMESTAMP>", "Version": "<VERSION>"}      # change with every run / commit


def mask_log(rows):
    """Rows of the log sheet with the values that change from run to run replaced by placeholders."""
    return [[row[0], f"s:{LOG_VOLATILE[row[0][3:-1]]!r}"] if len(row) == 2 and row[0][3:-1] in LOG_VOLATILE else row
            for row in rows]


def describe_xlsx(data):
    wb = openpyxl.load_workbook(io.BytesIO(data))
    sheets = {ws.title: [[_cell(c) for c in row] for row in ws.iter_rows()] for ws in wb.worksheets}
    if LOG_SHEET in sheets:
        sheets[LOG_SHEET] = mask_log(sheets[LOG_SHEET])
    return sheets


def describe_bytes(name, data):
    lower = name.lower()
    if lower.endswith(".xlsx"):
        return describe_xlsx(data)
    if lower.endswith(".json"):
        loaded = json.loads(data.decode("utf-8"))
        if os.path.basename(lower).endswith("log.json") and isinstance(loaded, dict):
            loaded.update({key: placeholder for key, placeholder in LOG_VOLATILE.items() if key in loaded})
        return scrub(loaded)
    if lower.endswith(".csv"):
        return data.decode("utf-8", errors="replace").splitlines()
    return {"binary_bytes_present": len(data) > 0}


def describe_file(path):
    with open(path, "rb") as f:
        data = f.read()
    if path.lower().endswith(".zip"):
        zf = zipfile.ZipFile(io.BytesIO(data))
        return {_clean_str(n): describe_bytes(n, zf.read(n)) for n in sorted(zf.namelist())}
    return describe_bytes(path, data)


def check(config, name, actual):
    """Compare `actual` to tests/golden/expected/<name>.json (or rewrite it with --update-golden)."""
    actual = json.loads(json.dumps(actual, default=str, sort_keys=True))
    path = os.path.join(EXPECTED, f"{name}.json")
    if config.getoption("--update-golden"):
        with open(path, "w", encoding="utf-8") as f:
            json.dump(actual, f, indent=1, sort_keys=True, ensure_ascii=False)
            f.write("\n")
        return
    assert os.path.exists(path), f"missing snapshot {name}.json (generate with --update-golden)"
    with open(path, encoding="utf-8") as f:
        expected = json.load(f)
    assert actual == expected, f"snapshot {name} differs"
